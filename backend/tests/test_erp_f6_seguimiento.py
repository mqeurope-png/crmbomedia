"""ERP-F6 — seguimiento de pedidos: campos nuevos, vista, exportación y
sincronización con la hoja de Drive (cuenta de servicio, incremental, sin
borrar filas ajenas ni pisar celdas manuales).
"""
from __future__ import annotations

import io
import json
import logging
from collections.abc import Generator
from datetime import UTC, datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

import app.main  # noqa: F401
from app.db.base import Base
from app.db.session import get_session
from app.erp import seguimiento as core
from app.erp.api.seguimiento import _rows
from app.erp.drive_sheets import DriveSyncError, sync_to_sheet
from app.erp.models import (
    Carrier,
    ErpDriveSyncRow,
    InvoiceStatus,
    Order,
    OrderSource,
    TransportStatus,
)
from app.main import app
from app.models.crm import Company
from tests._test_helpers import auth_headers, seed_test_users

# --- fixtures ----------------------------------------------------------------------


@pytest.fixture()
def session_factory() -> Generator[sessionmaker, None, None]:
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False}, poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    with factory() as seed:
        seed_test_users(seed)
        seed.commit()
    yield factory
    Base.metadata.drop_all(engine)


@pytest.fixture()
def http(session_factory) -> Generator[TestClient, None, None]:
    def override() -> Generator[Session, None, None]:
        with session_factory() as s:
            yield s
    app.dependency_overrides[get_session] = override
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _carrier(s: Session, name: str) -> str:
    c = Carrier(name=name, code=name.lower())
    s.add(c)
    s.flush()
    return c.id


def _order(
    s: Session,
    number: str,
    *,
    cliente: str | None = None,
    source: OrderSource = OrderSource.MANUAL,
    carrier_id: str | None = None,
    tracking: str | None = None,
    serial: str | None = None,
    whiterip: str | None = None,
    origin: str | None = None,
    factura: str | None = None,
    external_id: str | None = None,
    transport: TransportStatus = TransportStatus.NOT_SHIPPED,
    invoice: InvoiceStatus = InvoiceStatus.NOT_INVOICED,
    placed: str = "2026-09-01",
    notes: str | None = None,
) -> Order:
    company_id = None
    if cliente:
        comp = Company(name=cliente)
        s.add(comp)
        s.flush()
        company_id = comp.id
    o = Order(
        order_number=number, external_source=source, external_id=external_id,
        company_id=company_id, carrier_id=carrier_id, tracking_number=tracking,
        serial_number=serial, whiterip_license=whiterip, shipping_origin=origin,
        factusol_invoice_number=factura, transport_status=transport,
        invoice_status=invoice, notes=notes,
        placed_at=datetime.fromisoformat(placed).replace(tzinfo=UTC),
    )
    s.add(o)
    s.flush()
    return o


def _rows_for_sync(s: Session) -> list[dict]:
    return core.filter_rows(_rows(s), en_curso=True, sort="fecha", direction="asc")


# --- Parte A: campos nuevos --------------------------------------------------------


def test_order_new_fields_persist(session_factory, http) -> None:
    with session_factory() as s:
        oid = _order(s, "BOP-100001", cliente="Hirsch Armbänder GmbH").id
        s.commit()
    r = http.patch(f"/api/erp/orders/{oid}/seguimiento", json={
        "serial_number": "FBAP12613200249",
        "whiterip_license": "4829",
        "shipping_origin": "SAT",
    }, headers=auth_headers(http, "pedidos"))
    assert r.status_code == 200, r.text
    assert r.json()["serial_number"] == "FBAP12613200249"
    assert r.json()["whiterip_license"] == "4829"
    assert r.json()["shipping_origin"] == "SAT"
    # Persisten y salen en la ficha del pedido.
    r = http.get(f"/api/erp/orders/{oid}", headers=auth_headers(http, "user"))
    assert r.json()["serial_number"] == "FBAP12613200249"
    assert r.json()["whiterip_license"] == "4829"
    assert r.json()["shipping_origin"] == "SAT"
    # El nº de serie es texto LIBRE: en el Excel real también lleva notas.
    r = http.patch(f"/api/erp/orders/{oid}/seguimiento", json={
        "serial_number": "FINALIZADO, + RMA TRANSPORTE",
    }, headers=auth_headers(http, "pedidos"))
    assert r.json()["serial_number"] == "FINALIZADO, + RMA TRANSPORTE"
    # Vaciar un campo lo deja a NULL; el resto no se toca.
    r = http.patch(f"/api/erp/orders/{oid}/seguimiento", json={"whiterip_license": ""},
                   headers=auth_headers(http, "pedidos"))
    assert r.json()["whiterip_license"] is None
    assert r.json()["shipping_origin"] == "SAT"
    # Ver no basta para editar.
    r = http.patch(f"/api/erp/orders/{oid}/seguimiento", json={"shipping_origin": "OFI"},
                   headers=auth_headers(http, "user"))
    assert r.status_code == 403
    # Los orígenes configurables salen en ajustes (defaults del Excel real).
    r = http.get("/api/erp/settings", headers=auth_headers(http, "user"))
    assert r.json()["shipping_origins"] == ["SAT", "OFI", "TER", "directo",
                                           "INSITU", "ADR", "MAD"]
    r = http.patch("/api/erp/settings",
                   json={"shipping_origins": ["SAT", "OFI", "directo Conrexx"]},
                   headers=auth_headers(http, "admin"))
    assert r.json()["shipping_origins"] == ["SAT", "OFI", "directo Conrexx"]


def test_orden_column_content_goes_to_observations(session_factory, http) -> None:
    # «Orden» (37 valores en 7.743 filas, todos comentarios) NO es un campo:
    # su contenido va a las observaciones del pedido.
    with session_factory() as s:
        oid = _order(s, "BOP-100002", notes="Pedido urgente").id
        s.commit()
    r = http.patch(f"/api/erp/orders/{oid}/seguimiento", json={
        "orden": "ALBARÁN: RECOGE EL CLIENTE - SIN ENVÍO",
    }, headers=auth_headers(http, "pedidos"))
    assert r.status_code == 200, r.text
    assert r.json()["notes"] == "Pedido urgente\nALBARÁN: RECOGE EL CLIENTE - SIN ENVÍO"
    with session_factory() as s:
        o = s.get(Order, oid)
        assert "RECOGE EL CLIENTE" in o.notes
        assert not hasattr(o, "orden")          # no existe columna «orden»
    detail = http.get(f"/api/erp/orders/{oid}", headers=auth_headers(http, "user")).json()
    assert "orden" not in detail


# --- Parte B: vista ----------------------------------------------------------------


def test_seguimiento_list_filters_and_sorts(session_factory, http) -> None:
    with session_factory() as s:
        ups, mrw = _carrier(s, "UPS"), _carrier(s, "MRW")
        _order(s, "BOP-200001", cliente="Zeta SL", source=OrderSource.WOOCOMMERCE,
               carrier_id=ups, factura="5-260001", placed="2026-09-03")
        _order(s, "BOP-200002", cliente="Alfa SL", carrier_id=mrw,
               factura="1-260002", placed="2026-09-02", origin="SAT")
        _order(s, "BOP-200003", cliente="Beta SL", carrier_id=ups,
               placed="2026-09-01", origin="OFI")
        s.commit()
    headers = auth_headers(http, "user")
    r = http.get("/api/erp/seguimiento", headers=headers)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["total"] == 3
    assert body["columns"] == core.SEGUIMIENTO_COLUMNS
    # Orden por defecto: fecha desc.
    assert [i["order_number"] for i in body["items"]] == [
        "BOP-200001", "BOP-200002", "BOP-200003"]
    # La fila lleva lo que pinta la tabla (y el enlace via id).
    top = body["items"][0]
    assert top["cliente"] == "Zeta SL"
    assert top["vendedor"] == "WEB"                  # pedido web
    assert top["transportista"] == "UPS"
    assert top["serie"] == 5                          # del nº de factura 5-260001
    assert top["empresa"] == "Streamtec"
    assert top["factura"] == "5-260001"
    assert top["id"]
    # Filtros: transportista, serie, vendedor, origen, estado.
    r = http.get("/api/erp/seguimiento?transportista=UPS", headers=headers)
    assert {i["order_number"] for i in r.json()["items"]} == {"BOP-200001", "BOP-200003"}
    r = http.get("/api/erp/seguimiento?serie=1", headers=headers)
    assert [i["order_number"] for i in r.json()["items"]] == ["BOP-200002"]
    r = http.get("/api/erp/seguimiento?vendedor=WEB", headers=headers)
    assert [i["order_number"] for i in r.json()["items"]] == ["BOP-200001"]
    r = http.get("/api/erp/seguimiento?origen=ofi", headers=headers)
    assert [i["order_number"] for i in r.json()["items"]] == ["BOP-200003"]
    r = http.get("/api/erp/seguimiento?estado=facturado", headers=headers)
    assert {i["order_number"] for i in r.json()["items"]} == {"BOP-200001", "BOP-200002"}
    r = http.get("/api/erp/seguimiento?desde=2026-09-02&hasta=2026-09-02", headers=headers)
    assert [i["order_number"] for i in r.json()["items"]] == ["BOP-200002"]
    # Orden por cliente asc (misma infraestructura que E3-A-fix1).
    r = http.get("/api/erp/seguimiento?sort=cliente&dir=asc", headers=headers)
    assert [i["cliente"] for i in r.json()["items"]] == ["Alfa SL", "Beta SL", "Zeta SL"]
    # Sin sesión → 401.
    assert http.get("/api/erp/seguimiento").status_code == 401


def test_seguimiento_search_by_serial_and_tracking(session_factory, http) -> None:
    with session_factory() as s:
        _order(s, "BOP-300001", cliente="Uno SL", serial="FBAP12613200249",
               tracking="1Z999AA10123456784")
        _order(s, "BOP-300002", cliente="Dos SL", serial="ADO12611204959")
        _order(s, "BOP-300003", cliente="Tres SL", factura="5-260777")
        s.commit()
    headers = auth_headers(http, "user")
    for q, expected in [
        ("FBAP12613", ["BOP-300001"]),               # nº de serie
        ("1z999aa10", ["BOP-300001"]),               # tracking, sin mayúsculas
        ("ADO12611204959", ["BOP-300002"]),
        ("BOP-300003", ["BOP-300003"]),              # nº de pedido/albarán
        ("5-260777", ["BOP-300003"]),                # nº de factura
        ("dos sl", ["BOP-300002"]),                  # cliente
    ]:
        r = http.get(f"/api/erp/seguimiento?q={q}", headers=headers)
        assert [i["order_number"] for i in r.json()["items"]] == expected, q
    r = http.get("/api/erp/seguimiento?q=NOEXISTE", headers=headers)
    assert r.json()["items"] == []


def test_seguimiento_default_shows_open_orders(session_factory, http) -> None:
    with session_factory() as s:
        _order(s, "BOP-400001", cliente="Abierto SL")                   # en curso
        _order(s, "BOP-400002", cliente="Enviado SL",
               transport=TransportStatus.IN_TRANSIT)                    # en curso
        _order(s, "BOP-400003", cliente="Cerrado SL",
               transport=TransportStatus.DELIVERED,
               invoice=InvoiceStatus.INVOICED_BY_ERP, factura="1-260100")
        ext = _order(s, "BOP-400004", cliente="Externo SL")
        ext.externally_processed_at = datetime.now(UTC)
        s.commit()
    headers = auth_headers(http, "user")
    # Por defecto: SOLO los en curso — la parte de arriba del Excel.
    r = http.get("/api/erp/seguimiento", headers=headers)
    assert {i["order_number"] for i in r.json()["items"]} == {"BOP-400001", "BOP-400002"}
    assert all(i["en_curso"] for i in r.json()["items"])
    # Entregado+facturado y externalizado salen con en_curso=false.
    r = http.get("/api/erp/seguimiento?en_curso=false", headers=headers)
    assert r.json()["total"] == 4
    by_num = {i["order_number"]: i for i in r.json()["items"]}
    assert by_num["BOP-400003"]["en_curso"] is False
    assert by_num["BOP-400003"]["estado"] == "facturado"
    assert by_num["BOP-400004"]["en_curso"] is False
    # Enviado pero sin facturar sigue en curso (aún hay trabajo).
    assert by_num["BOP-400002"]["estado"] == "enviado"


# --- Parte C: hoja de Drive --------------------------------------------------------


class FakeSheet:
    """Hoja en memoria con la misma interfaz que `GoogleSheetsClient`."""

    def __init__(self, grid: list[list[str]]) -> None:
        self.grid = [list(r) for r in grid]

    def get_values(self) -> list[list[str]]:
        return [list(r) for r in self.grid]

    def update_cells(self, updates: list[tuple[int, int, str]]) -> None:
        for row, col, value in updates:
            r = self.grid[row - 1]
            while len(r) <= col:
                r.append("")
            r[col] = value

    def insert_rows_at(self, row: int, count: int) -> None:
        for _ in range(count):
            self.grid.insert(row - 1, [""] * len(core.SEGUIMIENTO_COLUMNS))

    def write_rows(self, start_row: int, rows: list[list[str]]) -> None:
        for i, values in enumerate(rows):
            self.grid[start_row - 1 + i] = list(values)


HEADER = list(core.SEGUIMIENTO_COLUMNS)


def test_drive_sync_updates_only_changed_rows(session_factory) -> None:
    with session_factory() as s:
        _order(s, "BOP-500001", cliente="Uno SL", origin="SAT")
        s.commit()
    sheet = FakeSheet([HEADER])
    with session_factory() as s:
        summary = sync_to_sheet(s, sheet, _rows_for_sync(s))
    assert summary["appended_rows"] == 1
    # ERP-F6-fix2: se escribe el número DESNUDO (formato de Bart), no la
    # referencia con prefijo.
    assert sheet.grid[1][11] == "500001"              # bajo la cabecera
    assert sheet.grid[1][4] == "SAT"
    # Sin cambios → NO se escribe nada (incremental de verdad).
    with session_factory() as s:
        summary = sync_to_sheet(s, sheet, _rows_for_sync(s))
    assert summary["updated_cells"] == 0
    assert summary["appended_rows"] == 0
    assert summary["conflicts"] == []
    # Cambia SOLO el tracking en BoHub → se actualiza SOLO esa celda.
    with session_factory() as s:
        o = s.scalar(select(Order).where(Order.order_number == "BOP-500001"))
        o.tracking_number = "1Z999AA10123456784"
        s.commit()
    before = [list(r) for r in sheet.grid]
    with session_factory() as s:
        summary = sync_to_sheet(s, sheet, _rows_for_sync(s))
    assert summary["updated_cells"] == 1
    assert sheet.grid[1][13] == "1Z999AA10123456784"
    for col in range(len(HEADER)):
        if col != 13:
            assert sheet.grid[1][col] == before[1][col]
    # La foto guardada refleja lo escrito (base del siguiente incremental).
    with session_factory() as s:
        snap = s.scalar(select(ErpDriveSyncRow))
        assert json.loads(snap.last_values_json)[13] == "1Z999AA10123456784"


def test_drive_sync_never_deletes_unknown_rows(session_factory) -> None:
    # La hoja real: sección de en-curso arriba, separador «^^^^», cabecera
    # repetida e histórico. Nada de eso es de BoHub y nada se toca.
    manual_top = ["BO", "1/9/2026", "CLIENTE MANUAL", "BART", "SAT", "MRW",
                  "", "", "", "Impresora vieja", "", "MANUAL-EXCEL-1", "", "",
                  "", "", "EL PRIMER ENVÍO HA LLEGADO ROTO"]
    separator = ["^^^^ Aquí arriba pedidos que están en curso"]
    hist = ["BO", "2/1/2020", "HISTORICO SL", "WEB", "SAT", "UPS",
            "", "", "", "", "", "BOP-000111", "1-200999", "", "", "", ""]
    sheet = FakeSheet([HEADER, manual_top, separator, list(HEADER), hist])
    with session_factory() as s:
        _order(s, "BOP-500010", cliente="Nuevo SL")
        s.commit()
        summary = sync_to_sheet(s, sheet, _rows_for_sync(s))
    assert summary["appended_rows"] == 1
    # El pedido nuevo entra justo bajo la PRIMERA cabecera (sección en curso),
    # con su número desnudo (ERP-F6-fix2).
    assert sheet.grid[1][11] == "500010"
    # Y todo lo que ya había sigue exactamente donde estaba, sin borrar nada.
    assert sheet.grid[2] == manual_top
    assert sheet.grid[3][0].startswith("^^^^")
    assert sheet.grid[4] == list(HEADER)
    assert sheet.grid[5] == hist
    assert len(sheet.grid) == 6


def test_drive_sync_does_not_overwrite_manual_cell(session_factory) -> None:
    with session_factory() as s:
        _order(s, "BOP-500020", cliente="Uno SL", origin="SAT",
               tracking="TRACK-1")
        s.commit()
        sheet = FakeSheet([HEADER])
        sync_to_sheet(s, sheet, _rows_for_sync(s))
    # Bart corrige el tracking A MANO en la hoja…
    sheet.grid[1][13] = "CORREGIDO A MANO"
    # …y además escribe una nota en «Orden» (columna que BoHub jamás toca).
    sheet.grid[1][16] = "OJO: RECOGE EL CLIENTE"
    # BoHub también cambia el tracking → conflicto: NO se pisa y se avisa.
    with session_factory() as s:
        o = s.scalar(select(Order).where(Order.order_number == "BOP-500020"))
        o.tracking_number = "TRACK-2"
        o.notes = "esto no debe llegar a la columna Orden"
        s.commit()
        summary = sync_to_sheet(s, sheet, _rows_for_sync(s))
    assert sheet.grid[1][13] == "CORREGIDO A MANO"
    assert len(summary["conflicts"]) == 1
    conflict = summary["conflicts"][0]
    assert conflict["order_number"] == "BOP-500020"
    assert conflict["column"] == "Tracking"
    assert conflict["sheet_value"] == "CORREGIDO A MANO"
    assert conflict["bohub_value"] == "TRACK-2"
    # «Orden» intacta aunque el pedido tenga observaciones.
    assert sheet.grid[1][16] == "OJO: RECOGE EL CLIENTE"
    # Una celda vacía sí se rellena (no es contenido manual).
    with session_factory() as s:
        o = s.scalar(select(Order).where(Order.order_number == "BOP-500020"))
        o.serial_number = "FBAP12613200249"
        s.commit()
        summary = sync_to_sheet(s, sheet, _rows_for_sync(s))
    assert sheet.grid[1][14] == "FBAP12613200249"
    # Si la hoja no tiene la cabecera esperada, no se escribe NADA.
    with session_factory() as s, pytest.raises(DriveSyncError):
        sync_to_sheet(s, FakeSheet([["cualquier", "cosa"]]), _rows_for_sync(s))


def test_view_works_without_drive_credentials(session_factory, http) -> None:
    with session_factory() as s:
        _order(s, "BOP-600001", cliente="Sin Drive SL")
        s.commit()
    # La vista funciona igual sin Drive configurado…
    r = http.get("/api/erp/seguimiento", headers=auth_headers(http, "user"))
    assert r.status_code == 200, r.text
    assert r.json()["total"] == 1
    assert r.json()["drive"]["configured"] is False
    assert r.json()["drive"]["service_account_email"] is None
    # …y la exportación también.
    r = http.get("/api/erp/seguimiento/export", headers=auth_headers(http, "user"))
    assert r.status_code == 200
    # El botón de sincronizar avisa de que falta configurarlo (409, no 500).
    r = http.post("/api/erp/seguimiento/drive-sync", headers=auth_headers(http, "pedidos"))
    assert r.status_code == 409
    assert r.json()["detail"]["code"] == "drive_not_configured"


# --- Parte D: exportación ----------------------------------------------------------


def test_export_xlsx_respects_filters_and_column_order(session_factory, http) -> None:
    with session_factory() as s:
        ups, mrw = _carrier(s, "UPS"), _carrier(s, "MRW")
        _order(s, "BOP-700001", cliente="Uno SL", carrier_id=ups, origin="SAT",
               serial="FBAP1", whiterip="4829", factura="5-260050",
               placed="2026-09-04", notes="nota del pedido")
        _order(s, "BOP-700002", cliente="Dos SL", carrier_id=mrw, placed="2026-09-03")
        s.commit()
    r = http.get("/api/erp/seguimiento/export?transportista=UPS",
                 headers=auth_headers(http, "user"))
    assert r.status_code == 200, r.text
    assert "seguimiento_pedidos_" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content), read_only=True).active
    grid = [list(row) for row in ws.iter_rows(values_only=True)]
    # Cabecera EXACTA del Excel de Bart, en su orden.
    assert list(grid[0]) == core.SEGUIMIENTO_COLUMNS
    # Solo la fila filtrada (transportista=UPS).
    assert len(grid) == 2
    row = grid[1]
    assert row[0] == "ST"                              # Empresa (código corto)
    assert row[1] == "4/9/2026"                        # fecha d/m/yyyy
    assert row[2] == "Uno SL"
    assert row[4] == "SAT"
    assert row[5] == "UPS"
    assert row[11] == "BOP-700001"
    assert row[12] == "5-260050"
    assert row[14] == "FBAP1"
    assert row[15] == "4829"
    assert row[16] == "nota del pedido"                # Orden = observaciones
    # Sin filtro salen las dos, con el orden de la vista.
    r = http.get("/api/erp/seguimiento/export", headers=auth_headers(http, "user"))
    ws = load_workbook(io.BytesIO(r.content), read_only=True).active
    assert sum(1 for _ in ws.iter_rows()) == 3


# --- Parte C: credenciales ---------------------------------------------------------

_FAKE_SA = {
    "type": "service_account",
    "project_id": "bohub-drive",
    "client_email": "bohub-seguimiento@bohub-drive.iam.gserviceaccount.com",
    "private_key": (
        "-----BEGIN PRIVATE KEY-----\nSECRETO-MARCADOR-XYZZY-99\n"
        "-----END PRIVATE KEY-----\n"
    ),
}


def test_service_account_credentials_never_logged(session_factory, http, caplog) -> None:
    marker = "SECRETO-MARCADOR-XYZZY-99"
    with caplog.at_level(logging.DEBUG):
        r = http.patch("/api/erp/settings", json={
            "drive_service_account_json": json.dumps(_FAKE_SA),
            "drive_spreadsheet_id": "1AbCdEfGhIjKlMnOpQrStUvWxYz",
        }, headers=auth_headers(http, "admin"))
        assert r.status_code == 200, r.text
        got = http.get("/api/erp/settings", headers=auth_headers(http, "user"))
        sync = http.post("/api/erp/seguimiento/drive-sync",
                         headers=auth_headers(http, "pedidos"))
    # La clave privada no aparece en NINGÚN log ni en ninguna respuesta.
    assert marker not in caplog.text
    assert marker not in r.text
    assert marker not in got.text
    assert marker not in sync.text
    # Se expone SOLO el client_email (para compartir la hoja) + el estado.
    body = got.json()
    assert body["drive_configured"] is True
    assert body["drive_service_account_email"] == _FAKE_SA["client_email"]
    assert body["drive_spreadsheet_id"] == "1AbCdEfGhIjKlMnOpQrStUvWxYz"
    assert "drive_service_account_json" not in body
    # En la base está CIFRADO (no en claro).
    with session_factory() as s:
        from app.erp.models import ERP_SETTINGS_SINGLETON_ID, ErpSettings
        cfg = s.get(ErpSettings, ERP_SETTINGS_SINGLETON_ID)
        assert marker not in (cfg.drive_service_account_json_encrypted or "")
    # Un JSON inválido se rechaza sin volcar su contenido en el error.
    r = http.patch("/api/erp/settings", json={
        "drive_service_account_json": "{esto-no-es-json " + marker,
    }, headers=auth_headers(http, "admin"))
    assert r.status_code == 400
    assert marker not in r.text
    # "" borra las credenciales.
    r = http.patch("/api/erp/settings", json={"drive_service_account_json": ""},
                   headers=auth_headers(http, "admin"))
    assert r.json()["drive_configured"] is False
