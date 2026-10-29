"""La conciliación se traga los extractos de verdad (28/10/2026).

La pantalla nunca se había usado: los dos extractos reales de Bart no
entraban. El de Belfius viene en neerlandés («no se encontró la fila de
columnas»); el de Sabadell trae la cuenta como CCC antiguo («El extracto es de
la cuenta (sin IBAN en cabecera)…») y parte cada movimiento en dos filas, y la
segunda, sin fecha, tumbaba el fichero entero («fila 139: fecha vacía»).

Aquí van FIGURANTES con la misma estructura que los reales (los extractos de
verdad no se suben al repositorio): preámbulo de cuatro filas con la cuenta en
CCC y las columnas `factura`/`pedido` rotuladas en la primera fila; filas de
continuación con el pagador; cabecera en neerlandés con la cuenta en cada fila
en dos formatos; filas antiguas sin fecha de operación; columnas vacías por
detrás. Los IBAN son de ejemplo.
"""
from __future__ import annotations

import csv
import io
from collections.abc import Generator
from datetime import datetime
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

import app.main  # noqa: F401
from app.db.session import get_session
from app.erp.bank import service
from app.erp.bank.parsing import cuenta_casa, leer_anotacion, read_statement
from app.erp.models import BankMovement, BankReconciliation
from app.erp.models.orders import Order, OrderSource
from app.main import app
from app.models.crm import Base
from tests._test_helpers import auth_headers, seed_test_users
from tests.test_erp_bank_reconciliation import _fac, _fake

#: Streamtec en Sabadell: el IBAN dado de alta y el CCC que trae el extracto.
STREAMTEC_IBAN = "ES91 2100 0418 4502 0005 1332"
STREAMTEC_CCC = "2100-0418-45-0200051332"
#: MQ Europe en Belfius: IBAN y el número belga antiguo de las filas viejas.
MQ_IBAN = "BE68 5390 0754 7034"
MQ_ANTIGUO = "539-0075470-34"


def _libro(filas: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for fila in filas:
        ws.append(fila)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _d(y: int, m: int, d: int) -> datetime:
    return datetime(y, m, d)


def _sabadell(datos: list[list]) -> bytes:
    """Como «Streamtec movimientos SABADELL …xlsx»: cuatro filas de preámbulo
    (la cuenta en CCC, y `factura`/`pedido` rotuladas en la PRIMERA fila), la
    de columnas en la quinta, cuatro filas vacías y los datos; columnas
    vacías por detrás hasta la 256."""
    relleno = [None] * 247
    return _libro([
        ["Cuenta: ", STREAMTEC_CCC, None, None, None, None, None, "factura", "pedido", *relleno],
        ["Divisa: ", "EUR"],
        ["Titular:", "STREAMTEC S.L"],
        ["Selección:"],
        ["FECHA OPER", "CONCEPTO", "FECHA VALOR", "IMPORTE", "SALDO",
         "REFERENCIA 1", "REFERENCIA 2"],
        [], [], [], [],
        *datos,
    ])


BELFIUS_COLUMNAS = [
    "Rekening", "Boekingsdatum", None, "Transactienummer", "Rekening tegenpartij",
    "Naam tegenpartij bevat", "Straat en nummer", "Postcode en plaats", "Transactie",
    "Valutadatum", "Bedrag", "Devies", "BIC", "Landcode", "Comments", None,
    "CHECKEO -BART", 7878.0, "factura", "pedido",
]


def _belfius_fila(
    fecha: datetime | None, importe: float | None, nombre: str, transactie: str, *,
    cuenta: str | None = MQ_IBAN, numero: float | None = None,
    valuta: datetime | str | None = None, factura: object = None, pedido: object = None,
) -> list:
    fila: list = [None] * len(BELFIUS_COLUMNAS)
    fila[0] = cuenta
    fila[1] = fecha
    fila[2] = 29.0
    fila[3] = numero
    fila[4] = "DE89 3704 0044 0532 0130 00"
    fila[5] = nombre
    fila[8] = transactie
    fila[9] = valuta or fecha
    fila[10] = importe
    fila[11] = "EUR"
    fila[18] = factura
    fila[19] = pedido
    return fila


def _belfius(datos: list[list]) -> bytes:
    """Como «2 MQ EUROPE Belfius extractos ok 2B.xlsx»: cabecera en
    neerlandés en la primera fila, una fila vacía, y la cuenta propia en cada
    fila (IBAN o el número belga antiguo)."""
    return _libro([BELFIUS_COLUMNAS, [], *datos])


# --- infraestructura -----------------------------------------------------------


@pytest.fixture()
def session_factory() -> Generator[sessionmaker, None, None]:
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False}, poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    with factory() as seed:
        seed_test_users(seed)
    yield factory
    Base.metadata.drop_all(engine)


@pytest.fixture()
def http(session_factory) -> Generator[TestClient, None, None]:
    def override() -> Generator[Session, None, None]:
        with session_factory() as s:
            yield s
    app.dependency_overrides[get_session] = override
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


@pytest.fixture()
def cuentas(session_factory) -> dict[str, str]:
    """Las cuentas como se dieron de alta el 28/10: con su serie y su
    contrapartida de cobro de FACTUSOL."""
    with session_factory() as s:
        st = service.create_account(s, {"name": "Sabadell Streamtec", "iban": STREAMTEC_IBAN,
                                        "serie": 5})
        mq = service.create_account(s, {"name": "Belfius MQ Europe", "iban": MQ_IBAN, "serie": 2})
        st.contrapartida_codigo, mq.contrapartida_codigo = "8", "2"
        s.commit()
        return {"streamtec": st.id, "mq": mq.id}


def _importar(session_factory, contenido: bytes, nombre: str = "extracto.xlsx",
              account_id: str | None = None) -> dict:
    with session_factory() as s:
        return service.import_statement(s, content=contenido, filename=nombre,
                                        account_id=account_id, user_id="u-bart")


def _movimiento(session_factory, fila: int) -> BankMovement:
    with session_factory() as s:
        mov = s.scalar(select(BankMovement).where(BankMovement.source_row == fila))
        assert mov is not None, f"no hay movimiento de la fila {fila}"
        s.expunge(mov)
        return mov


def _conciliaciones(session_factory, movement_id: str) -> list[BankReconciliation]:
    with session_factory() as s:
        rows = list(s.scalars(select(BankReconciliation).where(
            BankReconciliation.movement_id == movement_id)))
        for r in rows:
            s.expunge(r)
        return rows


SABADELL_DATOS = [
    # fila 10
    [_d(2026, 8, 24), "ABONO TRANSFERENCIA DE ROSA", _d(2026, 8, 24), 70.18, 6025.43, None,
     "ROSA BOSCH SALA", "5-260082", "BOP-099919"],
    # fila 11: el pagador cortado en el concepto; la fila 12, sin fecha, lo trae entero.
    [_d(2026, 8, 18), "ABONO TRANSFERENCIA DE CLARA", _d(2026, 8, 18), 145.2, 955.25,
     271815809.0, None, "5-260083", "BOP-099916"],
    [None, None, None, None, 956.46, None, "CLARA FIGUERAS M"],
    # fila 13: vacía de verdad (se salta sin decir nada)
    [],
    # fila 14: una anotación con nota: hay que mirarla
    [_d(2026, 8, 17), "ABONO TRANSFERENCIA DE MARTA", _d(2026, 8, 17), 408.48, 1483.51, None,
     "MARTA MOLINS", "5-260074 (cobro 408,48 vs factura 420,74 — confirmar)"],
    # fila 15: un pago con algo en la columna factura: no es un cobro, no se casa
    [_d(2026, 8, 14), "SEGUROS SOCIALES REG GENERAL", _d(2026, 8, 14), -802.34, 1210.05,
     260701.0, None, "5-260001"],
    # fila 16: sin anotar: lo que falta por conciliar
    [_d(2026, 8, 10), "ABONO TRANSFERENCIA DE SOLITIUM SL", _d(2026, 8, 10), 3623.95, 2012.39,
     None, "SOLITIUM SL"],
    # fila 17: años atrás, con la fecha como texto
    ["28/06/2012", "TRANSFERENCIA SOCIO EJEMPLO L", "28/06/2012", 2000.0, 2145.42,
     "0001259030"],
]


# --- 1.2 · Sabadell sin elegir la cuenta ---------------------------------------


def test_sabadell_se_importa_reconociendo_la_cuenta_por_su_ccc(
    session_factory, cuentas,
) -> None:
    assert cuenta_casa(STREAMTEC_CCC, STREAMTEC_IBAN)
    assert not cuenta_casa(STREAMTEC_CCC, MQ_IBAN)
    r = _importar(session_factory, _sabadell(SABADELL_DATOS))
    assert r["ok"], r
    assert (r["account_id"], r["formato"]) == (cuentas["streamtec"], "Sabadell")
    assert (r["imported"], r["continuaciones"], r["no_entendidas_total"]) == (6, 1, 0)
    viejo = _movimiento(session_factory, 17)
    assert viejo.fecha_oper.isoformat() == "2012-06-28"
    assert viejo.referencia1 == "0001259030"
    assert _movimiento(session_factory, 11).referencia1 == "271815809"   # no «271815809.0»


def test_sabadell_con_la_cuenta_de_otro_banco_no_se_importa_a_ciegas(
    session_factory, cuentas,
) -> None:
    r = _importar(session_factory, _sabadell(SABADELL_DATOS), account_id=cuentas["mq"])
    assert (r["ok"], r["code"]) == (False, "iban_mismatch")
    assert STREAMTEC_CCC in r["detail"]
    with session_factory() as s:
        service.delete_account(s, cuentas["streamtec"])
    r = _importar(session_factory, _sabadell(SABADELL_DATOS))
    assert (r["ok"], r["code"]) == (False, "unknown_account")
    assert r["detail"].startswith(f"El extracto es de la cuenta {STREAMTEC_CCC}, que no está")


# --- 1.3 · la fila de continuación -----------------------------------------------


def test_la_fila_sin_fecha_se_une_a_la_anterior_y_el_pagador_acaba_en_el_movimiento(
    session_factory, cuentas,
) -> None:
    _importar(session_factory, _sabadell(SABADELL_DATOS))
    mov = _movimiento(session_factory, 11)
    assert mov.referencia2 == "CLARA FIGUERAS M"       # venía en la fila 12
    assert mov.payer_name == "CLARA FIGUERAS M"        # no «CLARA», cortado
    assert float(mov.saldo) == 955.25                   # el saldo es el del movimiento
    with session_factory() as s:
        assert s.scalar(select(func.count(BankMovement.id)).where(
            BankMovement.source_row == 12)) == 0


# --- 1.1 · Belfius en neerlandés ---------------------------------------------------


def test_belfius_con_cabeceras_en_neerlandes_se_importa(session_factory, cuentas) -> None:
    contenido = _belfius([
        # fila 3
        _belfius_fila(_d(2026, 10, 19), 216.59, "Werkstatt Muster gGmbH",
                      "STORTING VAN DE89 Muster 14.10.26 NR. 2-526126 216,59 REF. : 100000001",
                      numero=187.0, factura="2-526126", pedido="ARTISJ-9696"),
        # fila 4: continuación con solo la cuenta (las hay en el real)
        _belfius_fila(None, None, None, None),
        # fila 5: el número de cuenta en el formato belga antiguo
        _belfius_fila(_d(2013, 3, 13), -520.8, "Voorbeeld Technology BV", "inv 1000000001",
                      cuenta=MQ_ANTIGUO, numero=92.0),
        # fila 6: 2010, sin Rekening ni Boekingsdatum: vale la fecha valor
        _belfius_fila(None, 1026.0, "EJEMPLO SA", "FACT. 500001 EJEMPLO\xa0", cuenta=None,
                      valuta=_d(2010, 1, 13), factura=526043.0),
    ])
    r = _importar(session_factory, contenido, "2 MQ EUROPE Belfius extractos ok 2B.xlsx")
    assert r["ok"], r
    assert (r["account_id"], r["formato"]) == (cuentas["mq"], "Belfius")
    assert (r["imported"], r["continuaciones"], r["no_entendidas_total"]) == (3, 1, 0)
    mov = _movimiento(session_factory, 3)
    assert (mov.referencia1, mov.referencia2, mov.payer_name) == (
        "187", "Werkstatt Muster gGmbH", "Werkstatt Muster gGmbH")
    assert "NR. 2-526126" in mov.concepto
    viejo = _movimiento(session_factory, 6)
    assert viejo.fecha_oper.isoformat() == "2010-01-13"
    assert viejo.concepto == "FACT. 500001 EJEMPLO"
    # Sin serie en la anotación: la de la cuenta (2).
    assert [c.numero for c in _conciliaciones(session_factory, viejo.id)] == ["2-526043"]


def test_una_fila_que_dice_ser_de_otra_cuenta_no_entra_en_esta(
    session_factory, cuentas,
) -> None:
    contenido = _belfius([
        _belfius_fila(_d(2026, 10, 19), 100.0, "A", "STORTING VAN A"),
        _belfius_fila(_d(2026, 10, 18), 200.0, "B", "STORTING VAN B",
                      cuenta="BE71 0961 2345 6769"),
    ])
    r = _importar(session_factory, contenido)
    assert r["imported"] == 1
    assert [(x["fila"], x["motivo"]) for x in r["no_entendidas"]] == [
        (4, "es de otra cuenta (BE71 0961 2345 6769)")]


# --- 1.3 · una fila mala no tumba el fichero ----------------------------------------


def test_un_fichero_con_tres_filas_ilegibles_se_importa_igual_y_las_nombra(
    session_factory, cuentas, http,
) -> None:
    datos = [
        *SABADELL_DATOS,
        # 18: el importe se fue a otra columna (las hay en el Belfius real)
        [_d(2026, 8, 31), "ABONO RARO", _d(2026, 8, 31), _d(2026, 8, 31), 100.0],
        # 19: fecha imposible
        ["31/13/2026", "OTRO", None, 5.0, 1.0],
        # 20: sin fecha ni importe, detrás de una fila que no se entendió
        [None, None, None, None, 7.0, None, "NADIE"],
    ]
    contenido = _sabadell(datos)
    r = _importar(session_factory, contenido)
    assert r["ok"] and r["imported"] == 6
    assert [x["fila"] for x in r["no_entendidas"]] == [18, 19, 20]
    assert r["no_entendidas"][0]["motivo"] == "importe no numérico: una fecha (2026-08-31)"
    assert r["no_entendidas"][1]["motivo"] == "fecha no reconocida: '31/13/2026'"
    assert r["no_entendidas"][2]["motivo"] == "sin fecha ni importe, y no sigue a un movimiento"
    assert r["no_entendidas"][2]["valores"] == ["7", "NADIE"]
    assert r["resumen"] == "6 movimientos importados, 3 filas no entendidas (18, 19, 20)"

    # La API lo devuelve igual (201, no un 422), para enseñarlo en la pantalla.
    with session_factory() as s:
        for m in s.scalars(select(BankMovement)).all():
            s.delete(m)
        s.commit()
    with patch("app.integrations.factusol.client.FactusolClient.from_settings",
               return_value=_fake([])):
        r = http.post(
            "/api/erp/bank/import",
            files={"file": ("Streamtec movimientos SABADELL.xlsx", contenido,
                            "application/octet-stream")},
            headers=auth_headers(http, "admin"),
        )
    assert r.status_code == 201, r.text
    cuerpo = r.json()
    assert cuerpo["account_name"] == "Sabadell Streamtec"
    assert cuerpo["resumen"].endswith("3 filas no entendidas (18, 19, 20)")
    assert cuerpo["anotadas"] == {"confirmadas": 2, "propuestas": 1, "sin_factura": 0,
                                  "nuevas_en_reimportados": 0}


# --- 1.4 · las columnas factura y pedido ---------------------------------------------


def test_una_fila_con_la_factura_escrita_casa_con_esa_factura(
    session_factory, cuentas,
) -> None:
    r = _importar(session_factory, _sabadell(SABADELL_DATOS))
    assert r["anotadas"] == {"confirmadas": 2, "propuestas": 1, "sin_factura": 0,
                             "nuevas_en_reimportados": 0}
    mov = _movimiento(session_factory, 11)
    assert (mov.status, mov.anotacion_factura, mov.anotacion_pedido) == (
        "reconciled", "5-260083", "BOP-099916")
    [rec] = _conciliaciones(session_factory, mov.id)
    assert (rec.serie, rec.codigo, rec.numero, float(rec.importe)) == (5, 260083, "5-260083", 145.2)
    assert (rec.status, rec.confidence, rec.origen) == ("confirmed", "alta", "extracto")
    assert rec.reason == "Venía anotado en el extracto («5-260083»)"
    assert rec.confirmed_by_user_id == "u-bart"

    # Con una nota al lado: propuesta para revisar, no conciliado.
    nota = _movimiento(session_factory, 14)
    assert nota.status == "pending"
    [prop] = _conciliaciones(session_factory, nota.id)
    assert (prop.numero, prop.status, prop.confidence) == ("5-260074", "proposed", "media")
    assert "Revísalo antes de confirmar" in prop.reason
    # Un pago con algo escrito: no se casa (la columna anota facturas de proveedor).
    assert _conciliaciones(session_factory, _movimiento(session_factory, 15).id) == []

    # El recálculo de propuestas no toca lo anotado; casa el resto.
    with session_factory() as s:
        stats = service.run_matching(s, service.pending_invoices(_fake([
            _fac(5, 260090, 3623.95, cliente="SOLITIUM SL", codcli=10),
            _fac(5, 260074, 420.74, cliente="Marta Molins", codcli=11),
        ]), "2026"))
    assert stats["anotadas"] == 1
    [prop] = _conciliaciones(session_factory, nota.id)
    assert prop.origen == "extracto" and prop.reason.startswith("Venía anotado")
    solitium = _movimiento(session_factory, 16)
    assert [c.numero for c in _conciliaciones(session_factory, solitium.id)] == ["5-260090"]
    # «Confirmar todas las de confianza alta» no se lleva la que tiene nota.
    with session_factory() as s:
        assert service.confirm_all_high(s, "u-bart") == 1     # solo SOLITIUM
    assert _movimiento(session_factory, 14).status == "pending"


def test_lo_que_se_lee_de_la_columna_factura() -> None:
    def leer(texto: object, serie: int | None = 2) -> tuple:
        a = leer_anotacion(texto, serie_cuenta=serie)
        return (a.facturas, a.limpia) if a else None

    assert leer("5-260083") == (((5, 260083),), True)
    assert leer("5-260088, 5-260089") == (((5, 260088), (5, 260089)), True)
    assert leer(526043.0) == (((2, 526043),), True)            # sin serie: la de la cuenta
    assert leer("5-250066, 250073") == (((5, 250066), (5, 250073)), True)
    assert leer("5-260067 (parcial, MQ Europe 18.813)") == (((5, 260067),), False)
    assert leer("525067 - parcial") == (((2, 525067),), False)
    assert leer("2-526082 (anticipo; cliente cita ref. 260720)") == (((2, 526082),), False)
    assert leer("Invoice No.: 25/F290808001") == ((), False)
    assert leer("falta facturar") == ((), False)
    assert leer("526043", serie=None) == ((), False)
    assert leer(None) is None


def test_varias_facturas_anotadas_reparten_el_importe(session_factory, cuentas) -> None:
    _importar(session_factory, _sabadell([
        [_d(2026, 8, 20), "ABONO TRANSFERENCIA DE PEPA", _d(2026, 8, 20), 100.0, 500.0, None,
         None, "5-260088, 5-260089"],
    ]))
    recs = _conciliaciones(session_factory, _movimiento(session_factory, 10).id)
    assert [(r.numero, float(r.importe), r.status) for r in recs] == [
        ("5-260088", 50.0, "confirmed"), ("5-260089", 50.0, "confirmed")]
    assert "repartido a partes iguales entre 2 facturas" in recs[0].reason


# --- duplicados --------------------------------------------------------------------


def test_importar_dos_veces_el_mismo_extracto_no_duplica(session_factory, cuentas) -> None:
    sabadell = _sabadell(SABADELL_DATOS)
    gemela = _belfius_fila(_d(2026, 10, 1), 25.0, "Mollie", "STORTING VAN Mollie")
    belfius = _belfius([gemela, list(gemela)])       # dos idénticas, sin saldo
    r1 = _importar(session_factory, sabadell)
    b1 = _importar(session_factory, belfius)
    assert (r1["imported"], b1["imported"]) == (6, 2)  # las gemelas son dos movimientos
    r2 = _importar(session_factory, sabadell)
    b2 = _importar(session_factory, belfius)
    assert (r2["imported"], r2["duplicates"]) == (0, 6)
    assert (b2["imported"], b2["duplicates"]) == (0, 2)
    with session_factory() as s:
        assert s.scalar(select(func.count(BankMovement.id))) == 8
        assert s.scalar(select(func.count(BankReconciliation.id))) == 3   # no se duplican


# --- 1.5 · el Excel conciliado ----------------------------------------------------


def test_el_excel_conciliado_sale_con_factura_y_pedido_rellenas(
    session_factory, cuentas, http,
) -> None:
    _importar(session_factory, _sabadell(SABADELL_DATOS))
    with session_factory() as s:
        # SOLITIUM se concilia en la pantalla; su pedido lo sabe BoHub.
        s.add(Order(external_source=OrderSource.WOOCOMMERCE, order_number="BOPRIN-100020",
                    factusol_invoice_number="260090", factusol_invoice_serie=5))
        s.commit()
        solitium = s.scalar(select(BankMovement).where(BankMovement.source_row == 16))
        service.reassign_movement(s, solitium.id, [
            {"serie": 5, "codigo": 260090, "importe": 3623.95, "cliente_nombre": "SOLITIUM SL"},
        ], user_id="u-bart")
    r = http.get(f"/api/erp/bank/export?account_id={cuentas['streamtec']}",
                 headers=auth_headers(http, "admin"))
    assert r.status_code == 200, r.text
    grid = [list(g) for g in load_workbook(io.BytesIO(r.content), read_only=True)
            .worksheets[0].iter_rows(values_only=True)]
    columnas = grid[5]
    assert columnas == ["FECHA OPER", "CONCEPTO", "FECHA VALOR", "IMPORTE", "SALDO",
                        "REFERENCIA 1", "REFERENCIA 2", "factura", "pedido"]  # sin las 247 vacías
    filas = {g[1]: g for g in grid[6:]}
    clara = filas["ABONO TRANSFERENCIA DE CLARA"]
    assert (clara[6], clara[7], clara[8]) == ("CLARA FIGUERAS M", "5-260083", "BOP-099916")
    soli = filas["ABONO TRANSFERENCIA DE SOLITIUM SL"]
    assert (soli[7], soli[8]) == ("5-260090", "BOPRIN-100020")       # el pedido, de BoHub
    marta = filas["ABONO TRANSFERENCIA DE MARTA"]
    assert marta[7].startswith("5-260074 (cobro")      # sin conciliar: lo que anotó Bart
    assert len(filas) == 6


# --- 1.6 · el CSV de cobros -------------------------------------------------------


def _factusol_con(*facturas: dict):
    return patch("app.integrations.factusol.client.FactusolClient.from_settings",
                 return_value=_fake(list(facturas)))


def test_el_csv_de_cobros_lleva_lo_que_cuadra_con_el_saldo_de_factusol(
    session_factory, cuentas, http,
) -> None:
    _importar(session_factory, _sabadell(SABADELL_DATOS))
    with session_factory() as s:
        # SOLITIUM, conciliado contra una factura de 5.000 €: es un cobro parcial.
        parcial = s.scalar(select(BankMovement).where(BankMovement.source_row == 16))
        s.add(BankReconciliation(movement_id=parcial.id, serie=5, codigo=260090,
                                 numero="5-260090", importe=3623.95, confidence="media",
                                 reason="cobro parcial de la factura", status="confirmed"))
        parcial.status = "reconciled"
        s.commit()
    with _factusol_con(
        _fac(5, 260083, 145.2, cliente="Clara Figueras", codcli=1),
        _fac(5, 260090, 5000.0, cliente="SOLITIUM SL", codcli=10),
        # 5-260082 (ROSA) ya no está pendiente: cobrada en FACTUSOL.
    ):
        r = http.get(f"/api/erp/bank/export-cobros?account_id={cuentas['streamtec']}",
                     headers=auth_headers(http, "admin"))
    assert r.status_code == 200, r.text
    assert r.headers["content-disposition"].endswith('cobros_conciliados.csv"')
    assert [r.headers[f"x-cobros-{k}"] for k in (
        "incluidos", "parciales", "no-pendientes", "no-cuadran", "sin-contrapartida",
    )] == ["1", "1", "1", "0", "0"]
    assert "X-Cobros-Parciales" in r.headers["access-control-expose-headers"]
    # UTF-8 con BOM, que Windows PowerShell 5.1 lee bien; observaciones en ASCII.
    assert r.content.startswith(b"\xef\xbb\xbf")
    filas = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert list(filas[0]) == ["serie", "codigo", "cuenta", "forma", "fecha", "observaciones"]
    assert [(f["serie"], f["codigo"], f["cuenta"], f["forma"], f["fecha"]) for f in filas] == [
        ("5", "260083", "8", "Transferencia", "18/08/2026"),
    ]
    obs = filas[0]["observaciones"]
    assert "anotado en el extracto" in obs and "145.20 EUR" in obs
    assert obs.isascii() and len(obs) <= 200
    # Un no autorizado no lo descarga.
    r = http.get("/api/erp/bank/export-cobros", headers=auth_headers(http, "user"))
    assert r.status_code == 403


def test_sin_factusol_no_hay_csv_de_cobros(session_factory, cuentas, http) -> None:
    _importar(session_factory, _sabadell(SABADELL_DATOS))
    with patch("app.integrations.factusol.client.FactusolClient.from_settings",
               side_effect=RuntimeError("DELSOL no responde")):
        r = http.get("/api/erp/bank/export-cobros", headers=auth_headers(http, "admin"))
    assert r.status_code == 503
    assert r.json()["detail"]["code"] == "factusol_unavailable"


def test_dos_propuestas_por_el_importe_entero_son_alternativas_y_no_se_confirman(
    session_factory, cuentas,
) -> None:
    _importar(session_factory, _sabadell(SABADELL_DATOS))
    with session_factory() as s:
        mov = s.scalar(select(BankMovement).where(BankMovement.source_row == 16))
        for codigo in (260090, 260091):
            s.add(BankReconciliation(movement_id=mov.id, serie=5, codigo=codigo,
                                     numero=f"5-{codigo}", importe=3623.95, confidence="alta",
                                     reason="mismo cliente, mismo importe", status="proposed"))
        s.commit()
        with pytest.raises(ValueError, match="alternativas"):
            service.confirm_movement(s, mov.id, "u-bart")
        assert service.confirm_all_high(s, "u-bart") == 0
    assert _movimiento(session_factory, 16).status == "pending"


# --- lo que salió de la revisión --------------------------------------------------


def test_una_fila_de_belfius_descolocada_no_se_importa_con_el_importe_equivocado(
    session_factory, cuentas,
) -> None:
    buena = _belfius_fila(_d(2026, 10, 19), 100.0, "A", "STORTING VAN A", numero=1.0)
    # Como en el Belfius real: una celda de más corre la fila una columna a la
    # derecha; el «importe» que queda en Bedrag no es el del movimiento.
    corrida = _belfius_fila(_d(2026, 10, 18), None, "B", "STORTING VAN B", numero=2.0)
    corrida[10], corrida[11], corrida[12] = 7.0, 200.0, "EUR"
    # Un tramo pegado partiendo por la coma decimal: «-90 | 75 | EUR». Los
    # céntimos no se adivinan («5» puede ser ,05 o ,50): se dice qué pasa.
    partida = _belfius_fila(_d(2026, 9, 22), None, "C", "STORTING VAN C", numero=3.0)
    partida[10], partida[11], partida[12] = -90.0, 75.0, "EUR"
    # Lo que no toca el importe sí entra: «##» (Excel sin ancho) en la fecha
    # valor y el código de país corrido a la columna del BIC.
    almohadillas = _belfius_fila(_d(2022, 1, 27), 12000.0, "D", "STORTING VAN D",
                                 numero=4.0, valuta="##")
    pais = _belfius_fila(_d(2023, 4, 11), 1695.0, "E", "STORTING VAN E", numero=5.0)
    pais[12] = "NL"
    r = _importar(session_factory, _belfius([buena, corrida, partida, almohadillas, pais]))
    assert r["imported"] == 3
    assert [(x["fila"], x["motivo"]) for x in r["no_entendidas"]] == [
        (4, "fila descolocada: divisa trae «200»"),
        (5, "importe partido en dos columnas («-90» y «75»): corrígelo en el fichero o "
            "vuelve a exportar ese tramo"),
    ]
    viejo = _movimiento(session_factory, 6)
    assert (float(viejo.importe), viejo.fecha_valor) == (12000.0, None)


def test_la_cuenta_de_la_cabecera_se_lee_celda_a_celda() -> None:
    # El IBAN en su celda y «factura» en la de al lado: unidos, el IBAN se
    # comería la palabra de al lado.
    contenido = _libro([
        ["Cuenta:", "ES91 2100 0418 4502 0005 1332", "factura", "pedido"],
        ["FECHA OPER", "CONCEPTO", "FECHA VALOR", "IMPORTE", "SALDO"],
        [_d(2026, 8, 1), "X", _d(2026, 8, 1), 1.0, 1.0],
    ])
    parsed = read_statement(contenido, "x.xlsx")
    assert parsed.header["iban"] == "ES9121000418450200051332"
    assert cuenta_casa(parsed.header["numero_cuenta"], STREAMTEC_IBAN)


def test_una_fila_con_concepto_pero_sin_fecha_ni_importe_no_se_pega_a_la_de_encima(
    session_factory, cuentas,
) -> None:
    r = _importar(session_factory, _sabadell([
        [_d(2026, 8, 24), "ABONO TRANSFERENCIA DE ROSA", _d(2026, 8, 24), 70.18, 6025.43],
        [None, "ABONO TRANSFERENCIA DE OTRA", None, None, 6100.0],
    ]))
    assert (r["imported"], r["continuaciones"]) == (1, 0)
    assert [(x["fila"], x["motivo"]) for x in r["no_entendidas"]] == [
        (11, "sin fecha de operación ni importe")]
    assert _movimiento(session_factory, 10).concepto == "ABONO TRANSFERENCIA DE ROSA"


def test_dos_movimientos_iguales_con_distinto_saldo_son_dos_y_reimportar_no_duplica(
    session_factory, cuentas,
) -> None:
    dos = _sabadell([
        [_d(2026, 8, 3), "COMISION MANTENIMIENTO", _d(2026, 8, 3), -5.0, 100.0],
        [_d(2026, 8, 3), "COMISION MANTENIMIENTO", _d(2026, 8, 3), -5.0, 95.0],
    ])
    assert _importar(session_factory, dos)["imported"] == 2
    r = _importar(session_factory, dos)
    assert (r["imported"], r["duplicates"]) == (0, 2)
    # Belfius sin saldo: el nº de transacción distingue dos iguales.
    a = _belfius_fila(_d(2026, 10, 1), 25.0, "Mollie", "STORTING VAN Mollie", numero=10.0)
    b = _belfius_fila(_d(2026, 10, 1), 25.0, "Mollie", "STORTING VAN Mollie", numero=11.0)
    assert _importar(session_factory, _belfius([a, b]))["imported"] == 2
    assert _importar(session_factory, _belfius([b]))["duplicates"] == 1


def test_lo_que_no_es_una_factura_en_la_columna_factura() -> None:
    def leer(texto: object, serie: int | None = 5) -> tuple:
        a = leer_anotacion(texto, serie_cuenta=serie)
        return (a.facturas, a.limpia) if a else None

    assert leer("5–260083") == (((5, 260083),), True)            # guion largo de Excel
    assert leer("proforma 260083") == ((), False)                # otro documento
    assert leer("PF 5-260083, 5-260084") == (((5, 260084),), False)
    assert leer("pedido BOP-100015") == ((), False)
    assert leer("2026-08-31") == ((), False)                     # una fecha
    assert leer("1.234,56") == ((), False)                       # un importe
    assert leer("5-260083/2") == ((), False)
    assert leer("5-260083 y 5-260084") == (((5, 260083), (5, 260084)), True)


def test_reimportar_el_extracto_ya_anotado_casa_lo_que_antes_no_traia(
    session_factory, cuentas,
) -> None:
    sin_anotar = [list(f[:7]) if f else f for f in SABADELL_DATOS]
    r1 = _importar(session_factory, _sabadell(sin_anotar))
    assert r1["anotadas"]["confirmadas"] == 0
    r2 = _importar(session_factory, _sabadell(SABADELL_DATOS))
    assert (r2["imported"], r2["duplicates"]) == (0, 6)
    assert r2["anotadas"] == {"confirmadas": 2, "propuestas": 1, "sin_factura": 0,
                              "nuevas_en_reimportados": 4}
    mov = _movimiento(session_factory, 11)
    assert (mov.status, mov.anotacion_factura) == ("reconciled", "5-260083")


def test_el_excel_conciliado_devuelve_las_columnas_sin_nombre(
    session_factory, cuentas, http,
) -> None:
    contenido = _belfius([
        _belfius_fila(_d(2026, 10, 19), 216.59, "Muster", "STORTING VAN Muster",
                      numero=187.0, factura="2-526126"),
    ])
    _importar(session_factory, contenido)
    r = http.get(f"/api/erp/bank/export?account_id={cuentas['mq']}",
                 headers=auth_headers(http, "admin"))
    assert r.status_code == 200, r.text
    grid = [list(g) for g in load_workbook(io.BytesIO(r.content), read_only=True)
            .worksheets[0].iter_rows(values_only=True)]
    columnas, fila = grid[0], grid[1]
    assert columnas[:4] == ["Rekening", "Boekingsdatum", None, "Transactienummer"]
    assert fila[2] == 29                       # la columna sin rótulo vuelve con lo que traía
    assert (columnas[18], fila[18]) == ("factura", "2-526126")
