"""ERP · Cuadre — API (`/api/erp/cuadre`) y Configuración ERP → «Cuadre».

Capacidad `erp.cuadre` (admin y ERP Pedidos). Revisar exige motivo; el Excel
lleva todos los abiertos; «Comprobar ahora» corre BoHub y encola FACTUSOL.
"""
from __future__ import annotations

import io
from collections.abc import Generator
from datetime import UTC, datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

import app.main  # noqa: F401
from app.db.base import Base
from app.db.session import get_session
from app.erp.cuadre import job
from app.erp.models import CuadreFinding, Order, PreparationStatus
from app.main import app
from tests._test_helpers import auth_headers, seed_test_users


@pytest.fixture()
def factory() -> Generator[sessionmaker, None, None]:
    engine = create_engine("sqlite+pysqlite:///:memory:",
                           connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    f = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    with f() as seed:
        seed_test_users(seed)
        seed.commit()
    yield f
    Base.metadata.drop_all(engine)


@pytest.fixture()
def http(factory, monkeypatch) -> Generator[TestClient, None, None]:
    def override() -> Generator[Session, None, None]:
        with factory() as s:
            yield s
    app.dependency_overrides[get_session] = override
    # Sin Redis en los tests: la pasada de FACTUSOL se «encola» en una lista.
    encolados: list = []
    monkeypatch.setattr(job, "_encolar", lambda func, *args: encolados.append(args))
    with TestClient(app) as c:
        c.encolados = encolados  # type: ignore[attr-defined]
        yield c
    app.dependency_overrides.clear()


def _pedido_sin_aprobar(factory, numero: str = "BOP-1", dias: int = 30) -> str:
    with factory() as s:
        hace = datetime.now(UTC) - timedelta(days=dias)
        o = Order(order_number=numero, preparation_status=PreparationStatus.PENDING_REVIEW,
                  placed_at=hace, created_at=hace)
        s.add(o)
        s.commit()
        return o.id


def test_sin_la_capacidad_no_se_entra(http):
    for role in ("comercial", "sat"):
        r = http.get("/api/erp/cuadre/resumen", headers=auth_headers(http, role))
        assert r.status_code == 403, role
    r = http.get("/api/erp/cuadre/resumen", headers=auth_headers(http, "pedidos"))
    assert r.status_code == 200


def test_comprobar_ahora_contadores_y_hallazgos(http, factory):
    oid = _pedido_sin_aprobar(factory)
    h = auth_headers(http, "admin")
    r = http.post("/api/erp/cuadre/comprobar", headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["lanzadas"]["mysql"]["estado"] == "ok"
    assert body["lanzadas"]["factusol"]["estado"] == "en_cola"     # a worker-sync
    assert body["lanzadas"]["woocommerce"]["estado"] == "en_cola"  # también (HTTP a tiendas)
    assert len(http.encolados) == 2
    res = body["resumen"]
    assert res["contadores"]["baja"] == 1 and res["contadores"]["total"] == 1
    assert sorted(p["fuente"] for p in res["en_curso"]) == ["factusol", "woocommerce"]
    assert res["ultima_pasada"]["fuente"] == "mysql"
    tarjeta = next(c for c in res["checks"] if c["id"] == "pedido_sin_aprobar")
    assert (tarjeta["abiertos"], tarjeta["nuevos"], tarjeta["severidad"]) == (1, 1, "baja")
    items = http.get("/api/erp/cuadre/hallazgos?check_id=pedido_sin_aprobar",
                     headers=h).json()["items"]
    assert [i["entidad_id"] for i in items] == [oid]
    assert items[0]["enlace"] == f"/erp/orders/{oid}"
    assert items[0]["arreglo_enlace"] == "/erp/orders?queue=por_revisar"
    assert http.get("/api/erp/cuadre/hallazgos?severidad=alta", headers=h).json()["items"] == []


def test_revisar_con_motivo_y_volver_a_incluir(http, factory):
    _pedido_sin_aprobar(factory)
    h = auth_headers(http, "pedidos")
    http.post("/api/erp/cuadre/comprobar", headers=h)
    fid = http.get("/api/erp/cuadre/hallazgos", headers=h).json()["items"][0]["id"]
    r = http.post(f"/api/erp/cuadre/hallazgos/{fid}/revisar", json={"motivo": " "}, headers=h)
    assert r.status_code == 400
    r = http.post(f"/api/erp/cuadre/hallazgos/{fid}/revisar",
                  json={"motivo": "Lo aprueba Bart el lunes"}, headers=h)
    assert r.status_code == 200 and r.json()["estado"] == "revisado"
    assert r.json()["motivo"] == "Lo aprueba Bart el lunes"
    assert http.get("/api/erp/cuadre/hallazgos", headers=h).json()["items"] == []
    con = http.get("/api/erp/cuadre/hallazgos?incluir_revisados=true", headers=h).json()
    assert [i["estado"] for i in con["items"]] == ["revisado"]
    assert http.get("/api/erp/cuadre/resumen", headers=h).json()["contadores"]["total"] == 0
    # Otra pasada con los mismos valores: sigue revisado.
    http.post("/api/erp/cuadre/comprobar", headers=h)
    assert http.get("/api/erp/cuadre/hallazgos", headers=h).json()["items"] == []
    r = http.post(f"/api/erp/cuadre/hallazgos/{fid}/reincluir", headers=h)
    assert r.status_code == 200 and r.json()["estado"] == "abierto"
    assert http.post(f"/api/erp/cuadre/hallazgos/{fid}/reincluir", headers=h).status_code == 400
    assert http.post("/api/erp/cuadre/hallazgos/nope/reincluir", headers=h).status_code == 404


def test_descargar_excel_con_los_abiertos(http, factory):
    _pedido_sin_aprobar(factory, "BOP-1")
    _pedido_sin_aprobar(factory, "BOP-2")
    h = auth_headers(http, "admin")
    http.post("/api/erp/cuadre/comprobar", headers=h)
    r = http.get("/api/erp/cuadre/export", headers=h)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "cuadre_descuadres_" in r.headers["content-disposition"]
    filas = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert filas[0][0] == "Severidad"
    assert sorted(f[3] for f in filas[1:]) == ["BOP-1", "BOP-2"]


def test_comprobaciones_y_el_cuadre_no_toca_los_pedidos(http, factory):
    oid = _pedido_sin_aprobar(factory)
    h = auth_headers(http, "admin")
    cat = http.get("/api/erp/cuadre/comprobaciones", headers=h).json()["items"]
    # 19 con «Lead sin contactar» (respuesta a leads · Fase 1).
    assert len(cat) == 19 and cat[0]["id"] == "factura_lineas_ajenas"
    http.post("/api/erp/cuadre/comprobar", headers=h)
    with factory() as s:
        o = s.get(Order, oid)
        assert o.preparation_status == PreparationStatus.PENDING_REVIEW   # solo lectura
        assert s.query(CuadreFinding).count() == 1


# --- Configuración ERP → «Cuadre» --------------------------------------------------


def test_configuracion_del_cuadre_en_ajustes(http):
    h = auth_headers(http, "admin")
    cfg = http.get("/api/erp/settings", headers=h).json()
    assert cfg["cuadre"]["nocturno_activo"] is False and cfg["cuadre"]["hora"] == "03:00"
    assert cfg["cuadre"]["checks"]["pedido_sin_aprobar"] == {"activo": True, "dias": 7}
    assert len(cfg["cuadre_catalogo"]) == 19
    r = http.patch("/api/erp/settings", headers=h, json={"cuadre": {
        "nocturno_activo": True, "hora": "02:15",
        "checks": {"pedido_sin_aprobar": {"activo": False, "dias": 10}},
    }})
    assert r.status_code == 200, r.text
    nuevo = r.json()["cuadre"]
    assert nuevo["nocturno_activo"] is True and nuevo["hora"] == "02:15"
    assert nuevo["checks"]["pedido_sin_aprobar"] == {"activo": False, "dias": 10}
    assert nuevo["checks"]["factura_sin_cobro"] == {"activo": True, "dias": 30}
    r = http.patch("/api/erp/settings", headers=h, json={"cuadre": {"hora": "3am"}})
    assert r.status_code == 400 and "HH:MM" in r.text


def test_desactivar_en_ajustes_saca_la_comprobacion_del_panel(http, factory):
    _pedido_sin_aprobar(factory)
    h = auth_headers(http, "admin")
    http.post("/api/erp/cuadre/comprobar", headers=h)
    http.patch("/api/erp/settings", headers=h, json={"cuadre": {
        "checks": {"pedido_sin_aprobar": {"activo": False}}}})
    res = http.get("/api/erp/cuadre/resumen", headers=h).json()
    assert "pedido_sin_aprobar" not in [c["id"] for c in res["checks"]]
    assert res["contadores"]["total"] == 0


# --- colas: reencolar y vaciar el registro de fallidos -----------------------


def _parche_colas(monkeypatch) -> dict:
    """Sin Redis: se sustituyen las dos operaciones y se mira qué se pidió."""
    from app.erp.api import cuadre as api_cuadre

    visto: dict = {}

    def _reencolar(cola, *, funcion=None, limite=0, probar=True):
        visto["reencolar"] = {"cola": cola, "funcion": funcion,
                              "limite": limite, "probar": probar}
        return {"cola": cola, "funcion": funcion, "probar": probar,
                "en_registro": 7, "elegidos": 2, "sin_datos": 0,
                "reencolados": 0 if probar else 2, "fallos_al_reencolar": 0,
                "tope": limite, "revisados": 7, "ventana_agotada": False}

    def _vaciar(cola, *, probar=True):
        visto["vaciar"] = {"cola": cola, "probar": probar}
        return {"cola": cola, "probar": probar, "en_registro": 7,
                "descartados": 0 if probar else 6, "sin_datos": 1,
                "fallos": 0, "quedan": 0, "tope": 20_000}

    monkeypatch.setattr(api_cuadre.registro_fallidos, "reencolar", _reencolar)
    monkeypatch.setattr(api_cuadre.registro_fallidos, "vaciar", _vaciar)
    monkeypatch.setattr(api_cuadre, "_cola_valida", lambda nombre: nombre)
    return visto


def test_vaciar_una_cola_no_lo_puede_hacer_quien_solo_ve_el_cuadre(
    http, factory, monkeypatch
):
    """El panel es de solo lectura y su capacidad la tiene también el rol de
    pedidos. Vaciar el registro es irreversible: solo admin."""
    _parche_colas(monkeypatch)
    r = http.post("/api/erp/cuadre/colas/vaciar",
                  json={"cola": "brevo:push_contact", "probar": False},
                  headers=auth_headers(http, "pedidos"))
    assert r.status_code == 403, r.text


def test_vaciar_una_cola_deja_fila_de_auditoria_con_cuantos(
    http, factory, monkeypatch
):
    from app.models.crm import AuditLog

    visto = _parche_colas(monkeypatch)
    previa = http.post("/api/erp/cuadre/colas/vaciar",
                       json={"cola": "brevo:push_contact"},
                       headers=auth_headers(http, "admin"))
    assert previa.status_code == 200, previa.text
    assert previa.json()["descartados"] == 0      # la vista previa no toca
    assert visto["vaciar"]["probar"] is True
    with factory() as s:
        assert s.query(AuditLog).filter(
            AuditLog.action == "queue.failed_cleared").count() == 0

    hecho = http.post("/api/erp/cuadre/colas/vaciar",
                      json={"cola": "brevo:push_contact", "probar": False},
                      headers=auth_headers(http, "admin"))
    assert hecho.status_code == 200, hecho.text
    assert hecho.json()["descartados"] == 6
    with factory() as s:
        fila = s.query(AuditLog).filter(
            AuditLog.action == "queue.failed_cleared").one()
        assert fila.target_id == "brevo:push_contact"
        assert "6" in (fila.metadata_json or "")


def test_vaciar_no_acepta_un_filtro_que_luego_se_ignora(http, monkeypatch):
    """`funcion`/`limite` en vaciar hacían parecer filtrada una petición que
    se lleva el registro entero."""
    _parche_colas(monkeypatch)
    r = http.post(
        "/api/erp/cuadre/colas/vaciar",
        json={"cola": "brevo:push_contact", "probar": False,
              "funcion": "push_contact_job"},
        headers=auth_headers(http, "admin"),
    )
    assert r.status_code == 422, r.text
