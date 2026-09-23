"""Rediseño 2026 del seguimiento: hoja simplificada, ordenada por Situación.

Comprueba que la Situación de cada fila se deriva de la cola de la línea de
vida (workflow), que la hoja se ordena por Situación (prioridad), que un pedido
«No requiere envío» muestra Preparación/Envío = «No aplica», que «Factura
enviada» sale del evento `erp.invoice_emailed`, que «Cobro» refleja el estado
contable de FACTUSOL, y que la pestaña Incidencias es el subconjunto EXACTO de
Situación=Incidencia.
"""
from __future__ import annotations

import io
from collections.abc import Generator
from datetime import UTC, datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

import app.main  # noqa: F401
from app.db.base import Base
from app.db.session import get_session
from app.erp.models import (
    ErpException,
    ExceptionStatus,
    ExceptionType,
    InvoiceStatus,
    Order,
    OrderSource,
    PaymentStatus,
    PreparationStatus,
    TransportStatus,
)
from app.main import app
from app.models.crm import AuditLog, Company
from tests._test_helpers import auth_headers, seed_test_users


@pytest.fixture()
def session_factory() -> Generator[sessionmaker, None, None]:
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False}, poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    with factory() as seed:
        seed_test_users(seed)
        seed.commit()
    yield factory
    Base.metadata.drop_all(engine)


@pytest.fixture()
def http(session_factory) -> Generator[TestClient, None, None]:
    def override() -> Generator[Session, None, None]:
        with session_factory() as s:
            yield s
    app.dependency_overrides[get_session] = override
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _dt(iso: str) -> datetime:
    return datetime.fromisoformat(iso).replace(tzinfo=UTC)


def _company(s: Session) -> str:
    # Vinculada a FACTUSOL → no genera la incidencia «empresa sin vincular».
    c = Company(name="Cliente SL", factusol_company_id="500")
    s.add(c)
    s.flush()
    return c.id


def _order(s: Session, number: str, company_id: str, **kw) -> Order:
    base = {
        "order_number": number,
        "external_source": OrderSource.MANUAL,
        "company_id": company_id,
        "total_amount": 100.0,
        "currency": "EUR",
        "transport_status": TransportStatus.NOT_SHIPPED,
        "invoice_status": InvoiceStatus.NOT_INVOICED,
        "preparation_status": PreparationStatus.IN_QUEUE,
        "approved_at": _dt("2026-09-01"),
        "placed_at": _dt("2026-09-01"),
    }
    base.update(kw)
    o = Order(**base)
    s.add(o)
    s.flush()
    return o


def _seed(s: Session) -> None:
    cid = _company(s)
    # A — facturado + cobrado, requiere envío → «Por enviar».
    _order(s, "ART-900001", cid, factusol_invoice_number="5-260001",
           factusol_cobro_status="cobrada")
    # B — facturado, sin cobrar → «Por cobrar». Factura enviada por email.
    b = _order(s, "ART-900002", cid, factusol_invoice_number="5-260002")
    s.add(AuditLog(
        action="erp.invoice_emailed", target_type="order", target_id=b.id,
        created_at=_dt("2026-09-10"),
    ))
    # C — aprobado, sin factura → «Por facturar».
    _order(s, "ART-900003", cid)
    # D — pedido WEB sin aprobar → «Por revisar». Es web a propósito: a un
    # pedido de la tienda no se le exige aprobación para estar en Seguimiento
    # (está vivo desde que entra); a un manual sin aprobar, sí.
    _order(s, "ART-900004", cid, approved_at=None,
           external_source=OrderSource.WOOCOMMERCE, external_id="900004",
           woo_status="processing",
           preparation_status=PreparationStatus.PENDING_REVIEW)
    # E — excepción abierta → «Incidencia» (aunque tenga factura).
    e = _order(s, "ART-900005", cid, factusol_invoice_number="5-260005")
    s.add(ErpException(
        order_id=e.id, type=ExceptionType.SAT_ISSUE, status=ExceptionStatus.OPEN,
        metadata_json='{"description": "Falta tornillería"}',
    ))
    # F — «No requiere envío», facturado + cobrado → «Listo».
    _order(s, "ART-900006", cid, factusol_invoice_number="5-260006",
           factusol_cobro_status="cobrada", shipping_not_required=True)
    s.commit()


def test_situacion_orden_y_campos(session_factory, http) -> None:
    with session_factory() as s:
        _seed(s)
    r = http.get("/api/erp/seguimiento", headers=auth_headers(http, "pedidos"))
    assert r.status_code == 200, r.text
    items = {i["order_number"]: i for i in r.json()["items"]}

    # Situación derivada de la cola de la línea de vida.
    assert items["ART-900001"]["situacion"] == "por_enviar"
    assert items["ART-900002"]["situacion"] == "por_cobrar"
    assert items["ART-900003"]["situacion"] == "por_facturar"
    assert items["ART-900004"]["situacion"] == "por_revisar"
    assert items["ART-900005"]["situacion"] == "incidencias"
    assert items["ART-900006"]["situacion"] == "listo"

    # Orden por Situación (prioridad): Incidencia → … → Listo.
    order = [i["order_number"] for i in r.json()["items"]]
    assert order == [
        "ART-900005", "ART-900004", "ART-900003",
        "ART-900002", "ART-900001", "ART-900006",
    ]

    # «No requiere envío» → Preparación/Envío = «No aplica».
    f = items["ART-900006"]
    assert f["preparacion"] == "No aplica"
    assert f["envio"] == "No aplica"

    # Cobro refleja el estado contable de FACTUSOL.
    assert items["ART-900001"]["cobro"] == "cobrado"
    assert items["ART-900002"]["cobro"] == "pendiente"
    assert items["ART-900003"]["cobro"] == "na"          # sin factura

    # Factura enviada = fecha del evento erp.invoice_emailed.
    assert items["ART-900002"]["factura_enviada"] == "2026-09-10"
    assert items["ART-900001"]["factura_enviada"] is None

    # Empresa (serie) e Importe.
    assert items["ART-900001"]["empresa_serie"].startswith("5 · ")
    assert items["ART-900001"]["importe"] == 100.0


def test_incidencias_sheet_es_subconjunto_exacto(session_factory, http) -> None:
    with session_factory() as s:
        _seed(s)
    r = http.get("/api/erp/seguimiento/export", headers=auth_headers(http, "pedidos"))
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    assert wb.sheetnames == ["Pedidos", "Incidencias"]

    inc = [list(row) for row in wb["Incidencias"].iter_rows(values_only=True)]
    # Solo el pedido con Situación=Incidencia (E), con su detalle.
    assert [r2[0] for r2 in inc[1:]] == ["ART-900005"]
    fila = inc[1]
    assert fila[2] == "Incidencia en SAT"                # Tipo
    assert fila[3] == "Falta tornillería"                # Motivo
    assert fila[6] == "Abierta"                          # Estado


def test_orden_por_situacion_funciona_como_defecto(session_factory, http) -> None:
    """La ordenación por Situación es la de por defecto (sin `sort`)."""
    with session_factory() as s:
        _seed(s)
    sin_sort = http.get("/api/erp/seguimiento", headers=auth_headers(http, "pedidos"))
    con_sort = http.get(
        "/api/erp/seguimiento?sort=situacion", headers=auth_headers(http, "pedidos")
    )
    assert (
        [i["order_number"] for i in sin_sort.json()["items"]]
        == [i["order_number"] for i in con_sort.json()["items"]]
    )


# --- solo pedidos VIVOS -------------------------------------------------------


def test_un_manual_sin_aprobar_no_sale_en_seguimiento(session_factory, http) -> None:
    """Seguimiento es la lista de pedidos vivos: un manual que nadie ha
    aprobado todavía no ha entrado al flujo. Sigue a un clic en «Ver ocultos
    por estado», y en cuanto se aprueba aparece."""
    with session_factory() as s:
        cid = _company(s)
        _order(s, "MANUAL-1", cid, approved_at=None,
               preparation_status=PreparationStatus.PENDING_REVIEW)
        s.commit()
    h = auth_headers(http, "pedidos")
    listado = http.get("/api/erp/seguimiento", headers=h).json()
    assert "MANUAL-1" not in {i["order_number"] for i in listado["items"]}

    ocultos = http.get("/api/erp/seguimiento?ver_ocultos_estado=true",
                       headers=h).json()["items"]
    fila = next(i for i in ocultos if i["order_number"] == "MANUAL-1")
    assert fila["estado_woo_motivo"] == "sin_aprobar"


def test_al_aprobarlo_vuelve_a_seguimiento(session_factory, http) -> None:
    with session_factory() as s:
        cid = _company(s)
        oid = _order(s, "MANUAL-2", cid, approved_at=None,
                     preparation_status=PreparationStatus.PENDING_REVIEW).id
        s.commit()
    h = auth_headers(http, "pedidos")
    assert "MANUAL-2" not in {
        i["order_number"] for i in http.get("/api/erp/seguimiento", headers=h).json()["items"]
    }
    with session_factory() as s:
        o = s.get(Order, oid)
        o.approved_at = _dt("2026-09-05")
        o.preparation_status = PreparationStatus.IN_QUEUE
        s.commit()
    assert "MANUAL-2" in {
        i["order_number"] for i in http.get("/api/erp/seguimiento", headers=h).json()["items"]
    }


def test_un_pedido_anulado_no_sale_en_seguimiento(session_factory, http) -> None:
    """El caso reportado: se anula y seguía saliendo como «Pendiente»."""
    with session_factory() as s:
        cid = _company(s)
        o = _order(s, "MANUAL-3", cid)
        o.cancelled_at = _dt("2026-09-06")
        s.commit()
    h = auth_headers(http, "pedidos")
    listado = http.get("/api/erp/seguimiento", headers=h).json()
    assert "MANUAL-3" not in {i["order_number"] for i in listado["items"]}
    ocultos = http.get("/api/erp/seguimiento?ver_ocultos_estado=true",
                       headers=h).json()["items"]
    assert next(i for i in ocultos
                if i["order_number"] == "MANUAL-3")["estado_woo_motivo"] == "anulado"


def test_aprobado_sin_pagar_sigue_en_seguimiento(session_factory, http) -> None:
    """No se excluye por falta de cobro."""
    with session_factory() as s:
        cid = _company(s)
        _order(s, "MANUAL-4", cid, payment_status=PaymentStatus.PENDING)
        s.commit()
    listado = http.get("/api/erp/seguimiento",
                       headers=auth_headers(http, "pedidos")).json()
    assert "MANUAL-4" in {i["order_number"] for i in listado["items"]}


def _web(s: Session, number: str, cid: str, woo_status: str, **kw):
    return _order(
        s, number, cid, external_source=OrderSource.WOOCOMMERCE,
        external_id=number.split("-")[-1], woo_status=woo_status, **kw,
    )


def test_pantalla_excel_y_drive_dicen_lo_mismo(session_factory, http) -> None:
    """Coherencia de las tres salidas: la pantalla, «Descargar Excel» y las
    filas que se vuelcan a la pestaña «Seguimiento (app)» de Drive salen de la
    misma consulta, así que la puerta web vale igual en las tres."""
    from app.erp import seguimiento as core  # noqa: PLC0415
    from app.erp.api.seguimiento import build_drive_sync_rows  # noqa: PLC0415

    with session_factory() as s:
        cid = _company(s)
        _web(s, "BOPRIN-99931", cid, "on-hold")        # fuera
        _web(s, "BOPRIN-99880", cid, "pending")        # fuera
        _web(s, "BOPRIN-99001", cid, "processing")     # dentro
        _web(s, "ARTISJ-9557", cid, "refunded",        # dentro, «Reembolsado»
             cancelled_at=_dt("2026-09-12"))
        s.commit()

    h = auth_headers(http, "pedidos")
    pantalla = {i["order_number"]
                for i in http.get("/api/erp/seguimiento", headers=h).json()["items"]}

    r = http.get("/api/erp/seguimiento/export", headers=h)
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    filas = [list(row) for row in wb["Pedidos"].iter_rows(values_only=True)]
    num = filas[0].index("Nº pedido")
    situacion_col = filas[0].index("Situación")
    excel = {f[num] for f in filas[1:]}

    with session_factory() as s:
        # Las mismas filas que `push_managed_tabs` vuelca a «Seguimiento (app)».
        vivas = core.sort_by_situacion(build_drive_sync_rows(s))
        drive_rows = [core.row_to_pedidos_values(r2) for r2 in vivas]
    drive = {f[num] for f in drive_rows}

    assert pantalla == excel == drive == {"BOPRIN-99001", "ARTISJ-9557"}
    # Y el reembolsado se dice por lo que es, también en el Excel.
    fila = next(f for f in filas[1:] if f[num] == "ARTISJ-9557")
    assert fila[situacion_col] == "Reembolsado"
