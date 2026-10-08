"""ERP · Cuadre — motor idempotente, configuración y job nocturno en worker-sync.

- Dos pasadas no duplican; lo que desaparece pasa a `resuelto`; un revisado no
  reaparece mientras su huella no cambie y vuelve como nuevo si cambia.
- Una comprobación que falla no toca sus descuadres; las desactivadas no corren.
- El job nocturno se arma en `cuadre:run` a la hora de Configuración (03:00 de
  Madrid por defecto), respeta el interruptor y los umbrales.
"""
from __future__ import annotations

import json
from collections.abc import Generator
from datetime import UTC, datetime, timedelta
from typing import Any
from unittest.mock import patch

import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

import app.main  # noqa: F401
from app.db.base import Base
from app.erp.cuadre import engine, job
from app.erp.cuadre.config import cuadre_config, validar_config
from app.erp.cuadre.registry import Comprobacion, Hallazgo, registro
from app.erp.models import (
    CuadreFinding,
    CuadreRun,
    ErpSettings,
    Order,
    OrderStatusHistory,
    StatusDomain,
    TransportStatus,
)
from app.erp.models.settings import ERP_SETTINGS_SINGLETON_ID

AHORA = datetime(2026, 10, 3, 10, 0, tzinfo=UTC)


@pytest.fixture()
def s() -> Generator[Session, None, None]:
    engine_ = create_engine("sqlite+pysqlite:///:memory:",
                            connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine_)
    with sessionmaker(bind=engine_)() as session:
        yield session
    Base.metadata.drop_all(engine_)


def _config(s: Session, **cuadre: Any) -> None:
    cfg = s.get(ErpSettings, ERP_SETTINGS_SINGLETON_ID)
    if cfg is None:
        cfg = ErpSettings(id=ERP_SETTINGS_SINGLETON_ID)
        s.add(cfg)
    cfg.factusol_series_json = json.dumps({"cuadre": cuadre})
    s.commit()


def _h(eid: str, valor: Any = 1) -> Hallazgo:
    return Hallazgo(entidad_tipo="pedido", entidad_id=eid, etiqueta=eid, detalle="x",
                    pista_de_arreglo="y", huella_datos={"v": valor})


COMP = Comprobacion(id="prueba", titulo="Prueba", descripcion="d", severidad="alta",
                    fuente="mysql", grupo="dinero", funcion=lambda ctx: [])


def _estados(s: Session) -> dict[str, str]:
    return {f.entidad_id: f.estado for f in s.scalars(select(CuadreFinding))}


# --- idempotencia ----------------------------------------------------------------------


def test_dos_pasadas_no_duplican(s):
    engine.aplicar(s, COMP, [_h("a"), _h("b")], AHORA)
    s.commit()
    r = engine.aplicar(s, COMP, [_h("a"), _h("b"), _h("a")], AHORA + timedelta(days=1))
    s.commit()
    assert r == {"hallazgos": 2, "nuevos": 0, "reabiertos": 0, "resueltos": 0}
    assert s.query(CuadreFinding).count() == 2
    a = s.scalars(select(CuadreFinding).where(CuadreFinding.entidad_id == "a")).one()
    assert a.primera_vez_at.replace(tzinfo=UTC) == AHORA
    assert a.ultima_vez_at.replace(tzinfo=UTC) == AHORA + timedelta(days=1)


def test_lo_que_desaparece_pasa_a_resuelto_y_si_vuelve_se_reabre(s):
    engine.aplicar(s, COMP, [_h("a"), _h("b")], AHORA)
    r = engine.aplicar(s, COMP, [_h("a")], AHORA + timedelta(days=1))
    s.commit()
    assert r["resueltos"] == 1
    assert _estados(s) == {"a": "abierto", "b": "resuelto"}       # nunca se borra
    b = s.scalars(select(CuadreFinding).where(CuadreFinding.entidad_id == "b")).one()
    assert b.resuelto_at is not None
    engine.aplicar(s, COMP, [_h("a"), _h("b")], AHORA + timedelta(days=2))
    s.commit()
    assert _estados(s) == {"a": "abierto", "b": "abierto"}
    assert b.resuelto_at is None and b.abierto_at.replace(tzinfo=UTC) == AHORA + timedelta(days=2)


def test_revisado_no_reaparece_hasta_que_cambia_la_huella(s):
    engine.aplicar(s, COMP, [_h("a", 100)], AHORA)
    s.commit()
    f = s.scalars(select(CuadreFinding)).one()
    engine.revisar(s, f.id, motivo="Es un abono pactado", user_id="u1")
    s.commit()
    engine.aplicar(s, COMP, [_h("a", 100)], AHORA + timedelta(days=1))   # misma huella
    s.commit()
    assert f.estado == "revisado" and f.motivo == "Es un abono pactado"
    r = engine.aplicar(s, COMP, [_h("a", 250)], AHORA + timedelta(days=2))  # cambian valores
    s.commit()
    assert r["reabiertos"] == 1
    assert f.estado == "abierto" and f.motivo is None and f.visto_por is None


def test_revisar_exige_motivo_y_reincluir_solo_revisados(s):
    engine.aplicar(s, COMP, [_h("a")], AHORA)
    s.commit()
    f = s.scalars(select(CuadreFinding)).one()
    with pytest.raises(engine.CuadreError):
        engine.revisar(s, f.id, motivo="  ", user_id="u1")
    with pytest.raises(engine.CuadreError):
        engine.reincluir(s, f.id)
    engine.revisar(s, f.id, motivo="ok, visto", user_id="u1")
    engine.reincluir(s, f.id)
    assert f.estado == "abierto" and f.motivo is None


# --- ejecutar ------------------------------------------------------------------------


@pytest.fixture()
def registro_de_prueba(monkeypatch):
    """Sustituye el registro por dos comprobaciones controladas por el test."""
    estado: dict[str, Any] = {"uno": [_h("a")], "dos": [_h("z")], "falla": False}

    def uno(ctx):
        return estado["uno"]

    def dos(ctx):
        if estado["falla"]:
            raise RuntimeError("boom")
        return estado["dos"]

    fake = {
        "uno": Comprobacion(id="uno", titulo="Uno", descripcion="", severidad="alta",
                            fuente="mysql", grupo="dinero", funcion=uno, orden=1),
        "dos": Comprobacion(id="dos", titulo="Dos", descripcion="", severidad="media",
                            fuente="mysql", grupo="envios", funcion=dos, orden=2),
    }
    registro()          # las comprobaciones reales se registran ANTES de sustituirlo
    monkeypatch.setattr("app.erp.cuadre.registry.REGISTRO", fake)
    return estado


def test_ejecutar_guarda_la_pasada_y_una_que_falla_no_resuelve_lo_suyo(s, registro_de_prueba):
    run = engine.ejecutar(s, fuente="mysql", origen="manual", ahora=AHORA)
    assert run.estado == "ok"
    assert _estados(s) == {"a": "abierto", "z": "abierto"}
    registro_de_prueba["falla"] = True
    registro_de_prueba["uno"] = []
    run = engine.ejecutar(s, fuente="mysql", origen="manual", ahora=AHORA + timedelta(hours=1))
    assert run.estado == "con_errores"
    assert "boom" in json.loads(run.resumen_json)["dos"]["error"]
    assert _estados(s) == {"a": "resuelto", "z": "abierto"}       # «dos» no se toca


def test_una_comprobacion_desactivada_ni_corre_ni_se_ve(s, registro_de_prueba):
    engine.ejecutar(s, fuente="mysql", origen="manual", ahora=AHORA)
    _config(s, checks={"dos": {"activo": False}})
    registro_de_prueba["dos"] = []
    engine.ejecutar(s, fuente="mysql", origen="manual", ahora=AHORA + timedelta(hours=1))
    assert _estados(s)["z"] == "abierto"                           # no ha corrido
    res = engine.resumen(s)
    assert [c["id"] for c in res["checks"]] == ["uno"]
    assert res["contadores"] == {"alta": 1, "media": 0, "baja": 0, "total": 1}
    assert [d["entidad_id"] for d in engine.listar(s)] == ["a"]


def test_resumen_y_solo_nuevos(s, registro_de_prueba):
    engine.ejecutar(s, fuente="mysql", origen="manual", ahora=AHORA)
    registro_de_prueba["uno"] = [_h("a"), _h("b")]
    engine.ejecutar(s, fuente="mysql", origen="nocturno", ahora=AHORA + timedelta(days=1))
    res = engine.resumen(s)
    assert res["contadores"]["total"] == 3
    assert res["ultima_pasada"]["origen"] == "nocturno"
    uno = next(c for c in res["checks"] if c["id"] == "uno")
    assert (uno["abiertos"], uno["nuevos"]) == (2, 1)
    assert [d["entidad_id"] for d in engine.listar(s, solo_nuevos=True)] == ["b"]
    f = s.scalars(select(CuadreFinding).where(CuadreFinding.entidad_id == "a")).one()
    engine.revisar(s, f.id, motivo="visto", user_id=None)
    s.commit()
    assert [d["entidad_id"] for d in engine.listar(s, check_id="uno")] == ["b"]
    con = engine.listar(s, check_id="uno", incluir_revisados=True)
    assert {d["entidad_id"]: d["estado"] for d in con} == {"a": "revisado", "b": "abierto"}


def test_exportar_excel_con_los_abiertos(s, registro_de_prueba):
    from io import BytesIO

    from openpyxl import load_workbook

    engine.ejecutar(s, fuente="mysql", origen="manual", ahora=AHORA)
    wb = load_workbook(BytesIO(engine.exportar_xlsx(s, base_url="https://bohub.test")))
    filas = list(wb.active.iter_rows(values_only=True))
    assert list(filas[0]) == engine.EXCEL_COLUMNAS
    assert {f[3] for f in filas[1:]} == {"a", "z"}
    assert filas[1][0] == "alta"                                    # ordenado por severidad


# --- umbrales de Configuración en una comprobación real ------------------------------


def test_los_umbrales_y_activaciones_de_configuracion_mandan(s):
    o = Order(order_number="BOP-1", transport_status=TransportStatus.IN_TRANSIT,
              created_at=datetime(2026, 1, 1, tzinfo=UTC),
              approved_at=datetime(2026, 1, 1, tzinfo=UTC))
    s.add(o)
    s.flush()
    s.add(OrderStatusHistory(order_id=o.id, domain=StatusDomain.TRANSPORT, to_status="in_transit",
                             changed_at=AHORA - timedelta(days=15), changed_by_user_id="u"))
    s.commit()
    engine.ejecutar(s, fuente="mysql", origen="manual", ahora=AHORA)
    assert s.scalars(select(CuadreFinding.check_id)).all() == ["envio_sin_entregar"]
    _config(s, checks={"envio_sin_entregar": {"activo": True, "dias": 20}})
    engine.ejecutar(s, fuente="mysql", origen="manual", ahora=AHORA)
    assert _estados(s) == {o.id: "resuelto"}


# --- configuración ---------------------------------------------------------------------


def test_config_por_defecto_y_validacion(s):
    cfg = cuadre_config(s)
    assert cfg["nocturno_activo"] is False and cfg["hora"] == "03:00"
    assert set(cfg["checks"]) == set(registro())
    assert len(registro()) == 18
    assert cfg["checks"]["factura_sin_cobro"] == {"activo": True, "dias": 30}
    assert cfg["checks"]["factura_lineas_ajenas"] == {"activo": True, "dias": None}
    # El embudo comercial no es un descuadre: apagada por defecto y a 90 días.
    assert cfg["checks"]["proforma_sin_convertir"] == {"activo": False, "dias": 90}
    with pytest.raises(ValueError, match="Hora"):
        validar_config({"hora": "25:00"})
    with pytest.raises(ValueError, match="desconocida"):
        validar_config({"checks": {"no_existe": {}}})
    with pytest.raises(ValueError, match="Días"):
        validar_config({"checks": {"factura_sin_cobro": {"dias": 0}}})
    ok = validar_config({"hora": "4:30", "nocturno_activo": True,
                         "checks": {"factura_sin_cobro": {"dias": 45, "activo": False}}})
    assert ok["hora"] == "04:30" and ok["checks"]["factura_sin_cobro"] == {
        "activo": False, "dias": 45}
    with pytest.raises(ValueError, match="sí o no"):
        validar_config({"nocturno_activo": "false"})
    with pytest.raises(ValueError, match="sí o no"):
        validar_config({"checks": {"factura_sin_cobro": {"activo": "no"}}})
    # Un cambio parcial se funde con lo guardado: lo que no viene se conserva.
    parcial = validar_config({"nocturno_activo": False}, actual=ok)
    assert parcial["hora"] == "04:30"
    assert parcial["checks"]["factura_sin_cobro"] == {"activo": False, "dias": 45}
    # Días vacíos = el de serie.
    vacio = validar_config({"checks": {"factura_sin_cobro": {"dias": None}}}, actual=ok)
    assert vacio["checks"]["factura_sin_cobro"]["dias"] == 30


# --- job nocturno ----------------------------------------------------------------------


def test_proxima_ejecucion_a_las_tres_de_madrid():
    # 3 de octubre, 12:00 UTC (14:00 en Madrid, CEST) → 4 de octubre 03:00 = 01:00 UTC.
    assert job.proxima_ejecucion("03:00", datetime(2026, 10, 3, 12, 0, tzinfo=UTC)) == \
        datetime(2026, 10, 4, 1, 0, tzinfo=UTC)
    # En invierno (CET): 03:00 de Madrid = 02:00 UTC.
    assert job.proxima_ejecucion("03:00", datetime(2026, 12, 1, 12, 0, tzinfo=UTC)) == \
        datetime(2026, 12, 2, 2, 0, tzinfo=UTC)
    # Un tic que llega un pelo antes no se re-programa para dentro de segundos.
    casi = datetime(2026, 10, 4, 0, 59, 50, tzinfo=UTC)
    assert job.proxima_ejecucion("03:00", casi) == datetime(2026, 10, 5, 1, 0, tzinfo=UTC)


def test_hora_de_madrid_sin_tzdata(monkeypatch):
    import zoneinfo

    def sin_tz(_name):
        raise zoneinfo.ZoneInfoNotFoundError("sin tzdata")

    monkeypatch.setattr(zoneinfo, "ZoneInfo", sin_tz)
    assert job.proxima_ejecucion("03:00", datetime(2026, 10, 3, 12, 0, tzinfo=UTC)) == \
        datetime(2026, 10, 4, 1, 0, tzinfo=UTC)
    assert job.proxima_ejecucion("03:00", datetime(2026, 12, 1, 12, 0, tzinfo=UTC)) == \
        datetime(2026, 12, 2, 2, 0, tzinfo=UTC)


def test_el_job_se_arma_en_cuadre_run_con_setnx(monkeypatch):
    captured: dict[str, Any] = {}

    class FakeRedis:
        def __init__(self) -> None:
            self.keys: dict[str, str] = {}

        def set(self, key, value, nx=False, ex=None):
            if nx and key in self.keys:
                return False
            self.keys[key] = value
            captured["key"], captured["ttl"] = key, ex
            return True

        def delete(self, key):
            self.keys.pop(key, None)

    class FakeQueue:
        def __init__(self, name, connection=None, default_timeout=None):
            captured["queue"], captured["timeout"] = name, default_timeout

        def enqueue_in(self, delay, fn):
            captured["delay"], captured["fn"] = delay, fn
            captured["n"] = captured.get("n", 0) + 1

    monkeypatch.setattr(job, "_hora_configurada", lambda: "03:00")
    conn = FakeRedis()
    ahora = datetime(2026, 10, 3, 12, 0, tzinfo=UTC)
    with patch("app.workers.queues.redis_connection", return_value=conn), \
            patch("rq.Queue", FakeQueue):
        job.schedule_nightly(ahora)
        job.schedule_nightly(ahora)                     # heartbeat vivo → no encola otro
    assert captured["queue"] == "cuadre:run"
    assert captured["delay"] == timedelta(hours=13)     # hasta las 01:00 UTC
    assert captured["fn"] is job._nightly_runner
    assert captured["key"] == "cuadre:nightly:heartbeat"
    assert captured["ttl"] == 13 * 3600 - 30
    assert captured["n"] == 1


def test_el_nocturno_respeta_el_interruptor_y_no_repite(s, monkeypatch):
    corridas: list[str] = []
    monkeypatch.setattr(job, "correr", lambda session, *, fuente, origen, **kw:
                        corridas.append(f"{fuente}:{origen}"))
    a_las_tres = datetime(2026, 10, 4, 1, 0, tzinfo=UTC)
    assert job.run_nightly(s, ahora=a_las_tres) is False            # apagado por defecto
    assert corridas == []
    _config(s, nocturno_activo=True)
    assert job.run_nightly(s, ahora=a_las_tres) is True
    assert corridas == ["mysql:nocturno", "factusol:nocturno", "woocommerce:nocturno"]
    # Si ya hubo pasada nocturna hace menos de 12 h (p. ej. al cambiar la hora,
    # el tic re-armado para la nueva), no repite.
    s.add(CuadreRun(fuente="mysql", origen="nocturno", estado="ok", resumen_json="{}",
                    started_at=a_las_tres))
    s.commit()
    assert job.run_nightly(s, ahora=a_las_tres + timedelta(hours=2)) is False
    assert job.run_nightly(s, ahora=a_las_tres + timedelta(hours=24)) is True


def test_comprobar_ahora_bohub_al_momento_y_factusol_a_la_cola(s, monkeypatch):
    encolados: list[Any] = []
    monkeypatch.setattr(job, "_encolar", lambda func, *args: encolados.append((func, args)))
    out = job.comprobar_ahora(s, user_id="u1")
    assert out["mysql"]["estado"] == "ok"
    assert out["factusol"]["estado"] == "en_cola"
    assert out["woocommerce"]["estado"] == "en_cola"      # pregunta a las tiendas: worker
    assert encolados == [
        (job._run_en_segundo_plano, (out["factusol"]["id"], "factusol")),
        (job._run_en_segundo_plano, (out["woocommerce"]["id"], "woocommerce")),
    ]
    assert s.get(CuadreRun, out["factusol"]["id"]).fuente == "factusol"
    assert s.get(CuadreRun, out["woocommerce"]["id"]).fuente == "woocommerce"
    # Mientras esas siguen en cola no se encolan otras.
    out2 = job.comprobar_ahora(s, user_id="u1")
    assert out2["factusol"]["id"] == out["factusol"]["id"] and len(encolados) == 2
    assert sorted(r["fuente"] for r in engine.resumen(s)["en_curso"]) == [
        "factusol", "woocommerce"]


def test_comprobar_ahora_sin_redis_deja_la_pasada_en_error(s, monkeypatch):
    def sin_redis(*_a):
        raise ConnectionError("redis caído")

    monkeypatch.setattr(job, "_encolar", sin_redis)
    out = job.comprobar_ahora(s, user_id=None)
    run = s.get(CuadreRun, out["factusol"]["id"])
    assert run.estado == "error" and "encolar" in run.error


def test_la_pasada_encolada_de_factusol_lee_con_el_cliente_y_termina(s, monkeypatch):
    from tests.test_erp_cuadre_checks import FakeFactusol, _fac, _vincular

    _vincular(s, (5, 1))

    client = FakeFactusol(
        F_FAC=[_fac(5, 1, net=100)],
        F_LFA=[{"TIPLFA": "5", "CODLFA": 99, "TOTLFA": 10}],          # de otra factura
        F_LCO=[{"TFALCO": "5", "CFALCO": 99, "IMPLCO": 10}],
        F_PRE=[{"TIPPRE": "1", "CODPRE": 1, "ESTPRE": 0, "FECPRE": "2026-07-01"}],
    )
    monkeypatch.setattr(engine, "_factusol_client", lambda session: (client, "2026"))
    run = engine.nueva_pasada(s, fuente="factusol", origen="manual")
    s.commit()
    run = job.correr(s, fuente="factusol", origen="manual", run_id=run.id)
    assert run.estado == "ok"
    # «Proforma aceptada sin convertir» viene apagada por defecto: no corre.
    assert set(json.loads(run.resumen_json)) == {
        "factura_lineas_ajenas", "cobro_descuadrado", "factura_sin_cobro",
        "factura_sin_vincular"}
    assert client.lecturas["F_PRE"] == 0
    assert client.lecturas["F_FAC"] == 1                              # una lectura por tabla
    vistos = {(f.check_id, f.entidad_id) for f in s.scalars(select(CuadreFinding))}
    assert vistos == {("factura_lineas_ajenas", "5-000001"),        # sin líneas
                      ("factura_sin_cobro", "5-000001")}            # y sin cobrar


def test_una_lectura_vacia_de_factusol_no_resuelve_ni_borra_el_revisado(s, monkeypatch):
    """F_LCO llega vacía (lectura rota): las comprobaciones que la usan fallan y
    sus descuadres —incluido un «revisado» con su motivo— se quedan como estaban."""
    from tests.test_erp_cuadre_checks import FakeFactusol, _fac, _vincular

    _vincular(s, (5, 74))

    tablas = {
        "F_FAC": [_fac(5, 74, net=100, estfac="2")],
        "F_LFA": [{"TIPLFA": "5", "CODLFA": 74, "TOTLFA": 100}],
        "F_LCO": [{"TFALCO": "5", "CFALCO": 1, "IMPLCO": 5}],
        "F_PRE": [{"TIPPRE": "1", "CODPRE": 1, "ESTPRE": 0, "FECPRE": "2026-07-01"}],
    }
    monkeypatch.setattr(engine, "_factusol_client",
                        lambda session: (FakeFactusol(**tablas), "2026"))
    engine.ejecutar(s, fuente="factusol", origen="manual", ahora=AHORA)
    f = s.scalars(select(CuadreFinding).where(
        CuadreFinding.check_id == "cobro_descuadrado")).one()        # cobrada sin líneas
    engine.revisar(s, f.id, motivo="Cobrada en caja, sin apunte", user_id="u1")
    s.commit()
    tablas["F_LCO"] = []
    run = engine.ejecutar(s, fuente="factusol", origen="manual", ahora=AHORA + timedelta(days=1))
    assert run.estado == "con_errores"
    assert "vacía" in json.loads(run.resumen_json)["cobro_descuadrado"]["error"]
    assert f.estado == "revisado" and f.motivo == "Cobrada en caja, sin apunte"


def test_al_cambiar_de_ejercicio_no_se_resuelve_lo_del_anterior(s, monkeypatch):
    """Las facturas de 2026 dejan de leerse al pasar a 2027: sus descuadres no
    se dan por resueltos (no se han mirado); los de 2027 van a su aire."""
    from tests.test_erp_cuadre_checks import FakeFactusol, _fac, _vincular

    _vincular(s, (5, 80), (5, 1))

    def pasada(ejercicio, facs, cuando):
        client = FakeFactusol(
            F_FAC=facs, F_LFA=[{"TIPLFA": "9", "CODLFA": 1, "TOTLFA": 1}],
            F_LCO=[{"TFALCO": "9", "CFALCO": 1, "IMPLCO": 1}],
            F_PRE=[{"TIPPRE": "1", "CODPRE": 1, "ESTPRE": 0, "FECPRE": "2026-07-01"}],
        )
        monkeypatch.setattr(engine, "_factusol_client", lambda session: (client, ejercicio))
        engine.ejecutar(s, fuente="factusol", origen="nocturno", ahora=cuando)

    pasada("2026", [_fac(5, 80, net=100, fecha="2026-08-01")], AHORA)
    pasada("2027", [_fac(5, 1, net=50, fecha="2027-01-02")], AHORA + timedelta(days=140))
    sin_cobro = {f.entidad_id: f.estado for f in s.scalars(select(CuadreFinding).where(
        CuadreFinding.check_id == "factura_sin_cobro"))}
    assert sin_cobro == {"5-000080": "abierto", "5-000001": "abierto"}


def test_la_pasada_carga_los_pedidos_una_vez(s):
    """Sin caducar los objetos entre comprobaciones: el nº de consultas no crece
    con los pedidos (antes, ~10 por pedido)."""
    from sqlalchemy import event

    for i in range(40):
        o = Order(order_number=f"BOP-{i}", transport_status=TransportStatus.IN_TRANSIT,
                  created_at=datetime(2026, 1, 1, tzinfo=UTC),
                  approved_at=datetime(2026, 1, 1, tzinfo=UTC))
        s.add(o)
        s.flush()
        s.add(OrderStatusHistory(order_id=o.id, domain=StatusDomain.TRANSPORT,
                                 to_status="in_transit", changed_by_user_id="u",
                                 changed_at=AHORA - timedelta(days=15)))
    s.commit()
    consultas: list[str] = []

    def contar(_conn, _cursor, sql, *_a, **_k):
        consultas.append(sql)

    bind = s.get_bind()
    event.listen(bind, "before_cursor_execute", contar)
    try:
        engine.ejecutar(s, fuente="mysql", origen="manual", ahora=AHORA)
    finally:
        event.remove(bind, "before_cursor_execute", contar)
    assert s.query(CuadreFinding).count() == 40
    assert len(consultas) < 150, len(consultas)
    assert s.expire_on_commit is True                                # se restaura


def test_factusol_sin_configurar_deja_error_sin_tocar_lo_guardado(s, monkeypatch):
    monkeypatch.setattr(engine, "_factusol_client", lambda session: (None, None))
    s.add(CuadreFinding(check_id="factura_sin_cobro", entidad_tipo="factura",
                        entidad_id="5-000009", huella="h", detalle_json="{}", severidad="alta",
                        estado="abierto", primera_vez_at=AHORA, ultima_vez_at=AHORA,
                        abierto_at=AHORA))
    s.commit()
    run = engine.ejecutar(s, fuente="factusol", origen="manual")
    assert run.estado == "error"
    assert _estados(s) == {"5-000009": "abierto"}


def test_una_config_guardada_antes_no_enciende_las_proformas(s):
    """La sección Cuadre se guarda entera: una config de antes del cambio de
    defecto traía las proformas activas y a 30 días. Ahí manda el defecto nuevo;
    lo demás que se eligió se conserva. Al volver a guardar, manda lo elegido."""
    _config(s, nocturno_activo=True, hora="02:00", checks={
        "proforma_sin_convertir": {"activo": True, "dias": 30},
        "factura_sin_cobro": {"activo": False, "dias": 45},
    })
    cfg = cuadre_config(s)
    assert cfg["version"] == 2
    assert cfg["checks"]["proforma_sin_convertir"] == {"activo": False, "dias": 90}
    assert cfg["checks"]["factura_sin_cobro"] == {"activo": False, "dias": 45}
    assert (cfg["nocturno_activo"], cfg["hora"]) == (True, "02:00")
    guardado = json.loads(s.get(ErpSettings, ERP_SETTINGS_SINGLETON_ID).factusol_series_json)
    elegido = validar_config(
        {"checks": {"proforma_sin_convertir": {"activo": True, "dias": 60}}},
        actual=guardado["cuadre"],
    )
    _config(s, **elegido)
    assert cuadre_config(s)["checks"]["proforma_sin_convertir"] == {"activo": True, "dias": 60}
