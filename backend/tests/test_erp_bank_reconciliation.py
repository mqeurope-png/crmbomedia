"""ERP-F4-A — conciliación bancaria: cuentas, importación del extracto,
casado con facturas pendientes, revisión (la persona decide) y export Excel.

Regla firme del PR: NADA se concilia solo y NO se escribe en FACTUSOL. Las
facturas vienen de F_FAC/F_LCO (solo lectura) a través del `FakeClient`.
"""
from __future__ import annotations

import io
from collections.abc import Generator
from datetime import date, datetime
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

import app.main  # noqa: F401
from app.db.base import Base
from app.db.session import get_session
from app.erp.bank import service
from app.erp.bank.matching import (
    DEFAULT_EXCLUSION_PATTERNS,
    extract_payer,
    is_excluded,
    names_match,
    normalize_name,
    propose,
)
from app.erp.bank.parsing import ParseError, parse_amount, parse_date, read_statement
from app.erp.models import BankLearnedRule, BankMovement, BankReconciliation
from app.main import app
from tests._test_helpers import auth_headers, seed_test_users
from tests.test_factusol_documents import FakeClient

# --- datos de referencia (los reales del extracto de Sabadell) ------------------

SABADELL_IBAN = "ES3300810202130001171918"
OPENBANK_IBAN = "ES6500730100510505555449"

HEADER_LINES = [
    "Cuenta: ES33 0081 0202 1300 0117 1918 · 337906622 · Boprint",
    "Divisa: EUR",
    "Titular: BOMEDIA S.L",
]
COLUMNS = [
    "FECHA OPER", "CONCEPTO", "FECHA VALOR", "IMPORTE", "SALDO",
    "REFERENCIA 1", "REFERENCIA 2", "FACTURA", "PRESUPUESTO", "PEDIDO",
]


def _row(fecha: str, concepto: str, importe: str, saldo: str,
         ref1: str = "", ref2: str = "") -> list:
    return [fecha, concepto, fecha, importe, saldo, ref1, ref2, None, None, None]


def _xlsx(rows: list[list], *, header: list[str] | None = None,
          columns: list[str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    for line in (HEADER_LINES if header is None else header):
        ws.append([line])
    ws.append([])
    ws.append(COLUMNS if columns is None else columns)
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fac(serie: int, codigo: int, total: float, *, cliente: str, codcli: int,
         fecha: str = "2026-08-01T00:00:00", estfac: str = "0") -> dict:
    return {
        "TIPFAC": str(serie), "CODFAC": codigo, "TOTFAC": total, "ESTFAC": estfac,
        "CLIFAC": codcli, "CNOFAC": cliente, "FECFAC": fecha, "REFFAC": f"BOP-{codigo}",
    }


# Facturas pendientes «tipo» (nombres/importes del caso real).
HIRSCH = _fac(1, 260649, 13900.45, cliente="Hirsch Armbänder GmbH", codcli=2101,
              fecha="2026-07-20T00:00:00")
SOLITIUM = _fac(1, 260720, 3623.95, cliente="SOLITIUM SL", codcli=10)
DM = _fac(1, 260700, 1314.47, cliente="DM Document Materiel SA", codcli=300)
MOVI_A = _fac(1, 260701, 100.00, cliente="MOVIATICOS SL", codcli=99)
MOVI_B = _fac(1, 260702, 250.50, cliente="MOVIATICOS SL", codcli=99,
              fecha="2026-08-05T00:00:00")
COBRADA = _fac(1, 260719, 192.39, cliente="Cliente cobrado", codcli=5, estfac="2")

INVOICES = [HIRSCH, SOLITIUM, DM, MOVI_A, MOVI_B, COBRADA]

F_BAN = [
    {"IBABAN": "ES33 0081 0202 1300 0117 1918", "BICBAN": "BSABESBB",
     "ENTBAN": "Banco Sabadell", "CUEBAN": "0117 1918", "DOMBAN": "Barcelona"},
    {"IBABAN": "ES65 0073 0100 5105 0555 5449", "BICBAN": "OPENESMM",
     "ENTBAN": "Open Bank", "CUEBAN": "0555 5449", "DOMBAN": "Madrid"},
]


def _fake(invoices: list[dict] | None = None) -> FakeClient:
    return FakeClient({
        "F_FAC": INVOICES if invoices is None else invoices,
        "F_LCO": [],
        "F_BAN": F_BAN,
        "F_FOP": [{"CODFOP": "002", "NOMFOP": "Transferencia"}],
    })


def _patched(fake: FakeClient):
    return patch(
        "app.integrations.factusol.client.FactusolClient.from_settings",
        return_value=fake,
    )


def _invoices(fake: FakeClient) -> list[dict]:
    return service.pending_invoices(fake, "2026")


# --- fixtures ----------------------------------------------------------------------


@pytest.fixture()
def session_factory() -> Generator[sessionmaker, None, None]:
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False}, poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    with factory() as seed:
        seed_test_users(seed)
    yield factory
    Base.metadata.drop_all(engine)


@pytest.fixture()
def http(session_factory) -> Generator[TestClient, None, None]:
    def override() -> Generator[Session, None, None]:
        with session_factory() as s:
            yield s
    app.dependency_overrides[get_session] = override
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


@pytest.fixture()
def account_id(session_factory) -> str:
    with session_factory() as s:
        acc = service.create_account(s, {
            "name": "Sabadell", "iban": "ES33 0081 0202 1300 0117 1918",
            "bank_name": "Banco Sabadell", "bic": "BSABESBB",
        })
        assert acc.iban == SABADELL_IBAN      # se normaliza sin espacios
        return acc.id


def _import(session_factory, content: bytes, filename: str = "extracto.xlsx",
            account_id: str | None = None) -> dict:
    with session_factory() as s:
        return service.import_statement(
            s, content=content, filename=filename, account_id=account_id,
        )


def _match(session_factory, fake: FakeClient) -> dict:
    with session_factory() as s:
        return service.run_matching(s, _invoices(fake))


def _movements(session_factory, **filters) -> dict:
    with session_factory() as s:
        return service.list_movements(s, **filters)


def _by_concepto(page: dict, needle: str) -> dict:
    return next(i for i in page["items"] if needle in i["concepto"])


# --- Parte A: cuentas --------------------------------------------------------------


def test_import_detects_account_by_iban(session_factory, http, account_id) -> None:
    # El IBAN de la cabecera («Cuenta: ES33 0081 …») identifica la cuenta.
    content = _xlsx([_row("3/9/2026", "ABONO TRANSFERENCIA DE SOLITIUM SL",
                          "3.623,95", "10.000,00")])
    summary = _import(session_factory, content)
    assert summary["ok"] is True
    assert summary["account_id"] == account_id
    assert summary["iban"] == SABADELL_IBAN
    assert summary["imported"] == 1

    # Extracto de una cuenta NO dada de alta → aviso, no se importa a ciegas.
    other = _xlsx(
        [_row("3/9/2026", "ABONO TRANSFERENCIA DE PEPE", "100,00", "500,00")],
        header=["Cuenta: ES65 0073 0100 5105 0555 5449 · 1234 · Open",
                "Divisa: EUR", "Titular: BOMEDIA S.L"],
    )
    res = _import(session_factory, other, "openbank.xlsx")
    assert res["ok"] is False
    assert res["code"] == "unknown_account"
    assert res["iban"] == OPENBANK_IBAN
    assert res["imported"] == 0
    with session_factory() as s:
        assert s.scalar(select(func.count(BankMovement.id))) == 1

    # Cuenta elegida a mano pero el IBAN del fichero es otro → se avisa.
    res = _import(session_factory, other, "openbank.xlsx", account_id=account_id)
    assert res["code"] == "iban_mismatch"

    # Por HTTP el aviso es un 409 con el resumen (sin importar nada).
    r = http.post(
        "/api/erp/bank/import",
        files={"file": ("openbank.xlsx", other, "application/octet-stream")},
        data={"run_match": "false"},
        headers=auth_headers(http, "pedidos"),
    )
    assert r.status_code == 409, r.text
    assert r.json()["detail"]["code"] == "unknown_account"

    # Las cuentas de F_BAN se SUGIEREN (sin las ya dadas de alta).
    with _patched(_fake()):
        r = http.get("/api/erp/bank/accounts/suggested", headers=auth_headers(http, "user"))
    assert r.status_code == 200, r.text
    assert [s["iban"] for s in r.json()["items"]] == [OPENBANK_IBAN]
    assert r.json()["items"][0]["bank_name"] == "Open Bank"

    # Dar de alta cuentas es de admin; el mapeo se guarda la 1ª importación.
    r = http.post("/api/erp/bank/accounts", json={"name": "x", "iban": OPENBANK_IBAN},
                  headers=auth_headers(http, "pedidos"))
    assert r.status_code == 403
    r = http.get("/api/erp/bank/accounts", headers=auth_headers(http, "user"))
    acc = r.json()["items"][0]
    assert acc["column_mapping"]["importe"] == "IMPORTE"
    assert acc["has_statement_header"] is True


# --- Parte B: importación -------------------------------------------------------------


def test_import_is_idempotent_on_overlapping_files(session_factory, account_id) -> None:
    rows = [
        _row(f"{d}/9/2026", f"ABONO TRANSFERENCIA DE CLIENTE {d}", f"{100 + d},00",
             f"{1000 + d},00")
        for d in range(1, 6)
    ]
    first = _import(session_factory, _xlsx(rows[:3]), "sept_1.xlsx")
    assert (first["imported"], first["duplicates"]) == (3, 0)
    # Rango solapado (filas 2..5): solo entran las 2 nuevas.
    second = _import(session_factory, _xlsx(rows[1:]), "sept_2.xlsx")
    assert (second["imported"], second["duplicates"]) == (2, 2)
    # Reimportar el primero: nada nuevo.
    again = _import(session_factory, _xlsx(rows[:3]), "sept_1.xlsx")
    assert (again["imported"], again["duplicates"]) == (0, 3)
    with session_factory() as s:
        assert s.scalar(select(func.count(BankMovement.id))) == 5
    # Misma fecha/importe/concepto pero SALDO distinto = dos movimientos reales
    # (dos cobros iguales el mismo día), no un duplicado.
    twins = [_row("9/9/2026", "ABONO TRANSFERENCIA DE GEMELO", "50,00", "2.000,00"),
             _row("9/9/2026", "ABONO TRANSFERENCIA DE GEMELO", "50,00", "2.050,00")]
    res = _import(session_factory, _xlsx(twins), "gemelos.xlsx")
    assert (res["imported"], res["duplicates"]) == (2, 0)


def test_import_parses_spanish_number_and_date_formats(session_factory, account_id) -> None:
    # Importes: español («1.314,47»), anglosajón («-1,500.00») y sin miles.
    assert parse_amount("1.314,47") == 1314.47
    assert parse_amount("-1,500.00") == -1500.0
    assert parse_amount("110,11") == 110.11
    assert parse_amount("6950.22") == 6950.22
    assert parse_amount("1.314") == 1314.0          # «.» con 3 dígitos = millar
    assert parse_amount("13.900,45") == 13900.45
    assert parse_amount(2278) == 2278.0
    assert parse_amount("2.278,00 €") == 2278.0
    with pytest.raises(ValueError):
        parse_amount("abc")
    # Fechas d/m/yyyy (con y sin cero), ISO y celdas datetime de openpyxl.
    assert parse_date("3/9/2026") == date(2026, 9, 3)
    assert parse_date("25/08/2026") == date(2026, 8, 25)
    assert parse_date("2026-09-03T00:00:00") == date(2026, 9, 3)
    assert parse_date(datetime(2026, 9, 3, 10, 0)) == date(2026, 9, 3)
    with pytest.raises(ValueError):
        parse_date("31/13/2026")

    # Fichero mixto (xlsx) con ambos formatos en la misma columna.
    content = _xlsx([
        _row("3/9/2026", "ABONO TRANSFERENCIA DE DM DOCUMENT MATERIEL", "1.314,47",
             "12.345,67", "", "DM DOCUMENT MATE"),
        _row("2/9/2026", "RECIBO LUZ", "-1,500.00", "11.031,20"),
        [datetime(2026, 9, 1), "ABONO TRANSFERENCIA DE HIRSCH ARMB-NDER GMBH",
         datetime(2026, 9, 1), 6950.22, 12531.2, "", "", None, None, None],
    ])
    res = _import(session_factory, content)
    assert res["imported"] == 3
    page = _movements(session_factory, only_incoming=False)
    got = {i["concepto"]: (i["fecha_oper"], i["importe"], i["saldo"]) for i in page["items"]}
    assert got["ABONO TRANSFERENCIA DE DM DOCUMENT MATERIEL"] == ("2026-09-03", 1314.47, 12345.67)
    assert got["RECIBO LUZ"] == ("2026-09-02", -1500.0, 11031.2)
    assert got["ABONO TRANSFERENCIA DE HIRSCH ARMB-NDER GMBH"] == ("2026-09-01", 6950.22, 12531.2)
    # REFERENCIA 2 y el pagador extraído del concepto se guardan enteros.
    dm = _by_concepto(page, "DM DOCUMENT")
    assert dm["referencia2"] == "DM DOCUMENT MATE"
    assert dm["payer_name"] == "DM DOCUMENT MATERIEL"

    # CSV (separador «;», mismas cabeceras) también se lee.
    csv_text = "\n".join([
        *HEADER_LINES, "",
        ";".join(COLUMNS),
        "4/9/2026;ABONO TRANSFERENCIA DE SOLITIUM SL;4/9/2026;3.623,95;15.000,00;;;;;",
        "4/9/2026;COMISION MANTENIMIENTO;4/9/2026;-12,50;14.987,50;;;;;",
    ])
    res = _import(session_factory, csv_text.encode("utf-8"), "extracto.csv")
    assert res["imported"] == 2
    page = _movements(session_factory, only_incoming=False)
    assert _by_concepto(page, "SOLITIUM")["importe"] == 3623.95
    assert _by_concepto(page, "COMISION")["importe"] == -12.5


def test_import_fails_visibly_on_unparseable_row(session_factory, http, account_id) -> None:
    content = _xlsx([
        _row("3/9/2026", "ABONO TRANSFERENCIA DE SOLITIUM SL", "3.623,95", "10.000,00"),
        _row("3/9/2026", "FILA ROTA", "tres mil", "10.000,00"),
    ])
    # Fila 7 del fichero: 3 líneas de cabecera + 1 en blanco + columnas + 1 dato.
    with pytest.raises(ParseError) as exc:
        _import(session_factory, content)
    assert exc.value.row == 7
    assert "fila 7" in str(exc.value)
    assert "no numérico" in str(exc.value)
    # No se importa NADA a medias: ni la fila buena.
    with session_factory() as s:
        assert s.scalar(select(func.count(BankMovement.id))) == 0

    r = http.post(
        "/api/erp/bank/import",
        files={"file": ("extracto.xlsx", content, "application/octet-stream")},
        headers=auth_headers(http, "pedidos"),
    )
    assert r.status_code == 422, r.text
    assert r.json()["detail"]["code"] == "parse_error"
    assert r.json()["detail"]["row"] == 7

    # Fecha imposible y concepto vacío también fallan con su fila.
    with pytest.raises(ParseError) as exc:
        _import(session_factory, _xlsx([_row("31/13/2026", "X", "1,00", "1,00")]))
    assert exc.value.row == 6
    with pytest.raises(ParseError) as exc:
        _import(session_factory, _xlsx([_row("1/9/2026", "", "1,00", "1,00")]))
    assert "concepto" in str(exc.value)
    # Sin fila de columnas reconocible → error claro (no un import vacío).
    with pytest.raises(ParseError):
        read_statement(_xlsx([], columns=["A", "B", "C"]), "raro.xlsx")


# --- Parte C: casado ----------------------------------------------------------------


def test_only_incoming_movements_are_candidates(session_factory, account_id) -> None:
    content = _xlsx([
        _row("3/9/2026", "ABONO TRANSFERENCIA DE SOLITIUM SL", "3.623,95", "10.000,00"),
        _row("2/9/2026", "RECIBO LUZ", "-1.500,00", "6.376,05"),
        # Un PAGO saliente con el mismo importe que una factura: jamás candidato.
        _row("1/9/2026", "TRANSFERENCIA A PROVEEDOR SOLITIUM SL", "-3.623,95", "7.876,05"),
    ])
    assert _import(session_factory, content)["imported"] == 3
    fake = _fake()
    stats = _match(session_factory, fake)
    assert stats["candidates"] == 1
    assert stats["proposed"] == 1
    with session_factory() as s:
        for mov in s.scalars(select(BankMovement)).all():
            n = s.scalar(select(func.count(BankReconciliation.id))
                         .where(BankReconciliation.movement_id == mov.id))
            assert (n > 0) == (float(mov.importe) > 0), mov.concepto
    # La bandeja de revisión solo lista entradas.
    page = _movements(session_factory)
    assert [i["importe"] for i in page["items"]] == [3623.95]
    assert page["counters"]["pending"] == 1
    # Y el motor, directamente, ignora importes ≤ 0.
    assert propose({"importe": -3623.95, "concepto": "X", "fecha_oper": date(2026, 9, 1)},
                   _invoices(fake)) == []


def test_excluded_concepts_are_not_proposed(session_factory, account_id) -> None:
    amt, bal = "3.623,95", "10.000,00"
    content = _xlsx([
        _row("3/9/2026", "ABONO TPV 123456 LIQUIDACION", amt, bal),
        _row("3/9/2026", "TRASPASO DESDE CUENTA 5449", amt, "13.623,95"),
        _row("3/9/2026", "SCALAPAY SETTLEMENT 2026-09", amt, "17.247,90"),
        _row("3/9/2026", "ANUL.ADEUDO RECIBO 000123", amt, "20.871,85"),
        _row("3/9/2026", "PAYPAL EUROPE S.A.R.L. ET CIE", amt, "24.495,80"),
        _row("3/9/2026", "DEVOLUCION TRANSFERENCIA", amt, "28.119,75"),
        _row("3/9/2026", "ABONO TRANSFERENCIA DE SOLITIUM SL", amt, "31.743,70"),
    ])
    assert _import(session_factory, content)["imported"] == 7
    fake = _fake()
    stats = _match(session_factory, fake)
    assert stats["candidates"] == 7
    assert stats["excluded"] == 6
    assert stats["proposed"] == 1
    page = _movements(session_factory)
    for item in page["items"]:
        assert bool(item["proposals"]) == ("SOLITIUM" in item["concepto"]), item["concepto"]

    # Las exclusiones son reglas CONFIGURABLES: borrar la de Scalapay hace que
    # ese abono vuelva a ser candidato en el siguiente casado.
    with session_factory() as s:
        rules = service.list_rules(s)
        assert {r["pattern"] for r in rules} == set(DEFAULT_EXCLUSION_PATTERNS)
        scalapay = next(r for r in rules if r["pattern"] == "SCALAPAY")
        service.delete_rule(s, scalapay["id"])
    stats = _match(session_factory, fake)
    assert stats["excluded"] == 5
    assert _by_concepto(_movements(session_factory), "SCALAPAY")["proposals"]
    # Y la función pura devuelve el patrón que excluye (o None).
    assert is_excluded("ABONO TPV 1", None, DEFAULT_EXCLUSION_PATTERNS) == "ABONO TPV"
    assert is_excluded("ABONO TRANSFERENCIA DE SOLITIUM SL", "SOLITIUM SL",
                       DEFAULT_EXCLUSION_PATTERNS) is None


def test_exact_amount_and_name_gives_high_confidence(session_factory, account_id) -> None:
    invoices = _invoices(_fake())
    mov = {"importe": 3623.95, "fecha_oper": date(2026, 9, 3),
           "concepto": "ABONO TRANSFERENCIA DE SOLITIUM SL"}
    props = propose(mov, invoices)
    assert len(props) == 1
    assert props[0]["numero"] == "1-260720"
    assert props[0]["confidence"] == "alta"
    assert props[0]["importe"] == 3623.95
    assert "importe exacto" in props[0]["reason"]
    assert "nombre" in props[0]["reason"]
    # Importe exacto pero el pagador no casa → media (se avisa en el motivo).
    props = propose({**mov, "concepto": "ABONO TRANSFERENCIA DE PEPE LOPEZ"}, invoices)
    assert props[0]["numero"] == "1-260720"
    assert props[0]["confidence"] == "media"
    assert "sin confirmación de nombre" in props[0]["reason"]
    # Cobro ANTERIOR a la factura → no plausible → sin propuesta (no se inventa).
    assert propose({**mov, "fecha_oper": date(2026, 7, 1)}, invoices) == []
    # Una factura ya cobrada (ESTFAC=2) nunca es candidata.
    assert propose({**mov, "importe": 192.39,
                    "concepto": "ABONO TRANSFERENCIA DE Cliente cobrado"}, invoices) == []
    # Persistido: la propuesta guarda confianza y motivo.
    content = _xlsx([_row("3/9/2026", "ABONO TRANSFERENCIA DE SOLITIUM SL",
                          "3.623,95", "10.000,00")])
    _import(session_factory, content)
    stats = _match(session_factory, _fake())
    assert stats["by_confidence"] == {"alta": 1, "media": 0, "baja": 0}
    item = _movements(session_factory)["items"][0]
    assert item["confidence"] == "alta"
    assert item["proposals"][0]["numero"] == "1-260720"
    assert item["proposals"][0]["status"] == "proposed"


def test_name_match_tolerates_broken_accents_and_truncation() -> None:
    # Acento roto en el extracto («Armb-nder» por «Armbänder»).
    assert names_match("Hirsch Armb-nder GmbH", "Hirsch Armbänder GmbH")
    assert names_match("HIRSCH ARMBANDER", "Hirsch Armbänder GmbH")
    # REFERENCIA 2 truncada a ~16 caracteres.
    assert names_match("DM DOCUMENT MATE", "DM Document Materiel SA")
    # Formas societarias y puntuación ignoradas.
    assert names_match("SOLITIUM, S.L.", "SOLITIUM SL")
    assert normalize_name("Hirsch Armbänder GmbH") == "HIRSCH ARMBANDER"
    # Y NO casa lo que no se parece (mejor sin propuesta que inventar).
    assert not names_match("Lorenz, Wolfgang", "SOLITIUM SL")
    assert not names_match("BOMEDIA SL", "BOPRINT SL")
    assert not names_match("", "SOLITIUM SL")
    # Pagador: del concepto o, si no, de REFERENCIA 2.
    assert extract_payer("ABONO TRANSFERENCIA DE Lorenz, Wolfgang") == "Lorenz, Wolfgang"
    assert extract_payer("TRANSFERENCIA RECIBIDA DE DM DOCUMENT MATERIEL SA") == (
        "DM DOCUMENT MATERIEL SA")
    assert extract_payer("ABONO TRANSFERENCIA", "DM DOCUMENT MATE") == "DM DOCUMENT MATE"
    assert extract_payer("RECIBO LUZ") is None
    # De punta a punta: nombre solo en REFERENCIA 2 (truncado) + importe exacto.
    invoices = _invoices(_fake())
    props = propose({"importe": 1314.47, "fecha_oper": date(2026, 9, 3),
                     "concepto": "ABONO TRANSFERENCIA", "referencia2": "DM DOCUMENT MATE"},
                    invoices)
    assert [p["numero"] for p in props] == ["1-260700"]
    assert props[0]["confidence"] == "alta"


def test_partial_payment_proposal(session_factory, account_id) -> None:
    # Caso real: Hirsch paga 6.950,22 de la 1-260649 (13.900,45).
    invoices = _invoices(_fake())
    props = propose({"importe": 6950.22, "fecha_oper": date(2026, 8, 13),
                     "concepto": "ABONO TRANSFERENCIA DE Hirsch Armb-nder GmbH"}, invoices)
    assert len(props) == 1
    assert props[0]["numero"] == "1-260649"
    assert props[0]["importe"] == 6950.22
    assert props[0]["confidence"] == "media"
    assert props[0]["reason"].startswith("cobro parcial (6950.22 de 13900.45")
    # Sin nombre que lo respalde, un importe menor NO se asigna a ciegas.
    assert propose({"importe": 6950.22, "fecha_oper": date(2026, 8, 13),
                    "concepto": "ABONO TRANSFERENCIA DE PEPE LOPEZ"}, invoices) == []
    # Persistido con su motivo.
    content = _xlsx([_row("13/8/2026", "ABONO TRANSFERENCIA DE Hirsch Armb-nder GmbH",
                          "6.950,22", "20.000,00")])
    _import(session_factory, content)
    stats = _match(session_factory, _fake())
    assert stats["by_confidence"]["media"] == 1
    item = _movements(session_factory)["items"][0]
    assert item["proposals"][0]["numero"] == "1-260649"
    assert "cobro parcial" in item["proposals"][0]["reason"]


def test_multiple_invoices_summing_to_transfer(session_factory, account_id) -> None:
    invoices = _invoices(_fake())
    props = propose({"importe": 350.50, "fecha_oper": date(2026, 9, 3),
                     "concepto": "ABONO TRANSFERENCIA DE MOVIATICOS SL"}, invoices)
    assert {p["numero"] for p in props} == {"1-260701", "1-260702"}
    assert sorted(p["importe"] for p in props) == [100.0, 250.5]
    assert all(p["confidence"] == "alta" for p in props)
    assert all("suma de 2 facturas" in p["reason"] for p in props)
    # Sin nombre: la suma sigue siendo única → media, avisando.
    props = propose({"importe": 350.50, "fecha_oper": date(2026, 9, 3),
                     "concepto": "ABONO TRANSFERENCIA DE DESCONOCIDO"}, invoices)
    assert {p["numero"] for p in props} == {"1-260701", "1-260702"}
    assert all(p["confidence"] == "media" for p in props)
    # Persistido: un movimiento, dos filas de propuesta; confirmar concilia las dos.
    content = _xlsx([_row("3/9/2026", "ABONO TRANSFERENCIA DE MOVIATICOS SL",
                          "350,50", "10.000,00")])
    _import(session_factory, content)
    _match(session_factory, _fake())
    item = _movements(session_factory)["items"][0]
    assert len(item["proposals"]) == 2
    with session_factory() as s:
        out = service.confirm_movement(s, item["id"], user_id=None)
    assert out["status"] == "reconciled"
    assert {r["numero"] for r in out["reconciled"]} == {"1-260701", "1-260702"}


# --- Parte D: revisión (la persona decide) --------------------------------------------


def test_nothing_is_auto_confirmed(session_factory, http) -> None:
    admin = auth_headers(http, "admin")
    r = http.post("/api/erp/bank/accounts", json={
        "name": "Sabadell", "iban": "ES33 0081 0202 1300 0117 1918", "bank_name": "Sabadell",
    }, headers=admin)
    assert r.status_code == 201, r.text
    content = _xlsx([
        _row("3/9/2026", "ABONO TRANSFERENCIA DE SOLITIUM SL", "3.623,95", "10.000,00"),
        _row("3/9/2026", "ABONO TRANSFERENCIA DE MOVIATICOS SL", "350,50", "10.350,50"),
        _row("13/8/2026", "ABONO TRANSFERENCIA DE Hirsch Armb-nder GmbH", "6.950,22",
             "17.300,72"),
        _row("2/9/2026", "ABONO TRANSFERENCIA DE NADIE CONOCIDO", "77,77", "17.378,49"),
    ])
    with _patched(_fake()):
        r = http.post(
            "/api/erp/bank/import",
            files={"file": ("extracto.xlsx", content, "application/octet-stream")},
            headers=auth_headers(http, "pedidos"),
        )
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["imported"] == 4
    assert body["matching"]["ok"] is True
    assert body["matching"]["by_confidence"] == {"alta": 2, "media": 1, "baja": 0}
    assert body["matching"]["no_proposal"] == 1

    # Tras importar + casar: TODO sigue pendiente, incluso lo de confianza alta.
    r = http.get("/api/erp/bank/movements", headers=auth_headers(http, "user"))
    page = r.json()
    assert page["total"] == 4
    assert all(i["status"] == "pending" for i in page["items"])
    assert all(i["reconciled"] == [] for i in page["items"])
    assert page["counters"] == {"pending": 4, "reconciled": 0, "discarded": 0,
                                "pending_amount": 11002.44}
    with session_factory() as s:
        assert s.scalar(select(func.count(BankReconciliation.id))
                        .where(BankReconciliation.status == "confirmed")) == 0
    # Filtro por confianza (incluida «none» = sin propuesta).
    r = http.get("/api/erp/bank/movements?confidence=alta", headers=auth_headers(http, "user"))
    assert {i["concepto"].split(" DE ")[1] for i in r.json()["items"]} == {
        "SOLITIUM SL", "MOVIATICOS SL"}
    r = http.get("/api/erp/bank/movements?confidence=none", headers=auth_headers(http, "user"))
    assert [i["importe"] for i in r.json()["items"]] == [77.77]

    # Confirmar un movimiento SIN propuesta → 409 (no se inventa).
    nadie = _by_concepto(page, "NADIE")
    r = http.post(f"/api/erp/bank/movements/{nadie['id']}/confirm",
                  headers=auth_headers(http, "pedidos"))
    assert r.status_code == 409
    assert r.json()["detail"]["code"] == "no_proposal"
    # Ver no basta para decidir: rol «user» no confirma.
    r = http.post("/api/erp/bank/movements/confirm-high", headers=auth_headers(http, "user"))
    assert r.status_code == 403

    # «Confirmar todas las de confianza alta» es un CLIC EXPLÍCITO de Bart.
    r = http.post("/api/erp/bank/movements/confirm-high", headers=auth_headers(http, "pedidos"))
    assert r.status_code == 200
    assert r.json()["confirmed"] == 2
    r = http.get("/api/erp/bank/movements", headers=auth_headers(http, "user"))
    page = r.json()
    assert page["counters"]["reconciled"] == 2
    assert page["counters"]["pending"] == 2
    assert _by_concepto(page, "Hirsch")["status"] == "pending"       # media: sigue
    # «No es cobro» con motivo, y reabrir.
    r = http.post(f"/api/erp/bank/movements/{nadie['id']}/discard",
                  json={"reason": "fianza devuelta"}, headers=auth_headers(http, "pedidos"))
    assert r.status_code == 200
    assert r.json()["status"] == "discarded"
    assert r.json()["discard_reason"] == "fianza devuelta"
    r = http.post(f"/api/erp/bank/movements/{nadie['id']}/reopen",
                  headers=auth_headers(http, "pedidos"))
    assert r.json()["status"] == "pending"
    # Recalcular propuestas NO toca lo ya confirmado.
    with _patched(_fake()):
        r = http.post("/api/erp/bank/match", headers=auth_headers(http, "pedidos"))
    assert r.status_code == 200
    assert r.json()["candidates"] == 2
    r = http.get("/api/erp/bank/movements?status=reconciled", headers=auth_headers(http, "user"))
    assert r.json()["total"] == 2


def test_learned_rules_apply_to_future_imports(session_factory, account_id) -> None:
    invoices = [
        _fac(1, 260710, 500.00, cliente="SOLITIUM SL", codcli=10),
        _fac(1, 260711, 500.00, cliente="MOVIATICOS SL", codcli=99),
        _fac(1, 260712, 700.00, cliente="SOLITIUM SL", codcli=10),
        _fac(1, 260713, 700.00, cliente="MOVIATICOS SL", codcli=99),
    ]
    fake = _fake(invoices)
    # 1) Un particular paga por SOLITIUM: dos facturas con ese importe → baja.
    _import(session_factory, _xlsx([_row("1/9/2026", "ABONO TRANSFERENCIA DE Lorenz, Wolfgang",
                                         "500,00", "10.500,00")]), "a.xlsx")
    _match(session_factory, fake)
    item = _movements(session_factory)["items"][0]
    assert item["confidence"] == "baja"
    assert {p["numero"] for p in item["proposals"]} == {"1-260710", "1-260711"}
    # Bart elige la de SOLITIUM y pide que se recuerde el pagador.
    with session_factory() as s:
        out = service.reassign_movement(
            s, item["id"],
            [{"serie": 1, "codigo": 260710, "numero": "1-260710", "importe": 500.0,
              "cliente_codigo": "10", "cliente_nombre": "SOLITIUM SL"}],
            user_id=None, learn_payer=True,
        )
        assert out["status"] == "reconciled"
        assert [r["numero"] for r in out["reconciled"]] == ["1-260710"]
        rule = s.scalar(select(BankLearnedRule).where(BankLearnedRule.kind == "payer_to_client"))
        assert rule is not None
        assert rule.pattern == "LORENZ WOLFGANG"
        assert rule.client_codcli == "10"
    # 2) Siguiente importación del mismo pagador: la regla decide → alta, única.
    _import(session_factory, _xlsx([_row("5/9/2026", "ABONO TRANSFERENCIA DE Lorenz, Wolfgang",
                                         "700,00", "11.200,00")]), "b.xlsx")
    stats = _match(session_factory, fake)
    assert stats["by_confidence"]["alta"] == 1
    item = _movements(session_factory, status="pending")["items"][0]
    assert [p["numero"] for p in item["proposals"]] == ["1-260712"]
    assert item["confidence"] == "alta"
    # Un reparto que no cuadra con el importe se rechaza.
    with session_factory() as s, pytest.raises(ValueError):
        service.reassign_movement(
            s, item["id"],
            [{"serie": 1, "codigo": 260712, "importe": 600.0}], user_id=None,
        )
    # 3) Exclusión aprendida: «no es cobro» + recordar → no vuelve a proponerse.
    _import(session_factory, _xlsx([_row("6/9/2026", "ABONO TRANSFERENCIA DE AMAZON PAYMENTS",
                                         "700,00", "11.900,00")]), "c.xlsx")
    _match(session_factory, fake)
    amazon = _by_concepto(_movements(session_factory, status="pending"), "AMAZON")
    assert amazon["proposals"]                       # antes: dos candidatas en baja
    with session_factory() as s:
        service.discard_movement(s, amazon["id"], "liquidación marketplace", None, learn=True)
        rules = service.list_rules(s)
        assert any(r["kind"] == "exclude_pattern" and r["pattern"] == "AMAZON PAYMENTS"
                   for r in rules)
    _import(session_factory, _xlsx([_row("7/9/2026", "ABONO TRANSFERENCIA DE AMAZON PAYMENTS",
                                         "700,00", "12.600,00")]), "d.xlsx")
    stats = _match(session_factory, fake)
    assert stats["excluded"] == 1
    later = _by_concepto(_movements(session_factory, desde=date(2026, 9, 7)), "AMAZON")
    assert later["proposals"] == []
    assert later["status"] == "pending"


# --- Parte E: exportación ---------------------------------------------------------------


def test_export_preserves_original_format_and_fills_columns(
    session_factory, http, account_id,
) -> None:
    rows = [
        _row("3/9/2026", "ABONO TRANSFERENCIA DE SOLITIUM SL", "3.623,95", "13.974,45",
             "REF-1", "SOLITIUM"),
        _row("2/9/2026", "ABONO TRANSFERENCIA DE MOVIATICOS SL", "350,50", "10.350,50"),
        _row("1/9/2026", "RECIBO LUZ", "-1.500,00", "10.000,00"),
    ]
    content = _xlsx(rows)
    original = bytes(content)
    _import(session_factory, content)
    _match(session_factory, _fake())
    assert content == original                  # el fichero subido no se toca
    page = _movements(session_factory)
    with session_factory() as s:
        # Solo se confirma MOVIATICOS (2 facturas); SOLITIUM queda propuesta.
        service.confirm_movement(s, _by_concepto(page, "MOVIATICOS")["id"], user_id=None)

    r = http.get(f"/api/erp/bank/export?account_id={account_id}",
                 headers=auth_headers(http, "user"))
    assert r.status_code == 200, r.text
    assert r.headers["content-disposition"].endswith('extracto_conciliado.xlsx"')
    ws = load_workbook(io.BytesIO(r.content), read_only=True).worksheets[0]
    grid = [list(row) for row in ws.iter_rows(values_only=True)]
    # Misma cabecera, misma fila de columnas en el MISMO orden.
    assert [g[0] for g in grid[:3]] == HEADER_LINES
    assert all(c is None for c in grid[3])
    assert [c for c in grid[4] if c is not None] == COLUMNS
    data = {g[1]: g for g in grid[5:]}
    fac = COLUMNS.index("FACTURA")
    movi = data["ABONO TRANSFERENCIA DE MOVIATICOS SL"]
    assert movi[fac] == "1-260701, 1-260702"           # varias → coma
    assert movi[COLUMNS.index("IMPORTE")] == "350,50"   # formato original intacto
    assert movi[COLUMNS.index("FECHA OPER")] == "2/9/2026"
    soli = data["ABONO TRANSFERENCIA DE SOLITIUM SL"]
    assert soli[fac] is None                             # propuesta ≠ confirmada
    assert soli[COLUMNS.index("REFERENCIA 2")] == "SOLITIUM"
    luz = data["RECIBO LUZ"]
    assert luz[COLUMNS.index("IMPORTE")] == "-1.500,00"
    assert luz[fac] is None
    assert len(data) == 3                                 # también los pagos, sin tocar

    # Rango de fechas.
    r = http.get(f"/api/erp/bank/export?account_id={account_id}&desde=2026-09-02"
                 "&hasta=2026-09-02", headers=auth_headers(http, "user"))
    assert r.headers["content-disposition"].endswith(
        'extracto_conciliado_2026-09-02_2026-09-02.xlsx"')
    ws = load_workbook(io.BytesIO(r.content), read_only=True).worksheets[0]
    grid = [list(row) for row in ws.iter_rows(values_only=True)]
    assert [g[1] for g in grid[5:]] == ["ABONO TRANSFERENCIA DE MOVIATICOS SL"]
