"""Seguimiento — columna «Courier» aparte y «Envío» solo con el estado.

- Envío: lista CERRADA de estados; Courier: la agencia de Genei, el courier
  apuntado en la Cola SAT («otro courier» si no se apuntó) o «—» sin envío.
- Excel y hoja con 20 columnas, «Courier» detrás de «Envío» y la «id» la última.
- Migración de la hoja (19 → 20 columnas) hecha por el espejo: una sola vez,
  idempotente, sin perder ninguna celda (recuento por columna antes/después),
  con el histórico manual intacto (no se reinterpreta). Y las filas que el
  espejo guarda en la BD, puestas al día en la misma pasada.

Sin red: el transporte de Sheets es un doble que simula `insertDimension`.
"""
from __future__ import annotations

import copy
import io
import json
from collections.abc import Generator
from datetime import UTC, date, datetime
from typing import Any

import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

import app.main  # noqa: F401
from app.db.base import Base
from app.erp.drive_managed import (
    DEFAULT_MANAGED_TAB,
    SEPARATOR_PEDIDOS,
    dates_to_serial,
    is_separator,
    push_managed_tabs,
    recuento_por_columna,
)
from app.erp.drive_sheets import DriveSyncError
from app.erp.integrations.genei.service import set_genei_state
from app.erp.models import (
    Order,
    SeguimientoLegacy,
    SeguimientoManual,
    SeguimientoOverride,
    SeguimientoSnapshot,
    SeguimientoSyncMeta,
    TransportStatus,
)
from app.erp.models.seguimiento_mirror import KIND_LEGACY, KIND_MANUAL, KIND_ORDER
from app.erp.seguimiento import (
    ID_INDEX,
    SEGUIMIENTO_COLUMNS_V2,
    SITUACION_LABELS,
    _envio_label,
    courier_label,
    envio_vocabulary,
    export_xlsx,
    filter_rows,
    row_to_pedidos_values,
)
from app.erp.seguimiento_mirror import (
    BLOQUEADAS,
    META_FORMATO_BD,
    _listas_cerradas,
    protection_requests,
)
from app.erp.shipping_courier import set_external_state

HISTORICA = "Pedidos Bomedia 2020-2026"
TAB = DEFAULT_MANAGED_TAB


def _c(nombre: str) -> int:
    return SEGUIMIENTO_COLUMNS_V2.index(nombre)


COURIER = _c("Courier")
ENVIO = _c("Envío")
#: Cabecera de la hoja ANTES de «Courier» (19 columnas, la «id» la última).
CABECERA_19 = [c for c in SEGUIMIENTO_COLUMNS_V2 if c != "Courier"]


def _a_19(fila: list[Any]) -> list[Any]:
    """Una fila del formato actual → la que habría escrito la versión anterior
    (sin la celda de «Courier»)."""
    r = list(fila)
    del r[COURIER]
    return r


# --- dobles ---------------------------------------------------------------------


class FakeTabs:
    """Pestañas en memoria que entienden `insertDimension` (COLUMNS) y el
    `updateCells` de una celda, como haría Sheets."""

    def __init__(self, tabs: dict[str, list[list[Any]]]) -> None:
        self.tabs = tabs
        self.written: dict[str, list[list[Any]]] = {}
        self.formats: dict[str, list[list[dict[str, Any]]]] = {}
        self.inserts = 0

    def tab_titles(self) -> list[str]:
        return list(self.tabs)

    def first_tab_title(self) -> str:
        return next(iter(self.tabs))

    def ensure_tab(self, title: str) -> int:
        self.tabs.setdefault(title, [])
        return list(self.tabs).index(title)

    def replace_tab(self, title: str, rows: list[list[Any]], *, raw: bool = False) -> None:
        self.tabs[title] = copy.deepcopy(rows)
        self.written[title] = copy.deepcopy(rows)

    def tab_values(self, title: str, *, raw: bool = False) -> list[list[Any]]:
        filas = self.tabs.get(title, [])
        return [list(r) for r in filas] if raw else [[str(c) for c in r] for r in filas]

    def format_tab(self, title: str, requests: list[dict[str, Any]]) -> None:
        self.formats.setdefault(title, []).append(requests)
        for req in requests:
            if "insertDimension" in req:
                rng = req["insertDimension"]["range"]
                assert rng["dimension"] == "COLUMNS"
                i0, i1 = rng["startIndex"], rng["endIndex"]
                for fila in self.tabs[title]:
                    if len(fila) > i0:
                        fila[i0:i0] = [""] * (i1 - i0)
                self.inserts += 1
            elif "updateCells" in req:
                rng = req["updateCells"]["range"]
                fila = self.tabs[title][rng["startRowIndex"]]
                col = rng["startColumnIndex"]
                valor = req["updateCells"]["rows"][0]["values"][0]["userEnteredValue"]
                fila.extend([""] * (col + 1 - len(fila)))
                fila[col] = valor["stringValue"]

    def requests(self, title: str) -> list[dict[str, Any]]:
        return [r for lote in self.formats.get(title, []) for r in lote]


@pytest.fixture()
def factory() -> Generator[sessionmaker, None, None]:
    engine = create_engine("sqlite+pysqlite:///:memory:",
                           connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    yield sessionmaker(bind=engine)
    Base.metadata.drop_all(engine)


def _order(s: Session, numero: str, **kw: Any) -> Order:
    o = Order(order_number=numero, payment_status="paid", **kw)
    s.add(o)
    s.flush()
    return o


def _genei(o: Order, *, courier: str | None = "Ctt Premium", step: str | None = None) -> None:
    state: dict[str, Any] = {"shipment_code": f"G-{o.order_number}", "courier": courier}
    if step:
        state["carrier_step"] = step
    set_genei_state(o, state)


# --- Envío (lista cerrada) y Courier, pedido a pedido ------------------------------


def test_genei_entregado_envio_entregado_y_courier_la_agencia(factory):
    with factory() as s:
        o = _order(s, "BOP-1", transport_status=TransportStatus.DELIVERED)
        _genei(o, step="delivered")
        assert _envio_label(o) == "Entregado"
        assert courier_label(o) == "Ctt Premium"          # tal cual la da Genei


def test_externo_con_ups_marcado_recogido(factory):
    with factory() as s:
        o = _order(s, "BOP-2", transport_status=TransportStatus.IN_TRANSIT,
                   tracking_number="1ZFV75016835455899")
        set_external_state(o, {"courier": "UPS"})
        assert _envio_label(o) == "Enviado"                # BoHub no ve sus escaneos
        assert courier_label(o) == "UPS"


def test_externo_sin_courier_es_otro_courier(factory):
    with factory() as s:
        o = _order(s, "BOP-3", transport_status=TransportStatus.IN_TRANSIT)
        assert _envio_label(o) == "Enviado"
        assert courier_label(o) == "otro courier"


def test_sin_envio_no_aplica_y_sin_courier(factory):
    with factory() as s:
        o = _order(s, "BOP-4", shipping_not_required=True,
                   transport_status=TransportStatus.NOT_SHIPPED)
        assert _envio_label(o) == "No aplica"
        assert courier_label(o) == "—"
        # Aunque llevara un courier apuntado: no se envía.
        set_external_state(o, {"courier": "MRW"})
        assert courier_label(o) == "—"


def test_aun_sin_salir_sin_enviar_y_sin_courier(factory):
    with factory() as s:
        o = _order(s, "BOP-5", transport_status=TransportStatus.NOT_SHIPPED)
        assert (_envio_label(o), courier_label(o)) == ("Sin enviar", "—")
        # La etiqueta subida a mano (otro courier) aún no es «enviado».
        o.transport_status = TransportStatus.LABEL_CREATED
        assert (_envio_label(o), courier_label(o)) == ("Sin enviar", "—")
        # Con el courier ya apuntado en la Cola SAT, se ve aunque no haya salido.
        set_external_state(o, {"courier": "MRW"})
        assert (_envio_label(o), courier_label(o)) == ("Sin enviar", "MRW")


def test_otro_courier_entregado_o_incidencia_marcado_en_la_ficha(factory):
    with factory() as s:
        o = _order(s, "BOP-6", transport_status=TransportStatus.DELIVERED)
        set_external_state(o, {"courier": "Seitrans"})
        assert (_envio_label(o), courier_label(o)) == ("Entregado", "Seitrans")
        o.transport_status = TransportStatus.INCIDENT
        assert _envio_label(o) == "Incidencia"
        o.transport_status = TransportStatus.RETURNED      # devuelto = incidencia
        assert _envio_label(o) == "Incidencia"


def test_genei_sin_escaneo_usa_el_transporte(factory):
    with factory() as s:
        o = _order(s, "BOP-7", transport_status=TransportStatus.LABEL_CREATED)
        _genei(o)
        assert _envio_label(o) == "Pendiente de entrada en red"   # etiqueta hecha
        o.transport_status = TransportStatus.IN_TRANSIT
        assert _envio_label(o) == "En tránsito"
        o.transport_status = TransportStatus.NOT_SHIPPED          # sin pagar aún
        assert _envio_label(o) == "Sin enviar"
        # Agencia desconocida (raro): se dice que es Genei, no «—».
        o2 = _order(s, "BOP-8", transport_status=TransportStatus.IN_TRANSIT)
        _genei(o2, courier=None)
        assert courier_label(o2) == "Genei"


@pytest.mark.parametrize("transporte", list(TransportStatus))
@pytest.mark.parametrize("tipo", ["genei", "externo", "nada", "sin_envio"])
def test_envio_siempre_de_la_lista_cerrada(factory, transporte, tipo):
    """Sea cual sea el transporte y el tipo de envío, Envío sale de la lista
    cerrada y nunca lleva el courier dentro."""
    with factory() as s:
        o = _order(s, "BOP-9", transport_status=transporte,
                   shipping_not_required=(tipo == "sin_envio"))
        if tipo == "genei":
            _genei(o)
        elif tipo == "externo":
            set_external_state(o, {"courier": "UPS"})
        assert _envio_label(o) in envio_vocabulary()
        assert "·" not in _envio_label(o)


def test_los_escaneos_reales_son_de_la_lista_cerrada():
    from app.erp.integrations.genei.tracking import CARRIER_STEP_LABELS

    assert set(CARRIER_STEP_LABELS.values()) <= set(envio_vocabulary())
    assert envio_vocabulary()[-2:] == ["Enviado", "No aplica"]
    assert len(envio_vocabulary()) == 10


# --- filtro «Transportista» = columna Courier ------------------------------------


def _fila_vista(numero: str, courier: str) -> dict[str, Any]:
    return {
        "order_number": numero, "courier": courier, "transportista": "Genei",
        "excluido": False, "oculto_por_estado": False, "forzado": False,
        "en_curso": True, "fecha": "2026-09-01", "situacion": "listo",
    }


def test_el_filtro_transportista_filtra_por_courier():
    filas = [_fila_vista("A", "Ctt Premium"), _fila_vista("B", "CTT Express"),
             _fila_vista("C", "UPS"), _fila_vista("D", "otro courier"),
             _fila_vista("E", "—")]

    def nums(texto: str) -> list[str]:
        return sorted(r["order_number"] for r in filter_rows(filas, transportista=texto))

    assert nums("ctt") == ["A", "B"]           # contiene, sin mayúsculas
    assert nums("UPS") == ["C"]
    assert nums("otro") == ["D"]
    assert nums("ctt premium") == ["A"]
    # El viejo `transportista` (el carrier de Genei) ya no es el criterio.
    assert nums("genei") == []


# --- Excel y hoja con 20 columnas -------------------------------------------------


def _row(order: Order, **extra: Any) -> dict[str, Any]:
    base = {
        "id": order.id, "situacion": "listo", "situacion_label": SITUACION_LABELS["listo"],
        "order_number": order.order_number, "fecha": "2026-09-01", "cliente": "Acme SL",
        "origen_label": "WEB", "productos": "Cabezal", "importe": 121.0,
        "empresa_serie": "1 · Bomedia", "factura": "", "fecha_factura": None,
        "factura_enviada": None, "cobro_label": "—", "preparacion": "Listo",
        "envio": "Enviado", "courier": "UPS", "recogido": "2026-09-02",
        "tracking": order.tracking_number or "", "serie_whiterip": "",
        "nota_incidencia": "", "envio_genei": False,
    }
    base.update(extra)
    return base


def test_excel_veinte_columnas_courier_detras_de_envio_e_id_la_ultima_oculta(factory):
    from openpyxl import load_workbook

    with factory() as s:
        o = _order(s, "BOP-10")
        datos = export_xlsx([_row(o)])
    ws = load_workbook(io.BytesIO(datos))["Pedidos"]
    cabecera = [c.value for c in ws[1]]
    assert len(cabecera) == 20 and cabecera == SEGUIMIENTO_COLUMNS_V2
    assert cabecera[ENVIO:ENVIO + 3] == ["Envío", "Courier", "Fecha recogido"]
    assert cabecera[-1] == "id"
    assert ws.column_dimensions["T"].hidden is True             # la «id», oculta
    assert ws.column_dimensions["O"].hidden is not True         # Courier, visible
    fila = [c.value for c in ws[2]]
    assert fila[ENVIO] == "Enviado" and fila[COURIER] == "UPS"
    assert fila[_c("Fecha recogido")].date() == date(2026, 9, 2)   # sigue siendo fecha
    assert fila[-1] == o.id


def test_hoja_veinte_columnas_y_courier_bloqueada_en_filas_de_bohub():
    assert len(SEGUIMIENTO_COLUMNS_V2) == 20 and SEGUIMIENTO_COLUMNS_V2[-1] == "id"
    assert ID_INDEX == 19
    assert "Courier" in BLOQUEADAS                   # la rellena BoHub
    # Validación de Envío = la lista cerrada; Courier, sin desplegable.
    assert _listas_cerradas()[ENVIO] == envio_vocabulary()
    assert COURIER not in _listas_cerradas()


def test_proteccion_courier_en_filas_de_bohub_y_libre_en_las_manuales():
    grid = [
        list(SEGUIMIENTO_COLUMNS_V2),
        ["Listo", "BOP-1"] + [""] * (ID_INDEX - 2) + ["ord-1"],
        ["Listo", "M-1", "", "", "MANUAL"] + [""] * (ID_INDEX - 5) + ["man-1"],
    ]
    req = protection_requests(grid, {"ord-1": KIND_ORDER, "man-1": KIND_MANUAL}, set())
    tramos = [
        (r["addProtectedRange"]["protectedRange"]["range"]["startRowIndex"],
         r["addProtectedRange"]["protectedRange"]["range"]["startColumnIndex"],
         r["addProtectedRange"]["protectedRange"]["range"]["endColumnIndex"])
        for r in req if "addProtectedRange" in r
    ]
    assert any(fila == 1 and c0 <= COURIER < c1 for fila, c0, c1 in tramos)
    assert not any(fila == 2 for fila, _c0, _c1 in tramos)       # la manual, abierta
    # La regla naranja apunta a la Nota en su sitio nuevo (columna S).
    regla = next(r for r in req if "addConditionalFormatRule" in r)
    formula = regla["addConditionalFormatRule"]["rule"]["booleanRule"]["condition"][
        "values"][0]["userEnteredValue"]
    assert formula.startswith("=REGEXMATCH($S2,")


def test_la_regla_naranja_de_antes_de_courier_tambien_se_quita():
    """Una pestaña migrada aún lleva la regla de la pasada anterior (apuntaba a
    la Nota en la R): se reconoce como del espejo y se sustituye, sin apilar."""
    vieja = {"booleanRule": {"condition": {"values": [
        {"userEnteredValue": '=REGEXMATCH($R2,"\\[⚠ revisar")'}]}}}
    de_persona = {"booleanRule": {"condition": {"values": [
        {"userEnteredValue": '=$A2="Incidencia"'}]}}}
    req = protection_requests([list(SEGUIMIENTO_COLUMNS_V2)], {}, set(),
                              meta={"conditional_formats": [de_persona, vieja]})
    borradas = [r["deleteConditionalFormatRule"]["index"]
                for r in req if "deleteConditionalFormatRule" in r]
    assert borradas == [1]


# --- migración de la hoja 19 → 20 columnas -----------------------------------------


def _hist(numero: str, envio: str, recogido: str, tracking: str, nota: str,
          rid: str) -> list[Any]:
    """Fila del histórico manual en el formato de 19 columnas."""
    fila = [""] * len(CABECERA_19)
    fila[0], fila[1], fila[3] = "Histórico", numero, "Cliente viejo"
    fila[CABECERA_19.index("Envío")] = envio
    fila[CABECERA_19.index("Fecha recogido")] = recogido
    fila[CABECERA_19.index("Tracking")] = tracking
    fila[CABECERA_19.index("Nota / Incidencia")] = nota
    fila[CABECERA_19.index("id")] = rid
    return fila


def _mundo_viejo(s: Session) -> tuple[Order, FakeTabs, dict[str, Any]]:
    """La hoja y la BD tal como las dejó la versión anterior (19 columnas):
    una fila de BoHub, una manual, el separador y tres del histórico manual que
    llevan el courier EN Envío (UPS, FEDEX, DSV)."""
    o = _order(s, "BOP-100", tracking_number="1ZFV75016835455899")
    # Lo que escribía la versión anterior para este pedido (Envío compuesto).
    bohub_20 = dates_to_serial(
        [row_to_pedidos_values(_row(o, envio="Enviado · UPS", courier=""))], (2, 9, 10, 15),
    )[0]
    bohub_19 = _a_19(bohub_20)
    manual_19 = [""] * len(CABECERA_19)
    manual_19[1], manual_19[3], manual_19[4] = "M-1", "Manual SL", "MANUAL"
    manual_19[CABECERA_19.index("Envío")] = "MRW"
    manual_19[CABECERA_19.index("Tracking")] = "T-1"
    manual_19[CABECERA_19.index("Nota / Incidencia")] = "hola"
    manual_19[CABECERA_19.index("id")] = "man-1"
    historico = [
        _hist("V-1", "UPS", "3/2/2023", "TRK-1", "nota vieja", "leg-1"),
        _hist("V-2", "FEDEX", "", "TRK-2", "", "leg-2"),
        _hist("V-3", "DSV", "4/2/2023", "", "llamar antes", "leg-3"),
    ]
    hoja = [CABECERA_19, bohub_19, manual_19,
            [SEPARATOR_PEDIDOS] + [""] * (len(CABECERA_19) - 1), *historico]
    sheets = FakeTabs({HISTORICA: [], TAB: copy.deepcopy(hoja)})
    # La BD del espejo en 19 columnas (sin marca de formato).
    s.add_all([
        SeguimientoSnapshot(row_id=o.id, kind=KIND_ORDER, values_json=json.dumps(bohub_19)),
        SeguimientoSnapshot(row_id="man-1", kind=KIND_MANUAL,
                            values_json=json.dumps(manual_19)),
        SeguimientoManual(row_id="man-1", values_json=json.dumps(manual_19)),
        *[SeguimientoSnapshot(row_id=f"leg-{i}", kind=KIND_LEGACY, values_json="[]")
          for i in (1, 2, 3)],
        *[SeguimientoLegacy(id=f"leg-{i}", row_index=i - 1, numero_raw=h[1],
                            cliente_raw=h[3], raw_json=json.dumps(h[:ID_INDEX - 1]),
                            match_status="synthetic")
          for i, h in enumerate(historico, start=1)],
    ])
    s.commit()
    return o, sheets, {"hoja": hoja, "historico": historico}


def _historico_escrito(sheets: FakeTabs) -> list[list[Any]]:
    filas = sheets.tabs[TAB]
    sep = next(i for i, f in enumerate(filas) if is_separator(f))
    return filas[sep + 1:]


def test_migracion_inserta_courier_sin_perder_celdas_y_el_historico_intacto(factory):
    with factory() as s:
        o, sheets, mundo = _mundo_viejo(s)
        antes = recuento_por_columna(mundo["hoja"])
        res = push_managed_tabs(s, sheets, [_row(o)])
        s.commit()

        # Una sola inserción de columna, con su nombre en la cabecera.
        assert sheets.inserts == 1
        mig = res["migracion_courier"]
        assert mig["estado"] == "hecha"
        assert mig["celdas_despues"] == mig["celdas_antes"] == sum(antes)
        # Recuento por columna: lo de delante igual, lo de detrás corrido una.
        assert mig["por_columna_despues"] == {
            **mig["por_columna_antes"], "Courier": 1,
        }
        escrito = sheets.written[TAB]
        assert escrito[0] == SEGUIMIENTO_COLUMNS_V2       # 20 columnas, id la última

        # Fila de BoHub: Envío solo el estado y Courier aparte.
        fila = next(f for f in escrito if len(f) > ID_INDEX and f[ID_INDEX] == o.id)
        assert fila[ENVIO] == "Enviado" and fila[COURIER] == "UPS"
        # Fila manual: lo tecleado, en su sitio; Courier vacío (editable).
        man = next(f for f in escrito if len(f) > ID_INDEX and f[ID_INDEX] == "man-1")
        assert man[ENVIO] == "MRW" and man[COURIER] == ""
        assert man[_c("Tracking")] == "T-1" and man[_c("Nota / Incidencia")] == "hola"

        # Histórico manual: el courier se queda en Envío (no se reinterpreta),
        # Courier vacío y todo lo de detrás en su columna nueva.
        hist = _historico_escrito(sheets)
        assert [h[ENVIO] for h in hist] == ["UPS", "FEDEX", "DSV"]
        assert [h[COURIER] for h in hist] == ["", "", ""]
        assert [h[_c("Tracking")] for h in hist] == ["TRK-1", "TRK-2", ""]
        assert [h[_c("Nota / Incidencia")] for h in hist] == [
            "nota vieja", "", "llamar antes"]
        assert [h[ID_INDEX] for h in hist] == ["leg-1", "leg-2", "leg-3"]
        assert [h[_c("Fecha recogido")] for h in hist] == [
            44960, "", 44961]                             # como fecha (serial)
        # Recuento del histórico, celda a celda: ni una perdida.
        hist_antes = recuento_por_columna(mundo["historico"])
        hist_despues = recuento_por_columna([list(h) for h in hist])
        assert sum(hist_despues) == sum(hist_antes)
        assert hist_despues[COURIER] == 0

        # El espejo no confunde la migración con ediciones a mano.
        assert res["espejo"]["ediciones_leidas"] == 0
        assert res["espejo"]["tracking_leidos"] == 0
        assert s.scalars(select(SeguimientoOverride)).all() == []
        s.refresh(o)
        assert o.tracking_number == "1ZFV75016835455899"

        # La BD del espejo, al formato nuevo (una vez) y con su marca.
        assert s.get(SeguimientoSyncMeta, META_FORMATO_BD).valor == "20"
        assert res["espejo"]["formato_bd_convertidas"] == 6   # 2 fotos + 1 manual + 3 hist.
        guardada = json.loads(s.scalars(select(SeguimientoManual)).one().values_json)
        assert guardada[ENVIO] == "MRW" and guardada[COURIER] == ""
        assert guardada[_c("Tracking")] == "T-1"
        leg1 = json.loads(s.get(SeguimientoLegacy, "leg-1").raw_json)
        assert leg1[ENVIO] == "UPS" and leg1[COURIER] == ""
        assert leg1[_c("Nota / Incidencia")] == "nota vieja"


def test_migracion_idempotente_la_segunda_pasada_no_inserta_otra_columna(factory):
    with factory() as s:
        o, sheets, _mundo = _mundo_viejo(s)
        push_managed_tabs(s, sheets, [_row(o)])
        s.commit()
        primera = copy.deepcopy(sheets.tabs[TAB])
        manual_bd = s.scalars(select(SeguimientoManual)).one().values_json
        res = push_managed_tabs(s, sheets, [_row(o)])
        s.commit()
        assert sheets.inserts == 1                        # no se vuelve a insertar
        assert res["migracion_courier"] is None
        assert sheets.tabs[TAB] == primera                # misma hoja
        assert sheets.tabs[TAB][0] == SEGUIMIENTO_COLUMNS_V2
        # La BD tampoco se convierte dos veces.
        assert "formato_bd_convertidas" not in res["espejo"]
        assert s.scalars(select(SeguimientoManual)).one().values_json == manual_bd
        assert res["espejo"]["ediciones_leidas"] == 0


def test_vista_previa_no_toca_la_hoja_ni_la_bd_y_avisa_de_la_migracion(factory):
    with factory() as s:
        o, sheets, mundo = _mundo_viejo(s)
        res = push_managed_tabs(s, sheets, [_row(o)], dry_run=True)
        assert sheets.inserts == 0 and sheets.written == {}
        assert sheets.tabs[TAB] == mundo["hoja"]           # ni una celda
        assert res["migracion_courier"]["estado"] == "pendiente"
        assert res["migracion_courier"]["celdas_antes"] == sum(
            recuento_por_columna(mundo["hoja"]))
        # Se lee como si ya estuviera migrada: la manual y el histórico, enteros.
        assert res["manuales"] == 1
        assert res["historico_preservado"] == 3
        assert res["espejo"]["ediciones_leidas"] == 0
        # La BD, sin convertir (ni marca).
        assert s.get(SeguimientoSyncMeta, META_FORMATO_BD) is None
        assert len(json.loads(s.scalars(select(SeguimientoManual)).one().values_json)) == 19


def test_si_el_recuento_no_cuadra_se_para_sin_escribir(factory):
    """Si alguien escribe mientras se inserta la columna (el recuento de después
    no cuadra con el de antes), la pasada se para: no se escribe nada más."""

    class HojaQueCambia(FakeTabs):
        def format_tab(self, title: str, requests: list[dict[str, Any]]) -> None:
            super().format_tab(title, requests)
            if any("insertDimension" in r for r in requests):
                self.tabs[title][-1][1] = ""          # una celda que desaparece

    with factory() as s:
        o, _sheets, mundo = _mundo_viejo(s)
        sheets = HojaQueCambia({HISTORICA: [], TAB: copy.deepcopy(mundo["hoja"])})
        with pytest.raises(DriveSyncError, match="recuento"):
            push_managed_tabs(s, sheets, [_row(o)])
        s.rollback()
        assert sheets.written == {}
        # La BD sigue sin convertir: la siguiente pasada lo repite todo junto.
        assert s.get(SeguimientoSyncMeta, META_FORMATO_BD) is None


def test_bd_vieja_con_la_hoja_ya_migrada_se_convierte_igual(factory):
    """Si una pasada insertó la columna pero no llegó a confirmar, la siguiente
    encuentra la hoja ya en 20 columnas y la BD aún en 19: convierte la BD (por
    su marca) y no vuelve a tocar la hoja."""

    class HojaQueFalla(FakeTabs):
        def replace_tab(self, title: str, rows: list[list[Any]], *, raw: bool = False) -> None:
            raise DriveSyncError("Google no responde")

    with factory() as s:
        o, _sheets, mundo = _mundo_viejo(s)
        sheets = HojaQueFalla({HISTORICA: [], TAB: copy.deepcopy(mundo["hoja"])})
        with pytest.raises(DriveSyncError):
            push_managed_tabs(s, sheets, [_row(o)])
        s.rollback()
        assert sheets.inserts == 1                     # la hoja SÍ se migró
        assert s.get(SeguimientoSyncMeta, META_FORMATO_BD) is None
        # Siguiente pasada, con Google ya respondiendo.
        sana = FakeTabs({HISTORICA: [], TAB: copy.deepcopy(sheets.tabs[TAB])})
        res = push_managed_tabs(s, sana, [_row(o)])
        s.commit()
        assert sana.inserts == 0 and res["migracion_courier"] is None
        assert s.get(SeguimientoSyncMeta, META_FORMATO_BD).valor == "20"
        assert res["espejo"]["ediciones_leidas"] == 0
        hist = _historico_escrito(sana)
        assert [h[ENVIO] for h in hist] == ["UPS", "FEDEX", "DSV"]
        assert [h[ID_INDEX] for h in hist] == ["leg-1", "leg-2", "leg-3"]


def test_una_hoja_nueva_nace_con_veinte_columnas(factory):
    with factory() as s:
        o = _order(s, "BOP-200")
        sheets = FakeTabs({HISTORICA: []})
        res = push_managed_tabs(s, sheets, [_row(o)])
        s.commit()
        assert res["migracion_courier"] is None and sheets.inserts == 0
        assert sheets.written[TAB][0] == SEGUIMIENTO_COLUMNS_V2
        assert sheets.written[TAB][1][COURIER] == "UPS"
        assert s.get(SeguimientoSyncMeta, META_FORMATO_BD).valor == "20"


# --- backfill: informe y agencias de Genei que faltan -------------------------------


class _GeneiFalso:
    """Cliente de Genei de mentira (sin red): solo `get_shipment`."""

    def __init__(self, agencias: dict[str, Any]) -> None:
        self.agencias = agencias
        self.pedidos: list[str] = []

    def get_shipment(self, code: str) -> dict[str, Any]:
        from app.erp.integrations.genei.client import GeneiError

        self.pedidos.append(code)
        valor = self.agencias.get(code)
        if isinstance(valor, Exception):
            raise valor
        if valor is None:
            raise GeneiError("no existe")
        return {"codigo_envio": code, "nombre_agencia": valor}


def test_backfill_informa_y_rellena_solo_las_agencias_que_faltan(factory):
    from app.erp.seguimiento_courier_backfill import informe, rellenar_agencias_genei

    with factory() as s:
        con = _order(s, "G-CON", transport_status=TransportStatus.DELIVERED)
        _genei(con, courier="UPS Standard")
        sin = _order(s, "G-SIN", transport_status=TransportStatus.DELIVERED)
        _genei(sin, courier=None)
        vacia = _order(s, "G-VACIA", transport_status=TransportStatus.IN_TRANSIT)
        _genei(vacia, courier=None)
        ext = _order(s, "E-UPS", transport_status=TransportStatus.IN_TRANSIT)
        set_external_state(ext, {"courier": "UPS"})
        _order(s, "E-SIN", transport_status=TransportStatus.IN_TRANSIT)
        _order(s, "NADA", transport_status=TransportStatus.NOT_SHIPPED)
        _order(s, "NO-ENV", shipping_not_required=True)
        s.commit()

        rep = informe(s)
        assert rep["pedidos"] == 7
        assert rep["por_tipo"] == {
            "genei_con_agencia": 1, "genei_sin_agencia": 2, "otro_courier_apuntado": 1,
            "otro_courier_sin_apuntar": 1, "sin_envio_aun": 1, "sin_envio": 1,
        }
        assert set(rep["por_envio"]) <= set(envio_vocabulary())
        assert rep["genei_sin_agencia"] == ["G-SIN", "G-VACIA"]

        # Sin --apply: no se pregunta a Genei ni se escribe nada.
        genei = _GeneiFalso({"G-G-SIN": "Ctt Premium", "G-G-VACIA": ""})
        res = rellenar_agencias_genei(s, genei)
        assert res["pendientes"] == 2 and res["rellenados"] == 0 and genei.pedidos == []
        assert courier_label(s.get(Order, sin.id)) == "Genei"

        # Con --apply: solo los que faltan; la que Genei tampoco sabe, se queda.
        pausas: list[int] = []
        res = rellenar_agencias_genei(s, genei, apply=True, pausa=lambda: pausas.append(1))
        assert res["rellenados"] == 1 and res["sin_agencia_en_genei"] == 1
        assert genei.pedidos == ["G-G-SIN", "G-G-VACIA"] and pausas == [1]
        assert courier_label(s.get(Order, sin.id)) == "Ctt Premium"
        assert courier_label(s.get(Order, con.id)) == "UPS Standard"     # no se tocó
        # Otra pasada: ya solo queda la que Genei no sabe.
        assert informe(s)["genei_sin_agencia"] == ["G-VACIA"]


def test_backfill_corta_si_genei_rechaza_las_credenciales(factory):
    from app.erp.integrations.genei.client import GeneiAuthError
    from app.erp.seguimiento_courier_backfill import rellenar_agencias_genei

    with factory() as s:
        for n in ("G-1", "G-2"):
            _genei(_order(s, n, transport_status=TransportStatus.DELIVERED), courier=None)
        s.commit()
        genei = _GeneiFalso({"G-G-1": GeneiAuthError("401"), "G-G-2": "MRW"})
        res = rellenar_agencias_genei(s, genei, apply=True)
        assert res["cortado"] == "credenciales rechazadas"
        assert res["rellenados"] == 0 and genei.pedidos == ["G-G-1"]


def test_un_cambio_de_bohub_entre_pasadas_no_se_deshace_al_migrar(factory):
    """El tracking de un pedido cambia en BoHub (1Z-VIEJO → 1Z-NUEVO) entre la
    última pasada de la versión anterior y la primera de esta (la que migra).
    La hoja y la foto dicen 1Z-VIEJO: eso no es una edición a mano, así que la
    hoja se pone al día con 1Z-NUEVO (no se vuelve a 1Z-VIEJO)."""
    with factory() as s:
        o, sheets, _mundo = _mundo_viejo(s)
        o.tracking_number = "1Z-NUEVO"                 # lo cambió BoHub (no la hoja)
        s.commit()
        res = push_managed_tabs(s, sheets, [_row(o)])
        s.commit()
        s.refresh(o)
        assert o.tracking_number == "1Z-NUEVO"
        assert res["espejo"]["tracking_leidos"] == 0
        fila = next(f for f in sheets.written[TAB] if len(f) > ID_INDEX and f[ID_INDEX] == o.id)
        assert fila[_c("Tracking")] == "1Z-NUEVO"


def test_sin_cabecera_el_formato_viejo_se_reconoce_por_donde_estan_las_ids(factory):
    """Si alguien borró la fila de cabecera, la pestaña vieja se reconoce por
    dónde están sus «id» (columna S): se migra igual, sin escribir «Courier»
    en una fila de datos, y la pasada vuelve a poner la cabecera."""
    with factory() as s:
        o, _sheets, mundo = _mundo_viejo(s)
        sin_cabecera = copy.deepcopy(mundo["hoja"][1:])
        sheets = FakeTabs({HISTORICA: [], TAB: sin_cabecera})
        res = push_managed_tabs(s, sheets, [_row(o)])
        s.commit()
        assert sheets.inserts == 1 and res["migracion_courier"]["estado"] == "hecha"
        lotes = [r for lote in sheets.formats[TAB] for r in lote if "updateCells" in r]
        assert lotes == []                             # nada escrito en una fila de datos
        escrito = sheets.written[TAB]
        assert escrito[0] == SEGUIMIENTO_COLUMNS_V2
        hist = _historico_escrito(sheets)
        assert [h[ENVIO] for h in hist] == ["UPS", "FEDEX", "DSV"]
        assert [h[ID_INDEX] for h in hist] == ["leg-1", "leg-2", "leg-3"]
        man = next(f for f in escrito if len(f) > ID_INDEX and f[ID_INDEX] == "man-1")
        assert man[_c("Tracking")] == "T-1" and man[COURIER] == ""
        assert res["espejo"]["ediciones_leidas"] == 0


# --- lo que la migración NO debe hacer (revisión) -------------------------------------


def test_courier_tecleado_en_otra_columna_no_hace_pasar_la_pestana_vieja_por_nueva(factory):
    """Alguien escribe «Courier» en T1 de la pestaña vieja (19 columnas). El
    formato se decide por lo que hay detrás de «Envío», no por si «Courier»
    aparece en alguna celda: la pestaña se migra igual y nada se descoloca."""
    with factory() as s:
        o, _sheets, mundo = _mundo_viejo(s)
        hoja = copy.deepcopy(mundo["hoja"])
        hoja[0] = [*hoja[0], "Courier"]                       # T1, a mano
        sheets = FakeTabs({HISTORICA: [], TAB: hoja})
        res = push_managed_tabs(s, sheets, [_row(o)])
        s.commit()
        assert sheets.inserts == 1 and res["migracion_courier"]["estado"] == "hecha"
        hist = _historico_escrito(sheets)
        assert [h[ENVIO] for h in hist] == ["UPS", "FEDEX", "DSV"]
        assert [h[_c("Tracking")] for h in hist] == ["TRK-1", "TRK-2", ""]
        assert [h[_c("Nota / Incidencia")] for h in hist] == [
            "nota vieja", "", "llamar antes"]
        assert [h[ID_INDEX] for h in hist] == ["leg-1", "leg-2", "leg-3"]
        man = next(f for f in sheets.written[TAB]
                   if len(f) > ID_INDEX and f[ID_INDEX] == "man-1")
        assert man[_c("Tracking")] == "T-1" and man[_c("Nota / Incidencia")] == "hola"
        assert res["espejo"]["borradas"] == 0 and res["espejo"]["ediciones_leidas"] == 0


def test_con_courier_insertada_dos_veces_no_se_escribe_nada(factory):
    """Si la columna se insertó dos veces (dos pasadas a la vez sin cerrojo, o a
    mano), las «id» quedan en la U: leerla como de 20 columnas descolocaría
    todo. Se para sin tocar la hoja ni la BD, y lo dice."""
    with factory() as s:
        o, _sheets, mundo = _mundo_viejo(s)
        hoja = copy.deepcopy(mundo["hoja"])
        for fila in hoja:
            if len(fila) > COURIER:
                fila[COURIER:COURIER] = ["", ""]
        hoja[0][COURIER:COURIER + 2] = ["Courier", "Courier"]
        sheets = FakeTabs({HISTORICA: [], TAB: copy.deepcopy(hoja)})
        for dry_run in (True, False):
            with pytest.raises(DriveSyncError, match="descolocadas"):
                push_managed_tabs(s, sheets, [_row(o)], dry_run=dry_run)
            s.rollback()
        assert sheets.inserts == 0 and sheets.written == {} and sheets.tabs[TAB] == hoja
        assert s.get(SeguimientoSyncMeta, META_FORMATO_BD) is None


def test_sin_la_bd_migrada_no_se_toca_la_hoja(factory):
    """Una pasada nueva contra una BD sin la migración 0123 (sin la tabla de
    estado del espejo) se para ANTES de insertar la columna en la hoja: si no,
    una versión vieja aún en marcha leería la hoja nueva descolocada."""
    with factory() as s:
        o, sheets, mundo = _mundo_viejo(s)
        SeguimientoSyncMeta.__table__.drop(s.get_bind())
        with pytest.raises(DriveSyncError, match="base de datos"):
            push_managed_tabs(s, sheets, [_row(o)])
        s.rollback()
        assert sheets.inserts == 0 and sheets.written == {}
        assert sheets.tabs[TAB] == mundo["hoja"]


def test_con_algo_en_la_columna_z_no_se_inserta_la_columna(factory):
    """Lo de la columna Z pasaría a AA al insertar «Courier»: fuera de lo que
    lee la app, y se quedaría suelto. La vista previa lo avisa y la pasada se
    para sin tocar nada."""
    with factory() as s:
        o, _sheets, mundo = _mundo_viejo(s)
        hoja = copy.deepcopy(mundo["hoja"])
        hoja[-1] = [*hoja[-1], *[""] * (25 - len(hoja[-1])), "apunte en Z"]
        sheets = FakeTabs({HISTORICA: [], TAB: copy.deepcopy(hoja)})
        previa = push_managed_tabs(s, sheets, [_row(o)], dry_run=True)
        assert previa["migracion_courier"]["celdas_que_no_caben"] == 1
        with pytest.raises(DriveSyncError, match="columna Z"):
            push_managed_tabs(s, sheets, [_row(o)])
        s.rollback()
        assert sheets.inserts == 0 and sheets.written == {} and sheets.tabs[TAB] == hoja
        assert s.get(SeguimientoSyncMeta, META_FORMATO_BD) is None


def test_si_la_hoja_cambia_antes_de_insertar_no_se_inserta(factory):
    """Entre la lectura y la inserción alguien escribe (u otra pasada acaba de
    migrar la pestaña): se relee justo antes y, si no es lo leído, se para sin
    insertar nada — nunca dos columnas «Courier»."""

    class HojaQueCambiaAlReleer(FakeTabs):
        lecturas = 0

        def tab_values(self, title: str, *, raw: bool = False) -> list[list[Any]]:
            self.lecturas += 1
            if title == TAB and self.lecturas == 2:
                self.tabs[TAB][2][CABECERA_19.index("Nota / Incidencia")] = "recién escrito"
            return super().tab_values(title, raw=raw)

    with factory() as s:
        o, _sheets, mundo = _mundo_viejo(s)
        sheets = HojaQueCambiaAlReleer({HISTORICA: [], TAB: copy.deepcopy(mundo["hoja"])})
        with pytest.raises(DriveSyncError, match="ha cambiado"):
            push_managed_tabs(s, sheets, [_row(o)])
        s.rollback()
        assert sheets.inserts == 0 and sheets.written == {}


def test_una_pestana_de_17_columnas_se_reescribe_y_no_queda_pendiente(factory):
    """Una pestaña del formato de 17 columnas (sin «Fecha recogido» ni «id») se
    pone al día al reescribirla: el resumen dice «reescrita», no «pendiente»."""
    from app.erp.drive_managed import _HEADER_V2_SIN_RECOGIDO

    with factory() as s:
        o = _order(s, "BOP-300")
        fila = [""] * 17
        fila[1], fila[3], fila[4] = "M-17", "Manual SL", "MANUAL"
        fila[13], fila[14], fila[16] = "MRW", "TRK-17", "nota 17"
        sheets = FakeTabs({HISTORICA: [], TAB: [list(_HEADER_V2_SIN_RECOGIDO), fila]})
        previa = push_managed_tabs(s, sheets, [_row(o)], dry_run=True)
        assert previa["migracion_courier"]["estado"] == "pendiente"
        assert previa["migracion_courier"]["formato"] == "sin_recogido"
        res = push_managed_tabs(s, sheets, [_row(o)])
        s.commit()
        assert res["migracion_courier"]["estado"] == "reescrita"
        assert sheets.inserts == 0
        assert sheets.written[TAB][0] == SEGUIMIENTO_COLUMNS_V2
        man = next(f for f in sheets.written[TAB] if len(f) > 1 and f[1] == "M-17")
        assert man[ENVIO] == "MRW" and man[COURIER] == ""
        assert man[_c("Tracking")] == "TRK-17" and man[_c("Nota / Incidencia")] == "nota 17"


def test_el_backfill_de_ids_antes_de_la_primera_pasada_no_corre_dos_veces_lo_guardado(factory):
    """`scripts.backfill_seguimiento_ids --apply` entre el despliegue y la
    primera pasada guarda el histórico ya en 20 columnas: antes pone la BD al
    día (con su marca, en la misma transacción), así que la pasada no vuelve a
    correr nada — tampoco lo que ya no está en la hoja."""
    from app.erp.drive_managed import historico_manual_rows, realinear_pestana
    from app.erp.seguimiento_backfill import import_legacy_rows

    with factory() as s:
        o, sheets, _mundo = _mundo_viejo(s)
        fuera = _hist("V-X", "MRW", "", "TRK-X", "borrada", "leg-x")
        s.add(SeguimientoLegacy(
            id="leg-x", row_index=50, numero_raw="V-X", cliente_raw="Cliente viejo",
            raw_json=json.dumps(fuera[:ID_INDEX - 1]), match_status="synthetic",
            deleted_at=datetime.now(UTC),
        ))
        s.commit()
        leidos = sheets.tab_values(TAB, raw=True)
        import_legacy_rows(s, historico_manual_rows(realinear_pestana(leidos)))
        s.commit()
        assert s.get(SeguimientoSyncMeta, META_FORMATO_BD).valor == "20"
        x = json.loads(s.get(SeguimientoLegacy, "leg-x").raw_json)
        assert x[ENVIO] == "MRW" and x[COURIER] == "" and x[_c("Tracking")] == "TRK-X"

        res = push_managed_tabs(s, sheets, [_row(o)])
        s.commit()
        assert "formato_bd_convertidas" not in res["espejo"]       # ya estaba al día
        assert json.loads(s.get(SeguimientoLegacy, "leg-x").raw_json) == x
        for rec in s.scalars(select(SeguimientoLegacy).where(SeguimientoLegacy.id != "leg-x")):
            raw = json.loads(rec.raw_json)
            assert raw[ENVIO] in ("UPS", "FEDEX", "DSV") and raw[COURIER] == ""
        man = json.loads(s.scalars(select(SeguimientoManual)).one().values_json)
        assert man[_c("Tracking")] == "T-1" and man[ID_INDEX] == "man-1"


def test_poner_la_bd_al_dia_no_corre_las_filas_que_ya_lo_estan(factory):
    """La foto y las filas manuales llevan su id: solo se convierten si la
    tienen en la posición vieja. Una fila ya de 20 columnas no se corre aunque
    falte la marca."""
    from app.erp.seguimiento_mirror import poner_bd_al_dia

    with factory() as s:
        nueva = [""] * len(SEGUIMIENTO_COLUMNS_V2)
        nueva[ENVIO], nueva[_c("Tracking")], nueva[ID_INDEX] = "MRW", "T-20", "man-20"
        vieja = _a_19([*nueva[:ID_INDEX], "man-19"])
        s.add_all([
            SeguimientoManual(row_id="man-20", values_json=json.dumps(nueva)),
            SeguimientoManual(row_id="man-19", values_json=json.dumps(vieja)),
            SeguimientoSnapshot(row_id="man-20", kind=KIND_MANUAL, values_json=json.dumps(nueva)),
        ])
        s.flush()
        assert poner_bd_al_dia(s) == 1
        por_id = {m.row_id: json.loads(m.values_json)
                  for m in s.scalars(select(SeguimientoManual))}
        assert por_id["man-20"] == nueva
        assert por_id["man-19"][_c("Tracking")] == "T-20"
        assert por_id["man-19"][ID_INDEX] == "man-19" and por_id["man-19"][COURIER] == ""
        assert json.loads(s.get(SeguimientoSnapshot, "man-20").values_json) == nueva
        assert poner_bd_al_dia(s) is None                   # con su marca: nada


# --- columnas insertadas / borradas / renombradas a mano (segunda revisión) -----------


def _sin_columna(filas: list[list[Any]], col: int) -> list[list[Any]]:
    """La pestaña con la columna `col` borrada (lo de detrás corre a la izquierda)."""
    return [[c for i, c in enumerate(f) if i != col] for f in filas]


def _con_columna(filas: list[list[Any]], col: int, valor: str = "") -> list[list[Any]]:
    """La pestaña con una columna insertada en `col` (lo de detrás corre a la derecha)."""
    return [[*f[:col], valor, *f[col:]] if len(f) > col else list(f) for f in filas]


def test_una_columna_borrada_a_mano_en_la_pestana_migrada_no_se_lee_corrida(factory):
    """Con la hoja ya en 20 columnas, alguien borra «Cliente» (D): la cabecera
    ya no es la de la app (todo lo de detrás ha corrido) y las «id» han pasado
    a la S. No se toma por la pestaña vieja (no se inserta otra «Courier») ni
    se lee corrida: no se escribe nada y se dice qué columna falla."""
    with factory() as s:
        o, sheets, _mundo = _mundo_viejo(s)
        push_managed_tabs(s, sheets, [_row(o)])
        s.commit()
        hoja = _sin_columna(sheets.tabs[TAB], _c("Cliente"))
        tocada = FakeTabs({HISTORICA: [], TAB: copy.deepcopy(hoja)})
        for dry_run in (True, False):
            with pytest.raises(DriveSyncError, match="descolocadas.*columna D"):
                push_managed_tabs(s, tocada, [_row(o)], dry_run=dry_run)
            s.rollback()
        assert tocada.inserts == 0 and tocada.written == {} and tocada.tabs[TAB] == hoja
        assert s.scalars(select(SeguimientoOverride)).all() == []


def test_una_columna_insertada_a_mano_en_la_pestana_vieja_no_se_lee_corrida(factory):
    """En la pestaña vieja (19 columnas) alguien inserta una columna detrás de
    «Origen»: las «id» quedan en la T, como en el formato nuevo, pero la
    cabecera delata que todo lo de detrás ha corrido. No se escribe nada."""
    with factory() as s:
        o, _sheets, mundo = _mundo_viejo(s)
        hoja = _con_columna(mundo["hoja"], _c("Origen") + 1)
        hoja[0][_c("Origen") + 1] = "Comercial"
        sheets = FakeTabs({HISTORICA: [], TAB: copy.deepcopy(hoja)})
        for dry_run in (True, False):
            with pytest.raises(DriveSyncError, match="descolocadas"):
                push_managed_tabs(s, sheets, [_row(o)], dry_run=dry_run)
            s.rollback()
        assert sheets.inserts == 0 and sheets.written == {} and sheets.tabs[TAB] == hoja
        assert s.get(SeguimientoSyncMeta, META_FORMATO_BD) is None


def test_una_celda_de_la_cabecera_renombrada_no_para_la_pasada(factory):
    """Renombrar una celda de la cabecera no descoloca nada: se tolera (una), la
    migración se hace igual y la pasada vuelve a escribir el nombre bueno."""
    with factory() as s:
        o, _sheets, mundo = _mundo_viejo(s)
        hoja = copy.deepcopy(mundo["hoja"])
        hoja[0][CABECERA_19.index("Cliente")] = "Cliente final"
        sheets = FakeTabs({HISTORICA: [], TAB: hoja})
        res = push_managed_tabs(s, sheets, [_row(o)])
        s.commit()
        assert sheets.inserts == 1 and res["migracion_courier"]["estado"] == "hecha"
        assert sheets.written[TAB][0] == SEGUIMIENTO_COLUMNS_V2
        hist = _historico_escrito(sheets)
        assert [h[ENVIO] for h in hist] == ["UPS", "FEDEX", "DSV"]
        assert [h[_c("Tracking")] for h in hist] == ["TRK-1", "TRK-2", ""]


def test_sin_cabecera_y_con_las_ids_fuera_de_sitio_no_se_escribe(factory):
    """Sin fila de cabecera, el formato sale de dónde están las «id»; si están
    en una columna que no es de ningún formato (aquí, a la izquierda de O, tras
    borrar varias columnas), no se sabe leer: no se escribe nada."""
    with factory() as s:
        o, sheets, _mundo = _mundo_viejo(s)
        push_managed_tabs(s, sheets, [_row(o)])
        s.commit()
        hoja = [list(f) for f in sheets.tabs[TAB][1:]]
        for _ in range(6):
            hoja = _sin_columna(hoja, 1)
        tocada = FakeTabs({HISTORICA: [], TAB: copy.deepcopy(hoja)})
        with pytest.raises(DriveSyncError):
            push_managed_tabs(s, tocada, [_row(o)])
        s.rollback()
        assert tocada.inserts == 0 and tocada.written == {}


def test_pestana_de_17_columnas_con_algo_en_y_no_se_reescribe(factory):
    """Al pasar de 17 a 20 columnas se abren DOS huecos: lo de Y y Z se saldría
    del rango de la app. La vista previa lo cuenta y la pasada no escribe."""
    from app.erp.drive_managed import _HEADER_V2_SIN_RECOGIDO

    with factory() as s:
        o = _order(s, "BOP-301")
        fila = [""] * 17
        fila[1], fila[3], fila[4] = "M-17", "Manual SL", "MANUAL"
        fila = [*fila, *[""] * 7, "apunte en Y"]                # Y = índice 24
        hoja = [list(_HEADER_V2_SIN_RECOGIDO), fila]
        sheets = FakeTabs({HISTORICA: [], TAB: copy.deepcopy(hoja)})
        previa = push_managed_tabs(s, sheets, [_row(o)], dry_run=True)
        assert previa["migracion_courier"]["celdas_que_no_caben"] == 1
        with pytest.raises(DriveSyncError, match="columnas Y y Z"):
            push_managed_tabs(s, sheets, [_row(o)])
        s.rollback()
        assert sheets.written == {} and sheets.tabs[TAB] == hoja
