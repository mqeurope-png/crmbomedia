"""ERP · WooCommerce — qué pedidos web entran en el seguimiento, según su estado.

La puerta es haber PASADO POR CAJA: `processing` / `completed` / `refunded`
entran; `pending`, `on-hold`, `cancelled`, `failed`, `draft` y `trash` quedan
ocultos por estado. `refunded` es el estado PROPIO «Reembolsado» y se ve
AUNQUE el pedido esté anulado en BoHub.

Cubre: la regla de visibilidad por estado (Parte A/D), el webhook de cambio de
estado (Parte B), la reconciliación de los ya importados (Parte C), el efecto
en Drive (Parte E) y que el estado NO se confunde con la exclusión manual de
F6-fix7 (Parte D).
"""
from __future__ import annotations

import json
from collections.abc import Generator
from datetime import UTC, datetime
from unittest.mock import MagicMock, patch

import pytest
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

import app.integrations.woocommerce.jobs as woo_jobs
import app.main  # noqa: F401
from app.core.crypto import encrypt
from app.db.base import Base
from app.db.session import get_session
from app.erp import seguimiento as core
from app.erp.api.seguimiento import _rows, build_drive_sync_rows
from app.erp.drive_sheets import sync_to_sheet
from app.erp.models import (
    IntegrationEvent,
    InvoiceStatus,
    Order,
    OrderSource,
    TransportStatus,
)
from app.integrations.woocommerce.mapper import import_woo_order
from app.main import app
from app.models.crm import Company, ExternalSystem
from app.models.integration_settings import (
    IntegrationAccount,
    IntegrationMode,
    IntegrationStatus,
)
from tests._test_helpers import auth_headers, seed_test_users

HEADER = list(core.SEGUIMIENTO_COLUMNS)
_IDX = {n: i for i, n in enumerate(core.SEGUIMIENTO_COLUMNS)}
KEY = _IDX["Albarán / Nº Pedido Web"]


@pytest.fixture()
def session_factory() -> Generator[sessionmaker, None, None]:
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False}, poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def _fk_on(dbapi_conn, _record):  # noqa: ANN001
        dbapi_conn.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    with factory() as seed:
        seed_test_users(seed)
        seed.commit()
    yield factory
    Base.metadata.drop_all(engine)


@pytest.fixture()
def http(session_factory) -> Generator:
    from fastapi.testclient import TestClient

    def override():
        with session_factory() as s:
            yield s
    app.dependency_overrides[get_session] = override
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _store(s: Session, slug: str = "boprint") -> IntegrationAccount:
    a = IntegrationAccount(
        system=ExternalSystem.WOOCOMMERCE, account_id=slug, display_name=slug,
        enabled=True, mode=IntegrationMode.LIVE, status=IntegrationStatus.CONFIGURED,
        base_url=f"https://{slug}.example",
        consumer_key_encrypted=encrypt("ck"), consumer_secret_encrypted=encrypt("cs"),
        credential_status="configured",
    )
    s.add(a)
    s.flush()
    return a


def _order(
    s: Session, *, woo_id: str, number: str, woo_status: str | None,
    store: IntegrationAccount, cliente: str = "Cliente SL",
    tracking: str | None = None, delivered: bool = False, invoiced: bool = False,
) -> Order:
    comp = Company(name=cliente)
    s.add(comp)
    s.flush()
    o = Order(
        external_source=OrderSource.WOOCOMMERCE, external_id=woo_id,
        store_id=store.id, order_number=number, company_id=comp.id,
        woo_status=woo_status, tracking_number=tracking,
        placed_at=datetime(2026, 9, 1, tzinfo=UTC),
    )
    if delivered:
        o.transport_status = TransportStatus.DELIVERED
    if invoiced:
        o.invoice_status = InvoiceStatus.GENERATED
        o.factusol_invoice_number = f"5-{woo_id}"
    s.add(o)
    s.flush()
    return o


def _rows_for(s: Session, **kw) -> list[dict]:
    return core.filter_rows(_rows(s), **kw)


def _numbers(rows: list[dict]) -> set[str]:
    return {r["order_number"] for r in rows}


# --- Parte A/D: la regla de visibilidad --------------------------------------------


def test_cancelled_order_removed_from_seguimiento(session_factory) -> None:
    with session_factory() as s:
        st = _store(s)
        _order(s, woo_id="1", number="BOPRIN-1", woo_status="cancelled", store=st)
        s.commit()
        default = _rows_for(s, en_curso=True)
        assert "BOPRIN-1" not in _numbers(default)
        # Aparece en la vista de «ocultados por estado».
        ocultos = _rows_for(s, ver_ocultos_estado=True)
        assert _numbers(ocultos) == {"BOPRIN-1"}
        assert ocultos[0]["oculto_por_estado"] is True
        assert ocultos[0]["estado_woo_motivo"] == "cancelled"


def test_failed_order_removed_from_seguimiento(session_factory) -> None:
    with session_factory() as s:
        st = _store(s)
        _order(s, woo_id="2", number="BOPRIN-2", woo_status="failed", store=st)
        s.commit()
        assert "BOPRIN-2" not in _numbers(_rows_for(s, en_curso=True))
        assert _numbers(_rows_for(s, ver_ocultos_estado=True)) == {"BOPRIN-2"}


@pytest.mark.parametrize("woo_status", ["processing", "completed"])
def test_un_web_que_paso_por_caja_esta_en_seguimiento(session_factory, woo_status) -> None:
    """La puerta web: `processing` (pagado) y `completed` (servido) entran."""
    with session_factory() as s:
        st = _store(s)
        _order(s, woo_id="3", number="BOPRIN-3", woo_status=woo_status, store=st)
        s.commit()
        rows = _rows_for(s, en_curso=True)
        assert "BOPRIN-3" in _numbers(rows)
        row = next(r for r in rows if r["order_number"] == "BOPRIN-3")
        assert row["oculto_por_estado"] is False
        assert row["pendiente_escribir"] is True


@pytest.mark.parametrize(
    ("woo_status", "motivo"),
    [
        ("pending", "pending"),
        ("on-hold", "on_hold"),
        ("wc-on-hold", "on_hold"),        # export/BD de la tienda, con prefijo
        ("failed", "failed"),
        ("draft", "draft"),
        ("checkout-draft", "checkout_draft"),
    ],
)
def test_un_web_que_no_paso_por_caja_no_sale_en_seguimiento(
    session_factory, woo_status, motivo,
) -> None:
    """El caso reportado: BOPRIN-99931/99922 (`on-hold`) y BOPRIN-99880/99878/
    99886 (`pending`) salían en Seguimiento sin haber entrado al flujo real.
    Quedan ocultos POR ESTADO, con su motivo y a un clic."""
    with session_factory() as s:
        st = _store(s)
        _order(s, woo_id="3", number="BOPRIN-3", woo_status=woo_status, store=st)
        s.commit()
        assert "BOPRIN-3" not in _numbers(_rows_for(s, en_curso=True))
        ocultos = _rows_for(s, ver_ocultos_estado=True)
        row = next(r for r in ocultos if r["order_number"] == "BOPRIN-3")
        assert row["oculto_por_estado"] is True
        assert row["estado_woo_motivo"] == motivo
        # Y con su etiqueta legible para «Ver ocultos por estado».
        assert row["estado_woo_motivo_label"]


@pytest.mark.parametrize("woo_status", [None, ""])
def test_un_web_sin_estado_se_queda(session_factory, woo_status) -> None:
    """Ocultar los NULL (#461) se llevó ~90 pedidos legítimos (BOPRIN-99915…,
    ARTISJ-9460…, FLUXLA-5742…): los web importados antes de que existiera el
    campo. Un web sin estado no se conoce → se queda; solo oculta un estado
    EXPLÍCITO de la tienda (o el `not_found` de la reconciliación)."""
    with session_factory() as s:
        st = _store(s)
        _order(s, woo_id="3b", number="BOPRIN-3B", woo_status=woo_status, store=st)
        s.commit()
        rows = _rows_for(s, en_curso=True)
        assert "BOPRIN-3B" in _numbers(rows)
        row = next(r for r in rows if r["order_number"] == "BOPRIN-3B")
        assert row["oculto_por_estado"] is False


def test_un_web_que_la_tienda_ya_no_tiene_queda_oculto(session_factory, http) -> None:
    """`not_found` lo pone «Poner al día estados Woo» al recibir el 404 de
    WooCommerce: fuera, con su motivo, y con «Reincluir» por si acaso."""
    with session_factory() as s:
        st = _store(s)
        o = _order(s, woo_id="3c", number="BOPRIN-3C", woo_status="not_found", store=st)
        s.commit()
        oid = o.id
        assert "BOPRIN-3C" not in _numbers(_rows_for(s, en_curso=True))
        row = next(r for r in _rows_for(s, ver_ocultos_estado=True)
                   if r["order_number"] == "BOPRIN-3C")
        assert row["estado_woo_motivo_label"] == "No encontrado en la tienda"
    r = http.post("/api/erp/seguimiento/force", json={"order_ids": [oid]},
                  headers=auth_headers(http, "pedidos"))
    assert r.status_code == 200, r.text
    with session_factory() as s:
        assert "BOPRIN-3C" in _numbers(_rows_for(s, en_curso=True))


def test_los_valores_reales_de_produccion_quedan_ocultos_en_los_tres_caminos(
    session_factory, http,
) -> None:
    """Los valores EXACTOS de la BD de producción (sin `wc-`, con guion):
    `on-hold` (BOPRIN-99922/99931) y `cancelled` (FLUXLA-5751/5786) tienen que
    desaparecer en el camino REAL de la pantalla, de «Descargar Excel» y de la
    zona viva de Drive — no solo en el helper de visibilidad."""
    import io  # noqa: PLC0415

    from openpyxl import load_workbook  # noqa: PLC0415

    from app.erp.api.seguimiento import drive_live_rows  # noqa: PLC0415

    with session_factory() as s:
        st = _store(s)
        _order(s, woo_id="99922", number="BOPRIN-99922", woo_status="on-hold", store=st)
        _order(s, woo_id="99931", number="BOPRIN-99931", woo_status="on-hold", store=st)
        _order(s, woo_id="5751", number="FLUXLA-5751", woo_status="cancelled", store=st)
        _order(s, woo_id="5786", number="FLUXLA-5786", woo_status="cancelled", store=st)
        _order(s, woo_id="99878", number="BOPRIN-99878", woo_status=None, store=st)
        _order(s, woo_id="5790", number="FLUXLA-5790", woo_status="not_found", store=st)
        _order(s, woo_id="1", number="BOPRIN-1", woo_status="processing", store=st)
        s.commit()
        drive = {r["order_number"] for r in drive_live_rows(s)}
    h = auth_headers(http, "pedidos")
    pantalla = {i["order_number"]
                for i in http.get("/api/erp/seguimiento", headers=h).json()["items"]}
    wb = load_workbook(io.BytesIO(
        http.get("/api/erp/seguimiento/export", headers=h).content,
    ), read_only=True)
    filas = [list(row) for row in wb["Pedidos"].iter_rows(values_only=True)]
    excel = {f[filas[0].index("Nº pedido")] for f in filas[1:]}
    # El NULL (BOPRIN-99878) se QUEDA: no se conoce su estado.
    assert pantalla == excel == drive == {"BOPRIN-1", "BOPRIN-99878"}
    ocultos = http.get("/api/erp/seguimiento?ver_ocultos_estado=true", headers=h).json()
    motivos = {i["order_number"]: i["estado_woo_motivo_label"] for i in ocultos["items"]}
    assert motivos == {
        "BOPRIN-99922": "En espera", "BOPRIN-99931": "En espera",
        "FLUXLA-5751": "Cancelado en la tienda", "FLUXLA-5786": "Cancelado en la tienda",
        "FLUXLA-5790": "No encontrado en la tienda",
    }


def test_actualizar_hoja_de_drive_vuelca_solo_lo_que_ensena_la_pantalla(
    session_factory, http,
) -> None:
    """El bug de producción: «Actualizar hoja de Drive» pasaba TODAS las filas
    (sin filtrar) a la pestaña gestionada, y la hoja enseñaba los web
    `on-hold`/`cancelled` que la pantalla escondía. El endpoint real tiene que
    volcar exactamente la zona viva (`drive_live_rows`)."""
    _fake_sa = {
        "type": "service_account", "project_id": "bohub-drive",
        "client_email": "bohub@bohub-drive.iam.gserviceaccount.com",
        "private_key": "-----BEGIN PRIVATE KEY-----\nX\n-----END PRIVATE KEY-----\n",
    }
    with session_factory() as s:
        st = _store(s)
        _order(s, woo_id="99922", number="BOPRIN-99922", woo_status="on-hold", store=st)
        _order(s, woo_id="5751", number="FLUXLA-5751", woo_status="cancelled", store=st)
        _order(s, woo_id="99878", number="BOPRIN-99878", woo_status="not_found", store=st)
        _order(s, woo_id="2", number="BOPRIN-2", woo_status="completed", store=st,
               tracking="1Z")
        _order(s, woo_id="1", number="BOPRIN-1", woo_status="processing", store=st)
        s.commit()
    r = http.patch("/api/erp/settings", json={
        "drive_service_account_json": json.dumps(_fake_sa),
        "drive_spreadsheet_id": "1AbCdEfGhIjKlMnOpQrStUvWxYz",
    }, headers=auth_headers(http, "admin"))
    assert r.status_code == 200, r.text
    captured: dict = {}

    def fake_push(session, client, rows, *, completados=None, dry_run=False):
        captured["rows"] = rows
        captured["completados"] = completados or []
        return {"mode": "managed_tab", "dry_run": dry_run, "written": not dry_run,
                "rows": len(rows), "por_situacion": {}}

    with patch("app.erp.drive_managed.push_managed_tabs", side_effect=fake_push), \
         patch("app.erp.drive_sheets.GoogleSheetsClient", return_value=MagicMock()):
        r = http.post("/api/erp/seguimiento/drive-sync?dry_run=true",
                      headers=auth_headers(http, "pedidos"))
    assert r.status_code == 200, r.text
    numeros = [row["order_number"] for row in captured["rows"]]
    assert set(numeros) == {"BOPRIN-1", "BOPRIN-2"}
    # Y ya ordenadas por fecha del pedido (ambas del mismo día: orden estable).
    assert all(row["fecha"] == "2026-09-01" for row in captured["rows"])
    # El `completed` de la TIENDA (BOPRIN-2) NO es «Marcar completado»
    # (`completed_at`) de BoHub: sigue en la zona viva, no en los completados.
    assert captured["completados"] == []


@pytest.mark.parametrize(
    ("raw", "canonical"),
    [
        ("on-hold", "on_hold"), ("wc-on-hold", "on_hold"), ("on_hold", "on_hold"),
        ("ON-HOLD", "on_hold"), (" Wc-On-Hold ", "on_hold"),
        ("cancelled", "cancelled"), ("wc-cancelled", "cancelled"), ("CANCELLED", "cancelled"),
        ("pending", "pending"), ("wc-pending", "pending"),
        ("refunded", "refunded"), ("wc-refunded", "refunded"),
        ("processing", "processing"), ("wc-processing", "processing"),
        ("completed", "completed"), ("wc-completed", "completed"),
        ("checkout-draft", "checkout_draft"), (None, ""), ("", ""),
    ],
)
def test_normalizacion_unica_de_estados(raw, canonical) -> None:
    """Guion, guion bajo, prefijo `wc-` y mayúsculas son el mismo estado."""
    from app.erp.woo_status import normalize  # noqa: PLC0415

    assert normalize(raw) == canonical


def test_al_pasar_a_processing_vuelve_al_seguimiento(session_factory) -> None:
    """Un `on-hold` que paga aparece en el siguiente refresco; si se cancela,
    desaparece otra vez."""
    with session_factory() as s:
        st = _store(s)
        o = _order(s, woo_id="3c", number="BOPRIN-3C", woo_status="on-hold", store=st)
        s.commit()
        assert "BOPRIN-3C" not in _numbers(_rows_for(s, en_curso=True))
        o.woo_status = "processing"
        s.commit()
        assert "BOPRIN-3C" in _numbers(_rows_for(s, en_curso=True))
        o.woo_status = "cancelled"
        s.commit()
        assert "BOPRIN-3C" not in _numbers(_rows_for(s, en_curso=True))


def test_pedido_anulado_sale_del_seguimiento(session_factory) -> None:
    """El caso reportado: un pedido que se anula en BoHub seguía saliendo como
    «Pendiente» / «Por revisar». Seguimiento es la lista de pedidos VIVOS."""
    with session_factory() as s:
        st = _store(s)
        o = _order(s, woo_id="10", number="BOPRIN-10", woo_status="processing",
                   store=st)
        o.cancelled_at = datetime(2026, 9, 10, tzinfo=UTC)
        s.commit()
        assert "BOPRIN-10" not in _numbers(_rows_for(s, en_curso=True))
        ocultos = _rows_for(s, ver_ocultos_estado=True)
        row = next(r for r in ocultos if r["order_number"] == "BOPRIN-10")
        assert row["oculto_por_estado"] is True
        assert row["estado_woo_motivo"] == "anulado"


def test_la_anulacion_manda_sobre_el_estado_de_la_tienda(session_factory) -> None:
    """Anulado en BoHub sale aunque en la tienda siga «completed»."""
    with session_factory() as s:
        st = _store(s)
        o = _order(s, woo_id="11", number="BOPRIN-11", woo_status="completed",
                   store=st, invoiced=True)
        o.cancelled_at = datetime(2026, 9, 10, tzinfo=UTC)
        s.commit()
        assert "BOPRIN-11" not in _numbers(_rows_for(s, en_curso=True))


def test_un_manual_aprobado_sin_pagar_sigue_en_seguimiento(session_factory) -> None:
    """La puerta de estado es SOLO para los web: un pedido manual aprobado se
    ve aunque no haya pagado (no tiene `woo_status` que mirar)."""
    from app.erp.models import PreparationStatus  # noqa: PLC0415

    with session_factory() as s:
        comp = Company(name="Cliente Manual SL")
        s.add(comp)
        s.flush()
        s.add(Order(
            external_source=OrderSource.MANUAL, order_number="MAN-12",
            company_id=comp.id, approved_at=datetime(2026, 9, 2, tzinfo=UTC),
            preparation_status=PreparationStatus.IN_QUEUE,
            placed_at=datetime(2026, 9, 1, tzinfo=UTC),
        ))
        s.commit()
        rows = _rows_for(s, en_curso=True)
        assert "MAN-12" in _numbers(rows)
        row = next(r for r in rows if r["order_number"] == "MAN-12")
        assert row["oculto_por_estado"] is False


def test_refunded_se_queda_con_su_propia_situacion(session_factory) -> None:
    """`refunded` es el estado PROPIO «Reembolsado» de BoHub: se pagó y luego
    se devolvió el dinero. Se queda a la vista, con su Situación — ni «Listo»
    ni oculto."""
    with session_factory() as s:
        st = _store(s)
        _order(s, woo_id="4", number="BOPRIN-4", woo_status="refunded", store=st)
        s.commit()
        rows = _rows_for(s, en_curso=True)
        row = next(r for r in rows if r["order_number"] == "BOPRIN-4")
        assert row["oculto_por_estado"] is False
        assert row["reembolsado"] is True
        assert row["situacion"] == "reembolsado"
        assert row["situacion_label"] == "Reembolsado"


def test_refunded_se_ve_aunque_este_anulado_en_bohub(session_factory) -> None:
    """El caso de ARTISJ-9557 / ARTISJ-9494 / FLUXLA-5749: la auto-anulación
    por reembolso los sacaba de la vista. Para un `refunded`, el estado de la
    tienda manda sobre la anulación — es SOLO visibilidad, el pedido sigue
    anulado para sus acciones."""
    with session_factory() as s:
        st = _store(s)
        o = _order(s, woo_id="5", number="ARTISJ-9557", woo_status="refunded",
                   store=st, tracking="1Z-ENVIADO", invoiced=True, delivered=True)
        o.cancelled_at = datetime(2026, 9, 10, tzinfo=UTC)
        s.commit()
        rows = _rows_for(s, en_curso=True)
        row = next(r for r in rows if r["order_number"] == "ARTISJ-9557")
        assert row["oculto_por_estado"] is False
        assert row["reembolsado"] is True
        assert row["situacion"] == "reembolsado"
        # Sigue anulado en el pedido: esto no reactiva nada.
        assert o.cancelled_at is not None


def test_un_anulado_no_reembolsado_sigue_oculto(session_factory) -> None:
    """La excepción es SOLO del reembolso."""
    with session_factory() as s:
        st = _store(s)
        o = _order(s, woo_id="6", number="BOPRIN-6", woo_status="completed",
                   store=st, invoiced=True)
        o.cancelled_at = datetime(2026, 9, 10, tzinfo=UTC)
        s.commit()
        assert "BOPRIN-6" not in _numbers(_rows_for(s, en_curso=True))
        ocultos = _rows_for(s, ver_ocultos_estado=True)
        row = next(r for r in ocultos if r["order_number"] == "BOPRIN-6")
        assert row["estado_woo_motivo"] == "anulado"


def test_un_reembolsado_entregado_y_facturado_sigue_en_curso(session_factory) -> None:
    """Un reembolso deja trabajo por delante (el abono): se queda «en curso»
    aunque esté entregado y facturado, hasta que se marque completado."""
    with session_factory() as s:
        st = _store(s)
        o = _order(s, woo_id="7", number="FLUXLA-5749", woo_status="refunded",
                   store=st, invoiced=True, delivered=True)
        s.commit()
        assert "FLUXLA-5749" in _numbers(_rows_for(s, en_curso=True))
        o.completed_at = datetime(2026, 9, 12, tzinfo=UTC)
        s.commit()
        assert "FLUXLA-5749" not in _numbers(_rows_for(s, en_curso=True))


# --- Parte B: webhook de cambio de estado ------------------------------------------


def test_status_change_webhook_updates_order(session_factory) -> None:
    # El pedido entra como processing; llega order.updated con cancelled.
    with session_factory() as s:
        st = _store(s)
        payload = {"id": 700, "number": "700", "status": "processing",
                   "total": "10", "currency": "EUR",
                   "date_created": "2026-09-01T00:00:00Z",
                   "billing": {"email": "a@b.com"}, "line_items": [], "meta_data": [],
                   "_store_slug": st.account_id}
        import_woo_order(s, store=st, woo_order=payload)
        s.commit()
        ev = IntegrationEvent(
            system="woocommerce", account_id=st.account_id,
            external_event_id="wh-700-upd", event_type="order.updated",
            payload_json=json.dumps({"id": 700}),
        )
        s.add(ev)
        s.commit()
        event_id = ev.id

    fake = MagicMock()
    fake.get_order.return_value = {"id": 700, "number": "700", "status": "cancelled",
                                   "total": "10", "currency": "EUR",
                                   "billing": {"email": "a@b.com"}, "line_items": []}
    with patch.object(woo_jobs, "_session_factory", return_value=session_factory), \
         patch.object(woo_jobs, "WooHTTPClient", return_value=fake):
        woo_jobs.process_webhook_event(event_id)

    with session_factory() as s:
        o = s.scalar(select(Order).where(Order.external_id == "700"))
        assert o.woo_status == "cancelled"
        assert "BOPRIN-700" not in _numbers(_rows_for(s, en_curso=True))


def test_order_deleted_webhook_marks_trash(session_factory) -> None:
    with session_factory() as s:
        st = _store(s)
        _order(s, woo_id="800", number="BOPRIN-800", woo_status="processing", store=st)
        s.commit()
        ev = IntegrationEvent(
            system="woocommerce", account_id=st.account_id,
            external_event_id="wh-800-del", event_type="order.deleted",
            payload_json=json.dumps({"id": 800}),
        )
        s.add(ev)
        s.commit()
        event_id = ev.id
    with patch.object(woo_jobs, "_session_factory", return_value=session_factory), \
         patch.object(woo_jobs, "WooHTTPClient") as MockClient:
        res = woo_jobs.process_webhook_event(event_id)
    MockClient.assert_not_called()   # no re-consulta un pedido borrado
    assert res["matched"] is True
    with session_factory() as s:
        o = s.scalar(select(Order).where(Order.external_id == "800"))
        assert o.woo_status == "trash"
        assert "BOPRIN-800" not in _numbers(_rows_for(s, en_curso=True))


# --- Parte C: reconciliación — ver test_erp_reconcile_woo.py (async + listado).


# --- Parte E: Drive ----------------------------------------------------------------


class FakeSheet:
    def __init__(self, grid):
        self.grid = [list(r) for r in grid]

    def get_values(self):
        return [list(r) for r in self.grid]

    def update_cells(self, updates):
        for row, col, value in updates:
            r = self.grid[row - 1]
            while len(r) <= col:
                r.append("")
            r[col] = value

    def insert_rows_at(self, row, count):
        for _ in range(count):
            self.grid.insert(row - 1, [""] * len(HEADER))

    def write_rows(self, start_row, rows):  # pragma: no cover
        for i, values in enumerate(rows):
            self.grid[start_row - 1 + i] = list(values)


def test_cancelled_not_inserted_into_drive(session_factory) -> None:
    with session_factory() as s:
        st = _store(s)
        _order(s, woo_id="1200", number="BOPRIN-1200", woo_status="cancelled", store=st)
        _order(s, woo_id="1201", number="BOPRIN-1201", woo_status="processing", store=st)
        s.commit()
        sheet = FakeSheet([HEADER])
        summary = sync_to_sheet(s, sheet, build_drive_sync_rows(s))
    # Solo se inserta el activo; el cancelado no.
    assert summary["appended_rows"] == 1
    written = {r[KEY] for r in sheet.grid if KEY < len(r) and r[KEY]}
    assert "1201" in written
    assert "1200" not in written


def test_already_written_drive_row_is_not_touched_on_cancel(session_factory) -> None:
    with session_factory() as s:
        st = _store(s)
        o = _order(s, woo_id="1300", number="BOPRIN-1300", woo_status="processing", store=st)
        s.commit()
        sheet = FakeSheet([HEADER])
        sync_to_sheet(s, sheet, build_drive_sync_rows(s))   # se inserta
        before = [list(r) for r in sheet.grid]
        # Ahora se cancela.
        o = s.get(Order, o.id)
        o.woo_status = "cancelled"
        s.commit()
        # Nueva sincronización: only-insert + ya no es candidato → no toca nada.
        summary = sync_to_sheet(s, sheet, build_drive_sync_rows(s))
    assert summary["appended_rows"] == 0
    assert sheet.grid == before   # la fila ya escrita se queda intacta


# --- Parte D: separado de la exclusión manual --------------------------------------


def test_woo_status_is_separate_from_manual_exclusion(session_factory, http) -> None:
    with session_factory() as s:
        st = _store(s)
        cancelled = _order(s, woo_id="1400", number="BOPRIN-1400",
                           woo_status="cancelled", store=st)
        active = _order(s, woo_id="1401", number="BOPRIN-1401",
                        woo_status="processing", store=st)
        s.commit()
        cancelled_id, active_id = cancelled.id, active.id

    # «Reincluir» a mano el cancelado NO lo devuelve al seguimiento: sigue
    # oculto por estado (la exclusión manual y el estado son cosas distintas).
    r = http.post("/api/erp/seguimiento/include", json={"order_ids": [cancelled_id]},
                  headers=auth_headers(http, "pedidos"))
    assert r.status_code == 200
    with session_factory() as s:
        assert "BOPRIN-1400" not in _numbers(_rows_for(s, en_curso=True))
        o = s.get(Order, cancelled_id)
        assert o.woo_status == "cancelled"          # el estado no se tocó
        assert o.seguimiento_excluded_at is None

    # …pero «forzarlo» desde «Ver ocultos por estado» SÍ lo devuelve a la vista
    # (es el otro eje: la decisión explícita de verlo pese a su estado).
    r = http.post("/api/erp/seguimiento/force", json={"order_ids": [cancelled_id]},
                  headers=auth_headers(http, "pedidos"))
    assert r.status_code == 200, r.text
    assert r.json()["changed"] == 1
    with session_factory() as s:
        rows = _rows_for(s, en_curso=True)
        assert "BOPRIN-1400" in _numbers(rows)
        row = next(r2 for r2 in rows if r2["order_number"] == "BOPRIN-1400")
        assert row["forzado"] is True
        assert row["oculto_por_estado"] is True     # sigue oculto POR ESTADO
        # Y se sigue listando en la vista de ocultos, para poder deshacerlo.
        assert "BOPRIN-1400" in _numbers(_rows_for(s, ver_ocultos_estado=True))
        o = s.get(Order, cancelled_id)
        assert o.woo_status == "cancelled"          # el estado no se tocó

    # Deshacer el forzado lo vuelve a esconder.
    r = http.post("/api/erp/seguimiento/force",
                  json={"order_ids": [cancelled_id], "forced": False},
                  headers=auth_headers(http, "pedidos"))
    assert r.status_code == 200
    with session_factory() as s:
        assert "BOPRIN-1400" not in _numbers(_rows_for(s, en_curso=True))

    # Excluir a mano un pedido activo NO le cambia el woo_status.
    r = http.post("/api/erp/seguimiento/exclude", json={"order_ids": [active_id]},
                  headers=auth_headers(http, "pedidos"))
    assert r.status_code == 200
    with session_factory() as s:
        o = s.get(Order, active_id)
        assert o.seguimiento_excluded_at is not None
        assert o.woo_status == "processing"         # el estado no se tocó
        # Excluido a mano → fuera; y aparece en «excluidos», no en «ocultos».
        assert _numbers(_rows_for(s, ver_excluidos=True)) == {"BOPRIN-1401"}
        assert "BOPRIN-1401" not in _numbers(_rows_for(s, ver_ocultos_estado=True))


# --- Parte A: auto-anular en la ingesta al reembolsar/cancelar ---------------


def _woo_payload(store, *, woo_id: int, number: str, status: str) -> dict:
    return {"id": woo_id, "number": number, "status": status, "total": "10",
            "currency": "EUR", "date_created": "2026-09-01T00:00:00Z",
            "billing": {"email": "a@b.com"}, "line_items": [], "meta_data": [],
            "_store_slug": store.account_id}


def test_import_refunded_autocancels_web_order(session_factory) -> None:
    """Un pedido web que pasa a `refunded` en Woo se auto-anula en BoHub y sale
    de las colas; el borrado del F_PCL se encola en factusol:writes."""
    with patch(
        "app.integrations.factusol.jobs.enqueue_autocancel_order_documents",
        return_value="job-x",
    ) as enq, session_factory() as s:
        st = _store(s)
        p = _woo_payload(st, woo_id=950, number="BOPRIN-950", status="processing")
        import_woo_order(s, store=st, woo_order=p)
        s.commit()
        import_woo_order(s, store=st, woo_order={**p, "status": "refunded"})
        s.commit()
        o = s.scalar(select(Order).where(Order.external_id == "950"))
        assert o.cancelled_at is not None
        assert o.woo_status == "refunded"
    enq.assert_called_once()


def test_el_sync_marca_reembolsado_no_anulado(session_factory) -> None:
    """En adelante el sync web sella el reembolso por lo que es: motivo
    «Reembolsado…» (no «Anulado automáticamente»), y el pedido SIGUE en
    Seguimiento con su Situación «Reembolsado»."""
    from app.erp.order_cancel import AUTO_CANCEL_REASON, AUTO_REFUND_REASON  # noqa: PLC0415

    with patch(
        "app.integrations.factusol.jobs.enqueue_autocancel_order_documents",
        return_value="job-x",
    ), session_factory() as s:
        st = _store(s)
        p = _woo_payload(st, woo_id=952, number="9494", status="processing")
        import_woo_order(s, store=st, woo_order=p)
        s.commit()
        import_woo_order(s, store=st, woo_order={**p, "status": "refunded"})
        s.commit()
        o = s.scalar(select(Order).where(Order.external_id == "952"))
        assert o.cancelled_reason == AUTO_REFUND_REASON
        assert o.cancelled_reason != AUTO_CANCEL_REASON
        # Anulado para sus acciones, pero a la vista en Seguimiento.
        assert o.cancelled_at is not None
        row = next(r for r in _rows_for(s, en_curso=True) if r["id"] == o.id)
        assert row["situacion"] == "reembolsado"


def test_un_cancelado_en_la_tienda_sigue_diciendose_anulado(session_factory) -> None:
    """La otra mitad: `cancelled` sí es una anulación."""
    from app.erp.order_cancel import AUTO_CANCEL_REASON  # noqa: PLC0415

    with patch(
        "app.integrations.factusol.jobs.enqueue_autocancel_order_documents",
        return_value="job-x",
    ), session_factory() as s:
        st = _store(s)
        p = _woo_payload(st, woo_id=953, number="953", status="processing")
        import_woo_order(s, store=st, woo_order=p)
        s.commit()
        import_woo_order(s, store=st, woo_order={**p, "status": "cancelled"})
        s.commit()
        o = s.scalar(select(Order).where(Order.external_id == "953"))
        assert o.cancelled_reason == AUTO_CANCEL_REASON
        assert o.id not in {r["id"] for r in _rows_for(s, en_curso=True)}


def test_import_still_processing_does_not_autocancel(session_factory) -> None:
    """Un reembolso PARCIAL deja el pedido en `processing`: NO se auto-anula."""
    with session_factory() as s:
        st = _store(s)
        p = _woo_payload(st, woo_id=951, number="BOPRIN-951", status="processing")
        import_woo_order(s, store=st, woo_order=p)
        s.commit()
        import_woo_order(s, store=st, woo_order={**p, "total": "8"})  # sigue processing
        s.commit()
        o = s.scalar(select(Order).where(Order.external_id == "951"))
        assert o.cancelled_at is None
