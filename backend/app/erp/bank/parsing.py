"""ERP-F4-A — lectura del extracto bancario (.xlsx / .csv).

Cada banco exporta distinto. Los formatos conocidos son DICCIONARIOS
(`FORMATOS`): qué cabecera de columna es cada campo. Añadir un banco es añadir
un `Formato`, no tocar el código. La fila de columnas se encuentra sola (la
primera que trae las cabeceras de fecha, concepto e importe de algún formato).

Lo que se sabe de los extractos reales (28/10/2026):

- **Sabadell** (Streamtec, Bomedia). Cuatro filas de preámbulo y la de
  columnas en la quinta::

      Cuenta:    | 0081-0202-17-0001259030 | … | factura | pedido
      Divisa:    | EUR
      Titular:   | STREAMTEC S.L
      Selección:
      FECHA OPER | CONCEPTO | FECHA VALOR | IMPORTE | SALDO | REFERENCIA 1 | REFERENCIA 2

  La cuenta puede venir como IBAN o como CCC antiguo (banco-oficina-dc-
  cuenta); las columnas `factura` y `pedido` que rellena Bart están rotuladas
  en la PRIMERA fila, no en la de columnas; y un movimiento puede ocupar dos
  filas: la segunda, sin fecha ni importe, trae el nombre completo del
  pagador (el concepto lo corta a 26 caracteres).
- **Belfius** (MQ Europe). Cabecera en neerlandés en la primera fila
  (`Rekening`, `Boekingsdatum`, `Transactie`, `Bedrag`…). La cuenta propia
  viene en CADA fila (`Rekening`), como IBAN o en el formato belga antiguo
  «068-2453123-20»; las filas antiguas no traen `Boekingsdatum` (vale la
  fecha valor) y hay filas de continuación con solo la cuenta o una nota.

Importes en formato español («1.314,47») o anglosajón («-1,500.00»); fechas
«d/m/yyyy», ISO o de Excel. **Una fila que no se entiende no aborta el
fichero**: se importa lo que se entiende y se devuelve qué filas no y por qué.
Solo un fichero sin fila de columnas reconocible se rechaza entero.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any


@dataclass(frozen=True)
class Formato:
    """Un formato de extracto: para cada campo interno, las cabeceras que lo
    nombran en el fichero (sin distinguir mayúsculas ni espacios)."""

    nombre: str
    columnas: dict[str, tuple[str, ...]]
    #: Campos cuyas cabeceras tienen que estar TODAS para reconocer la fila
    #: de columnas.
    identifican: tuple[str, ...] = ("fecha_oper", "concepto", "importe")
    #: Lo que tiene que parecer cada campo, si trae algo: una expresión
    #: regular o `FECHA`. Una fila que no cuadra está DESCOLOCADA (en el
    #: Belfius real hay ~170 con las celdas corridas: el importe partido por la
    #: coma decimal, los céntimos en `Devies`) y no se importa con un importe
    #: falso: va a las no entendidas. Con `divisa` en la coherencia, una fila
    #: corrida se intenta recolocar (ver `_recolocar`).
    coherencia: dict[str, str] = field(default_factory=dict)
    #: Cómo empieza el texto del movimiento (el campo `concepto`), para saber,
    #: al recolocar una fila corrida, qué celda se partió por una coma.
    transaccion: str | None = None


#: `Formato.coherencia`: el campo tiene que ser una fecha.
FECHA = "fecha"
#: Las divisas: la celda que la trae marca dónde acaba una fila corrida.
_DIVISAS = frozenset({
    "EUR", "USD", "GBP", "CHF", "SEK", "DKK", "NOK", "PLN", "CZK", "HUF", "RON",
    "JPY", "CAD", "AUD", "CNY",
})


#: Campos internos. `cuenta` es la cuenta propia cuando el banco la repite en
#: cada fila (Belfius); `divisa` e `iban_pagador` se leen pero no se exigen.
#: `factura`, `presupuesto` y `pedido` son las columnas que Bart rellena a mano.
FORMATOS: tuple[Formato, ...] = (
    Formato(
        "Sabadell",
        {
            # La hoja de Bart (un export antiguo) dice «FECHA OPER»; el export
            # que da hoy la web del banco (.xls), «F. Operativa». Una variante
            # más de rótulos se añade aquí, sin tocar código.
            "fecha_oper": ("FECHA OPER", "F. OPERATIVA"),
            "concepto": ("CONCEPTO",),
            "fecha_valor": ("FECHA VALOR", "F. VALOR"),
            "importe": ("IMPORTE",),
            "saldo": ("SALDO",),
            "referencia1": ("REFERENCIA 1",),
            "referencia2": ("REFERENCIA 2",),
            "factura": ("FACTURA",),
            "presupuesto": ("PRESUPUESTO",),
            "pedido": ("PEDIDO",),
        },
    ),
    Formato(
        "Belfius",
        {
            "cuenta": ("Rekening",),
            "fecha_oper": ("Boekingsdatum",),
            "referencia1": ("Transactienummer",),
            "iban_pagador": ("Rekening tegenpartij",),
            # El nombre del pagador: hace el papel de la REFERENCIA 2 de Sabadell.
            "referencia2": ("Naam tegenpartij bevat",),
            "concepto": ("Transactie",),
            "fecha_valor": ("Valutadatum",),
            "importe": ("Bedrag",),
            "divisa": ("Devies",),
            "bic": ("BIC",),
            "factura": ("FACTURA",),
            "presupuesto": ("PRESUPUESTO",),
            "pedido": ("PEDIDO",),
        },
        # Con la divisa en su sitio, el importe es el suyo, y lo que venga
        # corrido detrás (el BIC, el país) no importa.
        coherencia={"divisa": r"[A-Z]{3}", "fecha_valor": FECHA},
        transaccion=(
            r"(?:INSTANT\s+)?(?:STORTING|OVERSCHRIJVING|EUROPESE\s+DOMICILIERING|"
            r"DOMICILIERING|BETALING|TERUGBETALING|AFHALING)\b"
        ),
    ),
)

#: Mapeo campo → cabecera del formato Sabadell (el de siempre; valor por
#: defecto de la exportación).
SABADELL_MAPPING: dict[str, str] = {k: v[0] for k, v in FORMATOS[0].columnas.items()}

#: Campos obligatorios para que una fila sea un movimiento.
REQUIRED_FIELDS = ("fecha_oper", "concepto", "importe")

#: Columnas que Bart rellena a mano y que la exportación rellena con lo
#: conciliado.
FILLED_COLUMNS = ("factura", "presupuesto", "pedido")

#: Campos que una fila de continuación puede aportar al movimiento anterior.
_CAMPOS_CONTINUACION = ("referencia1", "referencia2", "factura", "presupuesto", "pedido")


class ParseError(ValueError):
    """El fichero no se puede leer (formato no soportado, sin fila de
    columnas). Lleva el nº de fila si aplica."""

    def __init__(self, row: int | None, message: str):
        self.row = row
        super().__init__(f"fila {row}: {message}" if row else message)


@dataclass
class ParsedStatement:
    """Resultado de leer el fichero: cabecera (cuenta/divisa/titular), las
    columnas tal cual y las filas crudas (dict por columna)."""

    header: dict[str, str] = field(default_factory=dict)
    header_lines: list[str] = field(default_factory=list)
    columns: list[str] = field(default_factory=list)
    rows: list[dict[str, Any]] = field(default_factory=list)
    #: Fila (1-based) del fichero donde está la cabecera de columnas.
    columns_row: int = 0
    #: Formato reconocido y su mapeo campo → columna del fichero.
    formato: str | None = None
    mapping: dict[str, str] = field(default_factory=dict)
    #: `Formato.coherencia` y `Formato.transaccion` del formato reconocido.
    coherencia: dict[str, str] = field(default_factory=dict)
    transaccion: str | None = None


@dataclass
class Lectura:
    """Lo que sale de interpretar las filas: los movimientos, las filas que no
    se entienden (con el porqué) y cuántas filas eran continuación de otra."""

    movimientos: list[dict[str, Any]] = field(default_factory=list)
    no_entendidas: list[dict[str, Any]] = field(default_factory=list)
    continuaciones: int = 0
    #: Filas corridas que se han recolocado (un nombre partido por una coma).
    recolocadas: int = 0
    #: Cuenta propia de cada fila, cuando el banco la repite (Belfius).
    cuentas: Counter[str] = field(default_factory=Counter)


# --- números y fechas --------------------------------------------------------


def parse_amount(value: Any) -> float:
    """«1.314,47» → 1314.47 · «-1,500.00» → -1500.0 · «110,11» → 110.11 ·
    «6950.22» → 6950.22 · 2278 → 2278.0. Decide el separador decimal por el
    ÚLTIMO separador presente; con uno solo, «,» siempre es decimal y «.» lo es
    si van 2 dígitos detrás (3 dígitos = millar, «1.314»)."""
    if value is None:
        raise ValueError("importe vacío")
    if isinstance(value, (datetime, date)):
        # Una fila descolocada: el importe se fue a otra columna.
        raise ValueError(f"importe no numérico: una fecha ({texto(value)})")
    if isinstance(value, bool):
        raise ValueError(f"importe no numérico: {value!r}")
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("€", "").replace(" ", "").replace("\xa0", "")
    if not text:
        raise ValueError("importe vacío")
    negative = text.startswith("-") or text.endswith("-")
    text = text.strip("-+")
    if not re.fullmatch(r"[\d.,]+", text):
        raise ValueError(f"importe no numérico: {value!r}")
    has_dot, has_comma = "." in text, "," in text
    if has_dot and has_comma:
        decimal = "," if text.rfind(",") > text.rfind(".") else "."
        thousands = "." if decimal == "," else ","
        text = text.replace(thousands, "").replace(decimal, ".")
    elif has_comma:
        text = text.replace(",", ".")
    elif has_dot:
        head, _, tail = text.rpartition(".")
        if len(tail) == 3 and head:  # «1.314» = millar
            text = head.replace(".", "") + tail
        # si no, «6950.22» ya es decimal
    try:
        number = float(text)
    except ValueError as exc:
        raise ValueError(f"importe no numérico: {value!r}") from exc
    return -number if negative else number


def parse_date(value: Any) -> date:
    """«3/9/2026» · «25/08/2026» · «2026-09-03» · datetime/date de openpyxl."""
    if value is None or (isinstance(value, str) and not value.strip("# \t")):
        # «##»: lo que Excel enseña cuando la fecha no cabe, pegado como texto.
        raise ValueError("fecha vacía")
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10] if fmt == "%Y-%m-%d" else text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"fecha no reconocida: {value!r}")


def texto(value: Any) -> str | None:
    """Una celda como texto limpio: «271815809.0» → «271815809», fechas en ISO,
    espacios duros fuera; vacía → None."""
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    s = str(value).replace("\xa0", " ").strip()
    return s or None


def _vacio(value: Any) -> bool:
    return texto(value) is None


# --- cuenta: IBAN o número antiguo -------------------------------------------

_IBAN_RE = re.compile(r"\b([A-Z]{2}\d{2}(?:\s?[A-Z0-9]{4}){2,7}(?:\s?[A-Z0-9]{1,4})?)\b")
#: Un número de cuenta sin IBAN: CCC español «0081-0202-17-0001259030», belga
#: antiguo «068-2453123-20»… Dígitos con guiones o espacios, al menos 10.
_NUMERO_CUENTA_RE = re.compile(r"(?<![\w-])\d[\d\s-]{8,}\d(?![\w-])")


def normalize_iban(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).upper()


#: Longitud del IBAN por país (los de los bancos y clientes habituales): el
#: patrón genérico se come lo que vaya pegado detrás («…9030 FACTURA»).
_LONGITUD_IBAN = {
    "ES": 24, "BE": 16, "DE": 22, "FR": 27, "NL": 18, "IT": 27, "PT": 25,
    "GB": 22, "IE": 22, "AT": 20, "CH": 21, "LU": 20, "DK": 18, "SE": 24,
    "PL": 28, "SI": 19,
}


def extract_iban(text: str) -> str | None:
    """IBAN de una línea de cabecera («Cuenta: ES33 0081 0202 1300 0117 1918 ·
    …»). Normalizado sin espacios y recortado a la longitud de su país."""
    match = _IBAN_RE.search(str(text or "").upper())
    if not match:
        return None
    iban = normalize_iban(match.group(1))
    largo = _LONGITUD_IBAN.get(iban[:2])
    return iban[:largo] if largo and len(iban) > largo else iban


def numero_de_cuenta(value: Any) -> str | None:
    """El número de cuenta de un texto: el IBAN si lo hay (normalizado) o, si
    no, el número antiguo tal cual («0081-0202-17-0001259030»)."""
    iban = extract_iban(str(value or ""))
    if iban:
        return iban
    for match in _NUMERO_CUENTA_RE.finditer(str(value or "")):
        candidato = match.group(0).strip()
        if len(re.sub(r"\D", "", candidato)) >= 10:
            return candidato
    return None


def cuenta_casa(numero: Any, iban: Any) -> bool:
    """Si el número de cuenta de un extracto es el de esta cuenta dada de alta.
    Un IBAN se compara entero; un número antiguo, por sus dígitos: tienen que
    ser el final del IBAN (`0081-0202-17-0001259030` es
    `ES11 0081 0202 17 0001259030`; `068-2453123-20` es `BE28 0682 4531 2320`).
    No hace falta calcular el dígito de control y no depende de él."""
    n = normalize_iban(numero)
    i = normalize_iban(iban)
    if not n or not i:
        return False
    if n == i:
        return True
    if re.fullmatch(r"[A-Z]{2}\d{2}[A-Z0-9]+", n):
        return False  # dos IBAN distintos
    digitos = re.sub(r"\D", "", n)
    return len(digitos) >= 10 and re.sub(r"\D", "", i).endswith(digitos)


def _norm_col(name: Any) -> str:
    return re.sub(r"\s+", " ", str(name or "").replace("\xa0", " ").strip().upper())


def _formato_en(cells: list[Any], formatos: tuple[Formato, ...]) -> Formato | None:
    normed = {_norm_col(c) for c in cells if not _vacio(c)}
    for formato in formatos:
        if all(
            any(_norm_col(alias) in normed for alias in formato.columnas.get(campo, ()))
            for campo in formato.identifican
        ):
            return formato
    return None


def mapping_de(formato: Formato, columns: list[str]) -> dict[str, str]:
    """Campo → nombre de la columna en el fichero, solo los que están."""
    by_norm = {_norm_col(c): c for c in columns if c}
    out: dict[str, str] = {}
    for campo, aliases in formato.columnas.items():
        for alias in aliases:
            if _norm_col(alias) in by_norm:
                out[campo] = by_norm[_norm_col(alias)]
                break
    return out


def formato_de_mapping(nombre: str, mapping: dict[str, str]) -> Formato:
    """Un mapeo guardado en la cuenta (campo → cabecera) como formato más."""
    return Formato(nombre, {k: (v,) for k, v in mapping.items() if v})


def detect_mapping(columns: list[str]) -> dict[str, str] | None:
    """El mapeo del formato conocido que encaja con estas columnas, o None."""
    formato = _formato_en(list(columns), FORMATOS)
    return mapping_de(formato, list(columns)) if formato else None


# --- lectura del fichero -----------------------------------------------------


def _rows_from_xlsx(content: bytes, pestana: str | None = None) -> list[list[Any]]:
    """La rejilla de la primera hoja del libro o, con `pestana`, de esa."""
    from openpyxl import load_workbook  # noqa: PLC0415

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    if pestana:
        if pestana not in wb.sheetnames:
            raise ParseError(
                None, f"el libro no tiene la pestaña «{pestana}» (tiene: "
                f"{', '.join(wb.sheetnames)})",
            )
        ws = wb[pestana]
    else:
        ws = wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _rows_from_xls(content: bytes) -> list[list[Any]]:
    """La primera hoja de un .xls (el Excel antiguo, el que da hoy la web de
    Banc Sabadell). Las fechas salen como fechas, como en el .xlsx."""
    import xlrd  # noqa: PLC0415

    try:
        libro = xlrd.open_workbook(file_contents=content)
    except xlrd.XLRDError as exc:
        raise ParseError(None, f"no se puede leer el .xls: {exc}") from exc
    hoja = libro.sheet_by_index(0)
    filas: list[list[Any]] = []
    for r in range(hoja.nrows):
        fila: list[Any] = []
        for c in range(hoja.ncols):
            celda = hoja.cell(r, c)
            if celda.ctype == xlrd.XL_CELL_DATE:
                fila.append(xlrd.xldate_as_datetime(celda.value, libro.datemode))
            elif celda.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                fila.append(None)
            else:
                fila.append(celda.value)
        filas.append(fila)
    return filas


#: Cómo empieza un .xlsx (zip) y un .xls (documento compuesto de OLE).
_FIRMA_XLSX = b"PK\x03\x04"
_FIRMA_XLS = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _rows_from_csv(content: bytes) -> list[list[Any]]:
    """El separador se decide por la fila de columnas (la que trae el importe):
    las líneas de cabecera («Cuenta: …») no llevan separador y confunden al
    Sniffer, y los importes «1.314,47» meten comas que no son separador."""
    text = content.decode("utf-8-sig", errors="replace")
    lines = text.splitlines()
    probe = next(
        (ln for ln in lines if "IMPORTE" in ln.upper() or "BEDRAG" in ln.upper()),
        lines[0] if lines else "",
    )
    delimiter = max(";,\t|", key=probe.count) if probe else ";"
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def read_statement(
    content: bytes, filename: str, *, extra: tuple[Formato, ...] = (),
    pestana: str | None = None,
) -> ParsedStatement:
    """Lee .xls, .xlsx o .csv. Separa la cabecera (líneas antes de la fila de
    columnas, donde vive la cuenta/divisa/titular) de las filas de datos. Una
    columna sin nombre en la fila de columnas toma el rótulo que tenga encima
    en el preámbulo si es uno de los del formato (`factura`/`pedido` de
    Sabadell, en la primera fila). Un .xlsx se lee por su primera pestaña o,
    con `pestana`, por esa."""
    name = (filename or "").lower()
    # Manda lo que es el fichero, no su nombre (un banco puede llamar .xls a un
    # .xlsx); la extensión decide solo si no se reconoce.
    if content.startswith(_FIRMA_XLSX) or (
        name.endswith((".xlsx", ".xlsm")) and not content.startswith(_FIRMA_XLS)
    ):
        grid = _rows_from_xlsx(content, pestana)
    elif content.startswith(_FIRMA_XLS) or name.endswith(".xls"):
        grid = _rows_from_xls(content)
    elif name.endswith((".csv", ".txt")):
        grid = _rows_from_csv(content)
    else:
        raise ParseError(None, "formato no soportado: usa .xls, .xlsx o .csv")

    formatos = (*FORMATOS, *extra)
    parsed = ParsedStatement()
    columns_idx = -1
    formato: Formato | None = None
    for idx, row in enumerate(grid):
        formato = _formato_en(list(row), formatos)
        if formato is not None:
            columns_idx = idx
            break
        cells = [c for c in row if not _vacio(c)]
        line = " · ".join(texto(c) or "" for c in cells)
        if line:
            parsed.header_lines.append(line)
            low = line.lower()
            if low.startswith(("cuenta", "account", "rekening")):
                # Celda a celda (la primera sin el rótulo): unidas, el IBAN se
                # comería lo que va en la celda de al lado («…9030 factura»).
                primera = re.sub(r"^\s*(cuenta|account|rekening)\s*:?", "",
                                 texto(cells[0]) or "", flags=re.IGNORECASE)
                numero = next(
                    (n for n in (numero_de_cuenta(t) for t in
                                 [primera, *(texto(c) or "" for c in cells[1:])]) if n),
                    None,
                )
                if numero:
                    parsed.header["numero_cuenta"] = numero
                    if re.fullmatch(r"[A-Z]{2}\d{2}[A-Z0-9]+", numero):
                        parsed.header["iban"] = numero
                parsed.header["cuenta"] = line
            elif low.startswith("divisa"):
                parsed.header["divisa"] = line.split(":", 1)[-1].strip(" ·")
            elif low.startswith("titular"):
                parsed.header["titular"] = line.split(":", 1)[-1].strip(" ·")
    if columns_idx < 0 or formato is None:
        conocidos = " · ".join(
            f"{f.nombre} ({' / '.join(f.columnas[c][0] for c in f.identifican)})"
            for f in FORMATOS
        )
        raise ParseError(
            None, f"no se encontró la fila de columnas de ningún formato conocido: {conocidos}",
        )
    parsed.formato = formato.nombre
    parsed.columns_row = columns_idx + 1
    columns = [texto(c) or "" for c in grid[columns_idx]]
    # Rótulos del preámbulo para columnas sin nombre (factura/pedido de Sabadell).
    rotulos = {_norm_col(a): a for campo in FILLED_COLUMNS for a in formato.columnas.get(campo, ())}
    for i, col in enumerate(columns):
        if col:
            continue
        for pre in grid[:columns_idx]:
            valor = texto(pre[i]) if i < len(pre) else None
            if valor and _norm_col(valor) in rotulos:
                columns[i] = valor
                break
    # Las columnas sin rótulo se guardan por posición («#3») si traen algo, para
    # que el Excel conciliado las devuelva; las vacías del final (el Sabadell
    # trae 256) fuera.
    datos = grid[columns_idx + 1:]
    ancho = max(
        [i + 1 for i, c in enumerate(columns) if c]
        + [i + 1 for fila in datos for i, c in enumerate(fila) if not _vacio(c)]
        + [0]
    )
    columns = (columns + [""] * ancho)[:ancho]
    columns = [c or f"{COLUMNA_SIN_ROTULO}{i + 1}" for i, c in enumerate(columns)]
    parsed.columns = columns
    parsed.mapping = mapping_de(formato, columns)
    parsed.coherencia = dict(formato.coherencia)
    parsed.transaccion = formato.transaccion
    width = len(columns)
    for offset, row in enumerate(grid[columns_idx + 1 :], start=columns_idx + 2):
        if all(_vacio(c) for c in row):
            continue
        cells = list(row) + [None] * (width - len(row))
        parsed.rows.append(
            {
                "__row__": offset,
                **{col: cells[i] for i, col in enumerate(columns) if col},
            }
        )
    return parsed


#: Prefijo de las columnas sin rótulo en la fila de columnas («#3»).
COLUMNA_SIN_ROTULO = "#"


def _valores(row: dict[str, Any]) -> list[str]:
    return [t for k, v in row.items() if k != "__row__" and (t := texto(v))]


def _descolocada(coherencia: dict[str, str], valor: Any) -> str | None:
    """Por qué la fila está descolocada (un campo que no parece lo que es), o
    None si cuadra. Solo mira los campos que traen algo."""
    for campo, regla in coherencia.items():
        v = valor(campo)
        if _vacio(v):
            continue
        if regla == FECHA:
            try:
                parse_date(v)
            except ValueError as exc:
                if str(exc) == "fecha vacía":
                    continue
                return f"fila descolocada: la fecha valor trae «{texto(v)}»"
        elif not re.fullmatch(regla, (texto(v) or "").upper()):
            return f"fila descolocada: {campo} trae «{texto(v)}»"
    return None


def _es_numero(value: Any) -> bool:
    if _vacio(value) or isinstance(value, (datetime, date)):
        return False
    try:
        parse_amount(value)
    except ValueError:
        return False
    return True


def _es_fecha(value: Any) -> bool:
    try:
        parse_date(value)
    except ValueError:
        return False
    return True


def _unir(celdas: list[Any]) -> str:
    """Las celdas en que una coma partió un texto, unidas otra vez por ella."""
    return ",".join(c if isinstance(c, str) else (texto(c) or "") for c in celdas)


#: Más columnas corridas que estas no se recolocan: ya no es una coma suelta.
_MAX_CORRIDA = 4


def _plural(n: int, palabra: str) -> str:
    return f"{n} {palabra}" + ("" if n == 1 else "s")


def _recolocar(
    parsed: ParsedStatement, col: Any, row: dict[str, Any],
) -> tuple[dict[str, Any], int] | str | None:
    """Una fila CORRIDA: un texto con una coma («Streamtec, Sl») se partió en
    dos celdas al pegar el extracto y empujó el resto de la fila a la
    derecha. La pista es la divisa: la celda que la trae dice cuántas
    posiciones se corrió la fila (`k`). Delante de la divisa van el importe y
    la fecha valor. Lo que se partió es el nombre del pagador, el texto del
    movimiento, o los dos (el nombre suele ir también dentro del texto): el
    texto del movimiento empieza en la celda que empieza como uno
    (`Formato.transaccion`); lo que sobra después de ella se une al texto.
    Lo que sobra antes es del nombre si, unido, sigue estando dentro del
    texto del movimiento; si no, de la dirección (calle y número). Si ni el
    nombre está en el texto, no se adivina. Solo se corre el tramo del banco:
    las columnas que rellena Bart (`factura`, `pedido`…) no se mueven.

    Devuelve la fila recolocada y `k`; un motivo si está corrida pero no se
    puede recolocar sin adivinar (el importe también partido por la coma
    decimal: «5» puede ser ,05 o ,50); o None si la fila no está corrida (o
    el formato no tiene las columnas en el orden que esto supone)."""
    columnas = parsed.columns

    def pos(campo: str) -> int | None:
        c = col(campo)
        return len(columnas) - 1 - columnas[::-1].index(c) if c in columnas else None

    i_div, i_imp, i_fv, i_con, i_nom = (
        pos(c) for c in ("divisa", "importe", "fecha_valor", "concepto", "referencia2")
    )
    if None in (i_div, i_imp, i_fv, i_con, i_nom) or not (
        i_nom < i_con and i_fv == i_con + 1 and i_imp == i_fv + 1 and i_div == i_imp + 1
    ):
        return None
    celdas = [row.get(c) for c in columnas]
    if (texto(celdas[i_div]) or "").upper() in _DIVISAS:
        return None
    # Hasta dónde llega el tramo del banco: lo de Bart va detrás y no se corre.
    fin = min((p for p in (pos(c) for c in FILLED_COLUMNS) if p is not None and p > i_div),
              default=len(celdas))
    j = next(
        (i for i in range(i_div + 1, fin) if (texto(celdas[i]) or "").upper() in _DIVISAS),
        None,
    )
    if j is None:
        return None
    k = j - i_div
    importe, antes = celdas[j - 1], celdas[j - 2]
    valuta = celdas[j - 3]
    if (
        re.fullmatch(r"\d{1,2}", texto(importe) or "") and _es_numero(antes)
        and (_es_fecha(valuta) or not (texto(valuta) or "").strip("# "))
    ):
        motivo = f"importe partido en dos columnas («{texto(antes)}» y «{texto(importe)}»)"
        if k > 1:
            motivo += f", y la fila corrida {_plural(k - 1, 'columna')} más"
        return motivo + ": corrígelo en el fichero o vuelve a exportar ese tramo"
    corrida = f"fila corrida {_plural(k, 'columna')} (la divisa está en «{columnas[j]}»)"
    if not (_es_numero(importe) and _es_fecha(antes)):
        return corrida + ": no se sabe dónde acaba el importe; corrígela en el fichero"
    no_se_sabe = corrida + ": no se sabe qué texto se partió; corrígela en el fichero"
    if not parsed.transaccion or k > _MAX_CORRIDA:
        return no_se_sabe

    def parece_movimiento(value: Any) -> bool:
        return bool(re.match(parsed.transaccion or "", texto(value) or "", re.IGNORECASE))

    # `a`: cuántas celdas de más hay antes del texto del movimiento.
    inicios = [a for a in range(k + 1) if parece_movimiento(celdas[i_con + a])]
    if len(inicios) != 1:
        return no_se_sabe
    a = inicios[0]
    concepto = _unir(celdas[i_con + a:i_con + k + 1])
    medio = celdas[i_nom:i_con + a]       # nombre, dirección…, con `a` celdas de más
    if a:
        # Cuántas de las celdas de más son del nombre: las que, unidas, aún
        # están dentro del texto del movimiento.
        letras = _solo_letras(concepto)
        del_nombre = [x for x in range(a + 1) if _solo_letras(_unir(medio[:x + 1])) in letras]
        if not del_nombre:
            return no_se_sabe
        x = max(del_nombre)
        resto = medio[x + 1:]
        medio = [_unir(medio[:x + 1])]
        if a - x:
            medio.append(_unir(resto[:a - x + 1]))   # la calle, partida por una coma
            resto = resto[a - x + 1:]
        medio += resto
    nuevas = [
        *celdas[:i_nom], *medio, concepto,
        *celdas[i_con + k + 1:fin], *([None] * k), *celdas[fin:],
    ]
    return {"__row__": row["__row__"], **dict(zip(columnas, nuevas, strict=True))}, k


def _solo_letras(value: str) -> str:
    return re.sub(r"[\W_]+", "", value).upper()


def interpretar(parsed: ParsedStatement, mapping: dict[str, str]) -> Lectura:
    """Filas crudas → movimientos normalizados, sin abortar por una fila.

    - Fila con fecha e importe → movimiento. Sin fecha de operación pero con
      importe, vale la fecha valor (las filas antiguas de Belfius).
    - Fila sin fecha ni importe justo detrás de un movimiento → es su
      CONTINUACIÓN (Sabadell parte así el pagador): se une a él y aporta lo
      que este no tenga (REFERENCIA 2 = el pagador, factura, pedido…).
    - Cualquier otra cosa que no se entienda → `no_entendidas`, con la fila,
      el motivo y sus valores, para enseñarla; el resto se importa igual."""
    col_by_norm = {_norm_col(c): c for c in parsed.columns if c}

    def col(field_name: str) -> str | None:
        wanted = mapping.get(field_name)
        return col_by_norm.get(_norm_col(wanted)) if wanted else None

    lectura = Lectura()
    anterior: dict[str, Any] | None = None   # el movimiento de la fila anterior

    def no_entendida(row: dict[str, Any], motivo: str) -> None:
        lectura.no_entendidas.append(
            {"fila": row["__row__"], "motivo": motivo, "valores": _valores(row)}
        )

    for row in parsed.rows:
        nrow = row["__row__"]
        recolocada = 0
        if parsed.coherencia.get("divisa"):
            arreglo = _recolocar(parsed, col, row)
            if isinstance(arreglo, str):
                anterior = None
                no_entendida(row, arreglo)
                continue
            if arreglo is not None:
                row, recolocada = arreglo
        raw = {k: v for k, v in row.items() if k != "__row__"}

        def valor(field_name: str, _row: dict[str, Any] = row) -> Any:
            c = col(field_name)
            return _row.get(c) if c else None

        fecha_raw, importe_raw = valor("fecha_oper"), valor("importe")
        sin_movimiento = _vacio(fecha_raw) and _vacio(importe_raw)
        if sin_movimiento and not (_vacio(valor("concepto")) and _vacio(valor("fecha_valor"))):
            # Trae concepto o fecha valor propios: es otro movimiento al que le
            # falta la fecha y el importe, no la continuación del de encima.
            anterior = None
            no_entendida(row, "sin fecha de operación ni importe")
            continue
        if sin_movimiento:
            aporta = {f: texto(valor(f)) for f in _CAMPOS_CONTINUACION if texto(valor(f))}
            otros = [
                k for k, v in raw.items()
                if not _vacio(v) and k not in {col("cuenta"), col("divisa")}
            ]
            if anterior is not None:
                for f, v in aporta.items():
                    if not anterior.get(f):
                        anterior[f] = v
                for k, v in raw.items():
                    if not _vacio(v) and _vacio(anterior["raw"].get(k)):
                        anterior["raw"][k] = v
                anterior["continuaciones"].append(nrow)
                lectura.continuaciones += 1
            elif otros:
                no_entendida(row, "sin fecha ni importe, y no sigue a un movimiento")
            continue

        anterior = None
        descolocada = _descolocada(parsed.coherencia, valor)
        if descolocada:
            no_entendida(row, descolocada)
            continue
        try:
            importe = parse_amount(importe_raw)
        except ValueError as exc:
            no_entendida(row, str(exc))
            continue
        try:
            fecha_oper = parse_date(fecha_raw)
        except ValueError as exc:
            try:
                fecha_oper = parse_date(valor("fecha_valor"))
            except ValueError:
                no_entendida(row, str(exc))
                continue
        fecha_valor: date | None = None
        try:
            fecha_valor = parse_date(valor("fecha_valor"))
        except ValueError:
            fecha_valor = None
        saldo: float | None = None
        if not _vacio(valor("saldo")):
            try:
                saldo = parse_amount(valor("saldo"))
            except ValueError:
                saldo = None
        cuenta = texto(valor("cuenta"))
        if cuenta:
            lectura.cuentas[cuenta] += 1
        movimiento = {
            "source_row": nrow,
            "fecha_oper": fecha_oper,
            "fecha_valor": fecha_valor,
            "concepto": texto(valor("concepto")) or texto(valor("referencia2")) or "(sin concepto)",
            "importe": round(importe, 2),
            "saldo": round(saldo, 2) if saldo is not None else None,
            "cuenta": cuenta,
            "continuaciones": [],
            "raw": raw,
            #: Columnas que se corrió la fila en el fichero (0: en su sitio).
            "recolocada": recolocada,
            **{f: texto(valor(f)) for f in _CAMPOS_CONTINUACION},
        }
        if recolocada:
            lectura.recolocadas += 1
        lectura.movimientos.append(movimiento)
        anterior = movimiento
    return lectura


def apply_mapping(parsed: ParsedStatement, mapping: dict[str, str]) -> list[dict[str, Any]]:
    """Solo los movimientos (compatibilidad). Las filas que no se entienden
    están en `interpretar(...).no_entendidas`."""
    return interpretar(parsed, mapping).movimientos


def dedupe_key(
    account_id: str,
    fecha_oper: date,
    importe: float,
    concepto: str,
    saldo: float | None,
    repeticion: int = 1,
    referencia: str | None = None,
) -> str:
    """cuenta + fecha oper + importe + concepto + saldo → hash estable. Un
    rango reimportado (solapado) produce las mismas claves → no duplica.
    Sin saldo (Belfius) dos movimientos pueden ser idénticos: los distingue
    su `referencia` (el nº de transacción) o, si no la traen, la `repeticion`
    dentro del fichero (el segundo lleva «#2»; reimportar el mismo fichero da
    las mismas claves). Con saldo, dos filas idénticas son la misma."""
    base = "|".join(
        [
            account_id,
            fecha_oper.isoformat(),
            f"{importe:.2f}",
            re.sub(r"\s+", " ", concepto.strip().upper()),
            f"{saldo:.2f}" if saldo is not None else "",
        ]
    )
    if referencia:
        base += f"|ref:{referencia}"
    if repeticion > 1:
        base += f"#{repeticion}"
    return hashlib.sha256(base.encode("utf-8")).hexdigest()


# --- identidad de un movimiento -------------------------------------------------


def _concepto_comparable(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().upper())


def mismo_movimiento(a: dict[str, Any], b: dict[str, Any]) -> bool:
    """¿Son el mismo movimiento? Lo que lo identifica sin ambigüedad es la
    fecha de operación, el importe y el SALDO (dos del mismo día e importe
    tienen saldos distintos). Sin saldo en los dos, el nº de transacción
    (`referencia1`, Belfius) y, si tampoco, el concepto recortado al prefijo
    común: la hoja de Bart lo trae cortado a 26 caracteres
    («ABONO TRANSFERENCIA DE DAN») y el export de hoy entero. `a` y `b`:
    dicts con `fecha_oper`, `importe`, `saldo`, `concepto`, `referencia1`."""
    if a.get("fecha_oper") != b.get("fecha_oper"):
        return False
    if abs(float(a["importe"]) - float(b["importe"])) > 0.005:
        return False
    if a.get("saldo") is not None and b.get("saldo") is not None:
        return abs(float(a["saldo"]) - float(b["saldo"])) <= 0.005
    ra, rb = texto(a.get("referencia1")), texto(b.get("referencia1"))
    if ra and rb:
        return ra == rb
    ca, cb = _concepto_comparable(a.get("concepto")), _concepto_comparable(b.get("concepto"))
    corto = min(len(ca), len(cb))
    return corto > 0 and ca[:corto] == cb[:corto]


class Identidades:
    """Movimientos ya conocidos (los de BoHub, o los de la hoja) para saber si
    uno leído es uno de ellos (`mismo_movimiento`). Cada conocido vale para
    UNO solo: dos movimientos idénticos del fichero no se tapan con uno."""

    def __init__(self, conocidos: list[tuple[Any, dict[str, Any]]]):
        #: (fecha, céntimos) → [(posición, ident, datos)]
        self._por_dia: dict[tuple[Any, int], list[tuple[int, Any, dict[str, Any]]]] = {}
        self._posicion: dict[Any, int] = {}
        self._usados: set[int] = set()
        for i, (ident, datos) in enumerate(conocidos):
            self._por_dia.setdefault(self._clave(datos), []).append((i, ident, datos))
            self._posicion.setdefault(ident, i)

    @staticmethod
    def _clave(datos: dict[str, Any]) -> tuple[Any, int]:
        return datos.get("fecha_oper"), round(float(datos["importe"]) * 100)

    def usar(self, ident: Any) -> None:
        """Da por casado un conocido (p. ej. por la clave exacta)."""
        if ident in self._posicion:
            self._usados.add(self._posicion[ident])

    def buscar(self, m: dict[str, Any]) -> Any | None:
        """El conocido que es este movimiento (y lo da por usado), o None."""
        for i, ident, datos in self._por_dia.get(self._clave(m), []):
            if i not in self._usados and mismo_movimiento(m, datos):
                self._usados.add(i)
                return ident
        return None


# --- columnas que rellena Bart -------------------------------------------------

#: «5-260083», «2 - 526140»: serie (1 a 9) y código de factura.
_FACTURA_RE = re.compile(r"(?<![\w.,/-])([1-9])\s*-\s*(\d{4,7})(?![\w/]|-\d|[.,]\d)")
#: «526043» suelto: código de factura sin serie (la de la cuenta). Un guion
#: detrás vale si sigue una palabra («526043-parcial»), no otro número.
_CODIGO_SUELTO_RE = re.compile(r"(?<![\w.,/-])(\d{6})(?![\w/]|-\d|[.,]\d)")
#: Un número que nombra OTRO documento: no es una factura.
_OTRO_DOCUMENTO_RE = re.compile(
    r"\b(?:proforma|pf|pedido|ped|presupuesto|pres|albar[aá]n|alb|order|bestell\w*)"
    r"\.?\s*(?:n[ºo°.]*\s*)?[:#]?\s*[\d-]+",
    re.IGNORECASE,
)
#: Lo que puede quedar entre números de factura sin que sea una nota.
_SOLO_SEPARADORES_RE = re.compile(r"[\s,;/+&]*(?:\by\b[\s,;/+&]*)*", re.IGNORECASE)


@dataclass(frozen=True)
class Anotacion:
    """Lo que Bart escribió a mano en la columna `factura` de una fila."""

    texto: str
    #: Facturas que nombra, en orden: (serie, código).
    facturas: tuple[tuple[int, int], ...]
    #: Solo números de factura (nada más que separadores): es la respuesta.
    #: Con una nota al lado («parcial», «confirmar», «50 %») hay que mirarla.
    limpia: bool


def leer_anotacion(value: Any, *, serie_cuenta: int | None = None) -> Anotacion | None:
    """La columna `factura` de una fila: «5-260083», «5-260088, 5-260089»,
    «526043» (sin serie: la de la cuenta), «5-260067 (parcial, …)»… None si
    está vacía."""
    t = texto(value)
    if not t:
        return None
    if re.fullmatch(r"\d+\.0", t):   # un número que Excel guardó como «526043.0»
        t = t[:-2]
    t = t.replace("\u2013", "-").replace("\u2014", "-")   # el guion largo de Excel
    original = t
    # Si nombra otro documento («PF 5-260083, 5-260084»), no es solo la
    # respuesta: se propone para mirarlo.
    otro_documento = bool(_OTRO_DOCUMENTO_RE.search(t))
    t = _OTRO_DOCUMENTO_RE.sub(" ", t)
    facturas: list[tuple[int, int]] = [
        (int(m.group(1)), int(m.group(2))) for m in _FACTURA_RE.finditer(t)
    ]
    resto = _FACTURA_RE.sub(" ", t)
    # Códigos sueltos, sin serie: la de la primera factura con serie o la de
    # la cuenta («526043», «5-250066, 250073», «525067 - parcial»). Si la
    # celda ya nombra facturas con serie, un número suelto DENTRO de la nota
    # («cliente cita ref. 260720») no es otra factura.
    serie = facturas[0][0] if facturas else serie_cuenta
    sueltos = [int(m.group(1)) for m in _CODIGO_SUELTO_RE.finditer(resto)]
    sin_sueltos = _CODIGO_SUELTO_RE.sub(" ", resto)
    if serie and sueltos and (not facturas or _SOLO_SEPARADORES_RE.fullmatch(sin_sueltos)):
        facturas += [(serie, codigo) for codigo in sueltos]
        resto = sin_sueltos
    unicas = tuple(dict.fromkeys(facturas))
    return Anotacion(
        texto=original,
        facturas=unicas,
        limpia=bool(unicas) and not otro_documento
        and bool(_SOLO_SEPARADORES_RE.fullmatch(resto)),
    )
